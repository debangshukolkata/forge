"""Change an existing workbook and keep its formulas (D-339).

openpyxl (3.1) reads charts and pictures back and writes them again, but drops macros and the chart style and colour parts Excel adds.
`edit_workbook` therefore compares the package before and after the save and refuses only a real loss, unless the caller says the loss is
fine (`allow_loss`); the message names what would be lost. Refusing every workbook with a chart up front (the first version) pushed the
model to set `allow_loss` for a chart that was in fact kept (doc-2 w13). Formulas stay formulas: a cell is only given a value when the
operation says so."""

from __future__ import annotations

import io
import zipfile
from typing import Any, Literal

from pydantic import BaseModel, Field

from forge.sheets.build import INVALID_TITLE
from forge.sheets.read import open_workbook, package_extras
from forge.sheets.spec import WorkbookError


class Operation(BaseModel):
    op: Literal[
        "set_cells",
        "append_rows",
        "format",
        "add_sheet",
        "delete_sheet",
        "rename_sheet",
        "define_name",
        "column_width",
    ]
    sheet: str = ""
    cells: dict[str, Any] = Field(default_factory=dict, description="set_cells: {'B2': 12, 'C2': '=B2*2'}")
    rows: list[list[Any]] = Field(
        default_factory=list, description="append_rows: rows to add under the last used row"
    )
    range: str = ""
    number_format: str | None = None
    bold: bool | None = None
    fill: str | None = None
    name: str = ""
    target: str = ""
    width: float | None = None


def _sheet(workbook: Any, name: str) -> Any:
    if name not in workbook.sheetnames:
        raise WorkbookError(f"No sheet called '{name}'. Sheets: {', '.join(workbook.sheetnames)}")
    return workbook[name]


def _apply(workbook: Any, operation: Operation) -> str:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName

    kind = operation.op
    if kind == "add_sheet":
        if not operation.name or len(operation.name) > 31 or INVALID_TITLE.search(operation.name):
            raise WorkbookError(f"Sheet name {operation.name!r}: 1 to 31 characters, none of [ ] : * ? / \\")
        if operation.name in workbook.sheetnames:
            raise WorkbookError(f"A sheet called '{operation.name}' exists already")
        workbook.create_sheet(operation.name)
        return f"added sheet '{operation.name}'"
    if kind == "define_name":
        if not operation.name or not operation.target:
            raise WorkbookError("define_name needs name and target (e.g. Assumptions!$B$2)")
        workbook.defined_names[operation.name] = DefinedName(operation.name, attr_text=operation.target)
        return f"named range {operation.name} -> {operation.target}"
    sheet = _sheet(workbook, operation.sheet)
    if kind == "delete_sheet":
        if len(workbook.sheetnames) == 1:
            raise WorkbookError("A workbook needs at least one sheet")
        workbook.remove(sheet)
        return f"deleted sheet '{operation.sheet}' (formulas elsewhere that referred to it now show #REF!; run check_workbook)"
    if kind == "rename_sheet":
        if not operation.name or len(operation.name) > 31 or INVALID_TITLE.search(operation.name):
            raise WorkbookError(f"Sheet name {operation.name!r} is not allowed")
        sheet.title = operation.name
        return f"renamed '{operation.sheet}' to '{operation.name}' (formulas in other sheets are NOT updated: check them)"
    if kind == "set_cells":
        for reference, value in operation.cells.items():
            sheet[reference] = value
        return f"set {len(operation.cells)} cell(s) on '{operation.sheet}'"
    if kind == "append_rows":
        for row in operation.rows:
            sheet.append(row)
        return f"appended {len(operation.rows)} row(s) to '{operation.sheet}'"
    if kind == "column_width":
        if not operation.range or operation.width is None:
            raise WorkbookError("column_width needs range (a column letter such as B) and width")
        sheet.column_dimensions[operation.range].width = operation.width
        return f"column {operation.range} width {operation.width}"
    if kind == "format":
        if not operation.range:
            raise WorkbookError("format needs range, e.g. B2:B20")
        cells = sheet[operation.range]
        flat = (
            [c for row in cells for c in row]
            if isinstance(cells, tuple) and cells and isinstance(cells[0], tuple)
            else (list(cells) if isinstance(cells, tuple) else [cells])
        )
        for cell in flat:
            if operation.number_format:
                cell.number_format = operation.number_format
            if operation.bold is not None:
                cell.font = Font(bold=operation.bold, color=cell.font.color)
            if operation.fill:
                cell.fill = PatternFill(
                    start_color=operation.fill, end_color=operation.fill, fill_type="solid"
                )
        return f"formatted {len(flat)} cell(s) in {operation.range}"
    raise WorkbookError(f"Unknown operation {kind}")


def _chart_styling(data: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return sum(
            1 for name in archive.namelist() if name.startswith(("xl/charts/style", "xl/charts/colors"))
        )


def edit_workbook(
    data: bytes, operations: list[Operation], allow_loss: bool = False
) -> tuple[bytes, list[str]]:
    workbook = open_workbook(data)
    done = []
    for operation in operations:
        done.append(_apply(workbook, operation))
    out = io.BytesIO()
    workbook.save(out)
    edited = out.getvalue()
    charts, pictures, macros = package_extras(data)
    kept_charts, kept_pictures, _ = package_extras(edited)
    styling = _chart_styling(data) - _chart_styling(edited)
    lost = [
        text
        for text, count in (
            (f"{charts - kept_charts} chart(s)", charts - kept_charts),
            (f"{pictures - kept_pictures} picture(s)", pictures - kept_pictures),
            ("macros", macros),
            (f"the Excel styling of {styling} chart part(s) (colours and style sets)", styling),
        )
        if count > 0
    ]
    if lost and not allow_loss:
        raise WorkbookError(
            "Saving this workbook would drop "
            + " and ".join(lost)
            + " (the editing library cannot keep them). Ask the user whether that is fine, "
            "then call again with allow_loss=true; or make the change in a copy built from an outline (build_workbook)."
        )
    if lost:
        done.append("LOST on save: " + " and ".join(lost))
    elif charts or pictures:
        done.append(f"kept {charts} chart(s) and {pictures} picture(s)")
    return edited, done
