"""Native charts and print setup for a built sheet (W6, W7).

Run doc-2 showed three builder faults, not model faults: a chart without `varyColors` is drawn by Excel with one colour per point and a
legend that lists the categories as if they were series; a data range laid out as a row (a monthly balance) was read as one series per
column; and charts ran off the printed page or on top of each other. The orientation is taken from the ranges, the category count must
match the points, overlapping charts move down and every sheet prints one page wide.
"""

from __future__ import annotations

from typing import Any

from forge.sheets.spec import ChartSpec, WorkbookError

CHART_WIDTH_CM, CHART_HEIGHT_CM = 18.0, 9.0
PIXELS_PER_CM = 96 / 2.54
DEFAULT_COLUMN_CHARS, DEFAULT_ROW_PIXELS = 8.43, 20.0
PORTRAIT_PRINTABLE_PIXELS = 680  # A4 or Letter minus the margins
LANDSCAPE_PRINTABLE_PIXELS = 1000
PRINTABLE_HEIGHT_PIXELS = {"portrait": 950, "landscape": 620}  # A4 or Letter minus the top and bottom margins
ONE_PAGE_UP_TO = 1.5  # a sheet with charts at most this many pages tall is fitted on one page
GRIDLINE = "D9D9D9"


def _is_label(value: Any) -> bool:
    return isinstance(value, str) and not value.startswith("=")


def _label_left_of(sheet: Any, row: int, column: int) -> str:
    for col in range(column - 1, 0, -1):
        value = sheet.cell(row=row, column=col).value
        if _is_label(value) and value.strip():
            return str(value)
    return ""


def _span(cells: Any) -> tuple[int, int, int, int]:
    first, last = cells[0][0], cells[-1][-1]
    return first.row, first.column, last.row, last.column


def _source(sheet: Any, reference: str) -> tuple[Any, str]:
    """The sheet a chart range reads from and the bare address: "Data!B1:B7" reads the Data sheet (a dashboard's charts sit
    beside the summary while the numbers live on a data sheet; doc-5b w14, D-543), "B1:B7" the chart's own sheet."""
    name, mark, address = reference.rpartition("!")
    if not mark:
        return sheet, reference
    name = name.strip("'").replace("''", "'")
    if name not in sheet.parent.sheetnames:
        raise WorkbookError(
            f"Chart range {reference!r} names a sheet that does not exist; the sheets are: {', '.join(sheet.parent.sheetnames)}"
        )
    return sheet.parent[name], address


def build_chart(sheet: Any, chart: ChartSpec) -> Any:
    from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    chart_sheet = sheet
    sheet, data_address = _source(chart_sheet, chart.data)  # from here on `sheet` is where the values live
    top, left, bottom, right = _span(sheet[data_address])
    kinds: dict[str, Any] = {
        "column": BarChart,
        "bar": BarChart,
        "line": LineChart,
        "pie": PieChart,
        "area": AreaChart,
        "scatter": ScatterChart,
    }
    made = kinds[chart.type]()
    if chart.type == "bar":
        made.type = "bar"
    elif chart.type == "column":
        made.type = "col"
    made.title = chart.title or None
    if chart.type == "scatter":
        for column in range(left + 1, right + 1):
            xs = Reference(sheet, min_col=left, min_row=top + 1, max_row=bottom)
            ys = Reference(sheet, min_col=column, min_row=top, max_row=bottom)
            made.series.append(Series(ys, xs, title_from_data=True))
    else:
        category_sheet, category_address = (
            _source(chart_sheet, chart.categories) if chart.categories else (sheet, "")
        )
        categories = _span(category_sheet[category_address]) if chart.categories else None
        horizontal_categories = (
            categories is not None and categories[0] == categories[2] and categories[1] < categories[3]
        )
        from_rows = (top == bottom and left < right) or (horizontal_categories and top < bottom)
        lead = sheet.cell(row=top, column=left).value
        titled = _is_label(lead)
        values = Reference(sheet, min_col=left, max_col=right, min_row=top, max_row=bottom)
        made.add_data(values, from_rows=from_rows, titles_from_data=titled)
        points = (right - left + 1 if from_rows else bottom - top + 1) - (1 if titled else 0)
        if categories:
            count = (
                (categories[3] - categories[1] + 1)
                if horizontal_categories
                else (categories[2] - categories[0] + 1)
            )
            if count != points:
                raise WorkbookError(
                    f"Chart {chart.title or chart.data!r}: {count} categories ({chart.categories}) but {points} value(s) per series "
                    f"({chart.data}, series read along the {'rows' if from_rows else 'columns'}). Give both ranges the same length; "
                    "the data range may include one header cell, the categories range none."
                )
            made.set_categories(
                Reference(
                    category_sheet,
                    min_row=categories[0],
                    min_col=categories[1],
                    max_row=categories[2],
                    max_col=categories[3],
                )
            )
        if from_rows and not titled and len(made.series) == 1:
            label = _label_left_of(sheet, top, left)
            if label:
                made.series[0].tx = SeriesLabel(v=label)
    if chart.type != "pie":
        made.varyColors = False  # Excel's default colours each point and lists the categories in the legend
        if chart.x_title:
            made.x_axis.title = chart.x_title
        if chart.y_title:
            made.y_axis.title = chart.y_title
        if chart.type in ("column", "bar"):
            # A negative bar kept its colour only with invertIfNegative off (Excel drew it white and outlined, read as a stray
            # selection box), and the category labels sit at the bottom, not on the zero line inside the negative bars (doc-6 w02, D-556).
            for series in made.series:
                series.invertIfNegative = False
            made.x_axis.tickLblPos = "low"
        made.x_axis.delete = False  # recent openpyxl versions hide the axes unless told otherwise
        made.y_axis.delete = False
        if made.y_axis.majorGridlines is not None:
            made.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRIDLINE))
        if len(made.series) == 1:
            made.legend = None  # the title already names the one measure
        else:
            made.legend.position = "b"
    if chart.type == "line":
        for series in made.series:
            series.smooth = False
    if made.title is not None:
        made.title.overlay = False  # otherwise Excel draws the title across the top gridline
    made.width, made.height = CHART_WIDTH_CM, CHART_HEIGHT_CM
    return made


def _column_pixels(sheet: Any, column: int) -> float:
    from openpyxl.utils import get_column_letter

    width = sheet.column_dimensions[get_column_letter(column)].width or DEFAULT_COLUMN_CHARS
    return float(int(width * 7 + 5))


def _left_pixels(sheet: Any, column: int) -> float:
    return sum(_column_pixels(sheet, index) for index in range(1, column))


def _chart_box(sheet: Any, column: int, row: int) -> tuple[float, float, float, float]:
    left = _left_pixels(sheet, column)
    top = (row - 1) * DEFAULT_ROW_PIXELS
    return left, top, left + CHART_WIDTH_CM * PIXELS_PER_CM, top + CHART_HEIGHT_CM * PIXELS_PER_CM


def place_charts(sheet: Any, charts: list[ChartSpec]) -> list[str]:
    """Add the charts; one that would cover an earlier chart moves below it. Returns a note per moved chart."""
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    placed: list[tuple[float, float, float, float]] = []
    notes: list[str] = []
    for chart in charts:
        letters, row = coordinate_from_string(chart.anchor.replace("$", ""))
        column = column_index_from_string(letters)
        box = _chart_box(sheet, column, row)
        moved = False
        for other in placed:
            if box[0] < other[2] and other[0] < box[2] and box[1] < other[3] and other[1] < box[3]:
                row = int(other[3] // DEFAULT_ROW_PIXELS) + 2
                box = _chart_box(sheet, column, row)
                moved = True
        if moved:
            notes.append(
                f"chart {chart.title or chart.data!r} moved from {chart.anchor} to {letters}{row} so it does not cover another chart"
            )
        placed.append(box)
        sheet.add_chart(build_chart(sheet, chart), f"{letters}{row}")
    return notes


def print_setup(sheet: Any, header_row: int | None) -> None:
    """One page wide, as many pages tall as needed; landscape when the table and charts are wider than a portrait page.
    A sheet with charts that would run a little over one page is fitted on one: otherwise the page break cuts a chart and the
    second page holds a strip of it under the repeated header (doc-5 w14, D-543)."""
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    from openpyxl.worksheet.properties import PageSetupProperties

    right = _left_pixels(sheet, (sheet.max_column or 1) + 1)
    bottom = (sheet.max_row or 1) * DEFAULT_ROW_PIXELS
    for chart in sheet._charts:
        anchor = chart.anchor  # a cell name until the workbook is saved
        if isinstance(anchor, str):
            letters, row = coordinate_from_string(anchor)
            column = column_index_from_string(letters)
        else:
            column, row = anchor._from.col + 1, anchor._from.row + 1
        right = max(right, _left_pixels(sheet, column) + chart.width * PIXELS_PER_CM)
        bottom = max(bottom, (row - 1) * DEFAULT_ROW_PIXELS + chart.height * PIXELS_PER_CM)
    orientation = "landscape" if right > PORTRAIT_PRINTABLE_PIXELS else "portrait"
    width = PORTRAIT_PRINTABLE_PIXELS if orientation == "portrait" else LANDSCAPE_PRINTABLE_PIXELS
    pages_tall = bottom * min(1.0, width / right) / PRINTABLE_HEIGHT_PIXELS[orientation]
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1 if sheet._charts and pages_tall <= ONE_PAGE_UP_TO else 0
    sheet.page_setup.orientation = orientation
    sheet.page_margins.left = sheet.page_margins.right = 0.5
    if header_row is not None:
        sheet.print_title_rows = f"{header_row}:{header_row}"  # a table longer than a page repeats its header
