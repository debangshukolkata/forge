"""Build an .xlsx from a `WorkbookSpec` with openpyxl (D-339)."""

from __future__ import annotations

import contextlib
import importlib
import io
import re
from datetime import date
from typing import Any

from forge.brand.kit import Resolved
from forge.sheets.charts import place_charts, print_setup
from forge.sheets.spec import Conditional, SheetSpec, WorkbookError, WorkbookSpec
from forge.sheets.titles import TITLE_ROWS, moved_spec

INVALID_TITLE = re.compile(r"[\[\]:*?/\\]")
HEADER_FILL = "1F3864"
SUBTITLE_GREY = "595959"
ISO_DATE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
DATE_FORMULA = re.compile(r"=DATE\((\d{4}),(\d{1,2}),(\d{1,2})\)", re.IGNORECASE)
MAX_FITTED_WIDTH = 50.0  # a column sized from its content stops here; longer text wraps
CHAR_WIDTH = (
    1.1  # column-width units per character of body text (Calibri 11 runs a little wider than one unit)
)


def load_openpyxl() -> Any:
    try:
        return importlib.import_module("openpyxl")
    except ImportError as error:
        raise WorkbookError(
            "openpyxl is not installed: pip install openpyxl (the forge[docs] extra)"
        ) from error


def _check_names(spec: WorkbookSpec) -> None:
    seen: set[str] = set()
    for sheet in spec.sheets:
        if not sheet.name or len(sheet.name) > 31 or INVALID_TITLE.search(sheet.name):
            raise WorkbookError(f"Sheet name {sheet.name!r}: 1 to 31 characters, none of [ ] : * ? / \\")
        if sheet.name.lower() in seen:
            raise WorkbookError(f"Two sheets are called {sheet.name!r}")
        seen.add(sheet.name.lower())
    if not spec.sheets:
        raise WorkbookError("The workbook needs at least one sheet")


def _conditional(sheet: Any, rule: Conditional) -> None:
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
    from openpyxl.styles import Font, PatternFill

    fill = PatternFill(start_color=rule.fill, end_color=rule.fill, fill_type="solid")
    font = Font(color=rule.font) if rule.font else None
    first_cell = rule.range.split(":")[0].replace("$", "")
    made: Any
    if rule.rule == "color_scale":
        made = ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        )
    elif rule.rule == "data_bar":
        made = DataBarRule(start_type="min", end_type="max", color="5B9BD5")
    elif rule.rule == "text_contains":
        made = FormulaRule(formula=[f'ISNUMBER(SEARCH("{rule.value}",{first_cell}))'], fill=fill, font=font)
    else:
        operator = {
            "greater_than": "greaterThan",
            "less_than": "lessThan",
            "between": "between",
            "equal": "equal",
        }[rule.rule]
        formula = [str(rule.value)] + ([str(rule.value2)] if rule.rule == "between" else [])
        made = CellIsRule(operator=operator, formula=formula, fill=fill, font=font)
    sheet.conditional_formatting.add(rule.range, made)


def _build_sheet(workbook: Any, spec: SheetSpec, first: bool, brand: Resolved) -> int | None:
    """Fills one sheet; returns its header row (None without columns). Charts come later, once every sheet exists."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = spec.name
    if spec.tab_colour:
        sheet.sheet_properties.tabColor = spec.tab_colour
    fill_colour = brand.header_fill or HEADER_FILL
    header_font = Font(bold=True, color=brand.header_text or "FFFFFF", name=brand.body_font)
    header_fill = PatternFill(start_color=fill_colour, end_color=fill_colour, fill_type="solid")
    header_row = 1 + (
        TITLE_ROWS if spec.title else 0
    )  # the outline's addresses were moved to match (titles.py)
    if spec.title:
        title = sheet.cell(row=1, column=1, value=spec.title)
        title.font = Font(
            bold=True, size=16, name=brand.title_font or brand.body_font, color=brand.heading or fill_colour
        )
        sheet.row_dimensions[1].height = 24
        if spec.subtitle:
            sheet.cell(row=2, column=1, value=spec.subtitle).font = Font(
                size=11, italic=True, name=brand.body_font, color=SUBTITLE_GREY
            )
    for index, column in enumerate(spec.columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column.header)
        if not spec.table:
            cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        longest = max(
            [len(column.header)]
            + [
                len(str(row[index - 1]))
                for row in spec.rows
                if index <= len(row)
                and not str(row[index - 1]).startswith("=")  # a formula's text is not what shows
            ]
        )
        sheet.column_dimensions[get_column_letter(index)].width = column.width or max(
            10.0, min(MAX_FITTED_WIDTH, longest + 4.0)
        )
    first_row = (
        header_row + 1 if spec.columns else header_row
    )  # the header takes its row only when there is one
    for row_number, row in enumerate(spec.rows, start=first_row):
        for index, value in enumerate(row, start=1):
            if (
                isinstance(value, str)
                and index <= len(spec.columns)
                and "yy" in (spec.columns[index - 1].format or "")
                and (iso := ISO_DATE.fullmatch(value) or DATE_FORMULA.fullmatch(value.replace(" ", "")))
            ):
                # A date column's ISO text or constant =DATE(y,m,d) becomes a real date (doc-4b w07, D-541).
                with contextlib.suppress(ValueError):  # an impossible date stays as written
                    value = date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))
            cell = sheet.cell(row=row_number, column=index, value=value)
            if brand.body_font:
                cell.font = Font(name=brand.body_font, color=brand.text)
            if index <= len(spec.columns) and spec.columns[index - 1].format:
                cell.number_format = spec.columns[index - 1].format
            width = sheet.column_dimensions[get_column_letter(index)].width or 10.0
            if isinstance(value, str) and not value.startswith("=") and len(value) * CHAR_WIDTH > width:
                # Text longer than its column wraps instead of being cut off on screen and in print (doc-4 w07, w14; D-541).
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    last_row = len(spec.rows) + first_row - 1
    if spec.total_row and spec.columns:
        total = last_row + 1
        for letter, what in spec.total_row.items():
            index = column_index_from_string(letter)
            if what.isupper() and what.isalpha():
                cell = sheet.cell(
                    row=total, column=index, value=f"={what}({letter}{first_row}:{letter}{last_row})"
                )
                if index <= len(spec.columns) and spec.columns[index - 1].format:
                    cell.number_format = spec.columns[index - 1].format
            else:
                cell = sheet.cell(row=total, column=index, value=what)
            cell.font = Font(bold=True)
    if spec.table and spec.columns and spec.rows:
        reference = f"A{header_row}:{get_column_letter(len(spec.columns))}{last_row}"
        table = Table(displayName=re.sub(r"\W", "_", spec.name)[:30] or "Table1", ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    if spec.freeze:
        sheet.freeze_panes = spec.freeze
    elif spec.columns:
        sheet.freeze_panes = f"A{header_row + 1}"
    for validation in spec.validations:
        if validation.choices:
            made = DataValidation(
                type="list", formula1='"' + ",".join(validation.choices) + '"', allow_blank=True
            )
        elif validation.whole:
            made = DataValidation(
                type="whole",
                operator="between",
                formula1=str(validation.whole[0]),
                formula2=str(validation.whole[1]),
            )
        elif validation.decimal:
            made = DataValidation(
                type="decimal",
                operator="between",
                formula1=str(validation.decimal[0]),
                formula2=str(validation.decimal[1]),
            )
        else:
            raise WorkbookError(f"Validation for {validation.range} needs choices, whole or decimal")
        if validation.message:
            made.error, made.errorTitle, made.showErrorMessage = validation.message, "Not allowed", True
        made.add(validation.range)
        sheet.add_data_validation(made)
    for rule in spec.conditional:
        _conditional(sheet, rule)
    for (
        reference,
        shape,
    ) in spec.formats.items():  # after the column formats and the total row, so a single cell wins
        cells = sheet[reference]
        for row in cells if isinstance(cells, tuple) else ((cells,),):
            for cell in row if isinstance(row, tuple) else (row,):
                cell.number_format = shape
    for reference, text in spec.notes.items():
        sheet[reference].comment = Comment(text, "DebCraft")
    return header_row if spec.columns else None


def build_workbook(spec: WorkbookSpec, notes: list[str] | None = None) -> bytes:
    """The .xlsx bytes; `notes` collects what the builder changed on its own (a chart moved off another)."""
    openpyxl = load_openpyxl()
    _check_names(spec)
    workbook = openpyxl.Workbook()
    brand = Resolved(spec.brand)
    try:
        spec = moved_spec(spec)  # a titled sheet's table starts below its title
        header_rows = [
            _build_sheet(workbook, sheet, index == 0, brand) for index, sheet in enumerate(spec.sheets)
        ]
        for sheet, header_row in zip(spec.sheets, header_rows, strict=True):
            made = workbook[sheet.name]
            moved = place_charts(made, sheet.charts)
            print_setup(made, header_row=header_row)
            if notes is not None:
                notes.extend(moved)
        from openpyxl.workbook.defined_name import DefinedName

        for name, target in spec.names.items():
            workbook.defined_names[name] = DefinedName(name, attr_text=target)
        if spec.title:
            workbook.properties.title = spec.title
        if spec.mark_inputs:
            from forge.sheets.inputs import mark_inputs

            mark_inputs(workbook)
    except WorkbookError:
        raise
    except Exception as error:  # a bad range or format in the outline
        raise WorkbookError(
            f"The workbook could not be built: {type(error).__name__}: {str(error)[:300]}"
        ) from error
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
