"""Charts and tables inside a template's sample slides (TODO P2.23, D-333): copy the sample, keep its design, put the new data in.

A company template often has a finished chart (a donut with its colours, a dark bar chart) and a finished table (its header colours,
its stripes). The data of a chart is replaced in place (python-pptx keeps the colours, fonts and layout of the chart); a table gets
the rows and columns it needs by cloning its own last row and last column, so the borders and fills carry on."""

from __future__ import annotations

from copy import deepcopy
from typing import Any

from forge.slides.open_deck import DeckError

FAMILIES = {
    "column": ("COLUMN",),
    "bar": ("BAR",),
    "line": ("LINE",),
    "area": ("AREA",),
    "pie": ("PIE",),
    "donut": ("DOUGHNUT",),
}
KIND_FAMILY = {
    "column": "column",
    "stacked_column": "column",
    "bar": "bar",
    "stacked_bar": "bar",
    "line": "line",
    "area": "area",
    "stacked_area": "area",
    "pie": "pie",
    "donut": "donut",
}


def chart_family(chart_type_name: str) -> str:
    """'column', 'bar', 'line', 'area', 'pie', 'donut' or 'other' from python-pptx's chart type name (COLUMN_CLUSTERED, DOUGHNUT, ...)."""
    name = chart_type_name.upper()
    for family, markers in FAMILIES.items():
        if any(marker in name for marker in markers):
            return family
    return "other"


def requested_family(kind: str) -> str:
    return KIND_FAMILY.get(kind, "column")


def fill_chart(shape: Any, categories: list[str], series: list[tuple[str, list[float]]]) -> None:
    """Replaces the data of the chart in `shape`; the chart's look stays."""
    from pptx.chart.data import CategoryChartData

    if not getattr(shape, "has_chart", False) or not shape.has_chart:
        raise DeckError(f'"{shape.name}" is not a chart')
    for name, values in series:
        if len(values) != len(categories):
            raise DeckError(
                f"series '{name}' has {len(values)} values but there are {len(categories)} categories"
            )
    factory: Any = CategoryChartData
    data = factory()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    shape.chart.replace_data(data)


def _set_cell(cell: Any, text: str) -> None:
    """Writes the text into the cell's first run so its font stays; other paragraphs and runs go."""
    frame = cell.text_frame
    paragraphs = frame.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    first = frame.paragraphs[0]
    runs = first.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        first.add_run().text = text


def fill_table(shape: Any, columns: list[str], rows: list[list[Any]]) -> None:
    """Makes the table in `shape` hold `columns` as its header and `rows` below it, growing or shrinking it like a person would."""
    if not getattr(shape, "has_table", False) or not shape.has_table:
        raise DeckError(f'"{shape.name}" is not a table')
    table = shape.table
    wanted_rows, wanted_columns = len(rows) + 1, len(columns)
    if wanted_columns < 1 or any(len(row) > wanted_columns for row in rows):
        raise DeckError("a row has more cells than there are columns")
    tbl = table._tbl
    grid = tbl.tblGrid
    while len(grid.gridCol_lst) < wanted_columns:  # a new column copies the last one
        last = grid.gridCol_lst[-1]
        grid.append(deepcopy(last))
        for row in tbl.tr_lst:
            row.append(deepcopy(row.tc_lst[-1]))
    while len(grid.gridCol_lst) > wanted_columns:
        grid.remove(grid.gridCol_lst[-1])
        for row in tbl.tr_lst:
            row.remove(row.tc_lst[-1])
    for column in grid.gridCol_lst:  # the table keeps its width: the columns share it
        column.w = shape.width // wanted_columns
    rows_xml = tbl.tr_lst
    while len(rows_xml) < wanted_rows:  # a new row copies the last body row (its stripes and borders)
        clone = deepcopy(rows_xml[-1])
        rows_xml[-1].addnext(clone)
        shape.height = shape.height + clone.h
        rows_xml = tbl.tr_lst
    while len(rows_xml) > wanted_rows and len(rows_xml) > 1:
        shape.height = max(shape.height - rows_xml[-1].h, 1)
        tbl.remove(rows_xml[-1])
        rows_xml = tbl.tr_lst
    for column, text in enumerate(columns):
        _set_cell(table.cell(0, column), str(text))
    for index, row in enumerate(rows, start=1):
        for column in range(wanted_columns):
            value = row[column] if column < len(row) else ""
            _set_cell(table.cell(index, column), "" if value is None else str(value))
