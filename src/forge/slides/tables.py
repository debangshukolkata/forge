"""Tables at full strength (TODO P2.5, D-285): merged cells, column widths, per-column alignment and number formats, per-cell
fill, colour, bold and size, four border styles, heat-map shading, and splitting a long table over several slides. Cells are
plain values or `{"text": ..., "bold": true, "fill": "accent", "colspan": 2}` objects. Text colour follows the fill
automatically (dark or light, whichever reads), so a custom fill never produces unreadable text."""

from __future__ import annotations

import re
from typing import Any

from forge.slides import fit
from forge.slides.errors import BuildError
from forge.slides.themes import mix, readable_on

MAX_TABLE_ROWS = 10
BORDER_WIDTH_PT = 0.75
NUMBER_FORMAT = re.compile(r"^(?P<prefix>[^#0,.%]*)(?P<number>[#0,]*(?:\.0+)?)(?P<percent>%?)(?P<suffix>.*)$")


def format_number(value: float, pattern: str) -> str:
    """Excel-like number formats: `#,##0`, `0.0`, `0%`, `0.0%`, `$#,##0.00`, `0.0"k"`; anything else is shown as given."""
    if not pattern:
        return f"{value:g}" if isinstance(value, float) else str(value)
    match = NUMBER_FORMAT.match(pattern.replace('"', ""))
    if not match or not match.group("number"):
        return str(value)
    number = match.group("number")
    decimals = len(number.split(".")[1]) if "." in number else 0
    shown = value * 100 if match.group("percent") else value
    text = f"{shown:,.{decimals}f}" if "," in number else f"{shown:.{decimals}f}"
    return f"{match.group('prefix')}{text}{match.group('percent')}{match.group('suffix')}"


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def _cell_parts(cell: Any) -> dict[str, Any]:
    """A cell as a dict: plain values become {"text": value}."""
    if isinstance(cell, dict):
        return dict(cell)
    if hasattr(cell, "model_dump"):
        return dict(cell.model_dump(exclude_none=True))
    return {"text": cell}


def lay_out_rows(rows: list[list[Any]], columns: int, where: str) -> list[list[dict[str, Any] | None]]:
    """The rows as a grid of `columns` entries each: a cell dict where a cell starts, None where a merged cell covers it.
    Checks that every row fills exactly `columns` columns, counting column and row spans."""
    covered: dict[tuple[int, int], bool] = {}
    grid: list[list[dict[str, Any] | None]] = []
    for row_number, row in enumerate(rows):
        line: list[dict[str, Any] | None] = [None] * columns
        column = 0
        for cell in row:
            while column < columns and covered.get((row_number, column)):
                column += 1
            parts = _cell_parts(cell)
            colspan, rowspan = int(parts.get("colspan", 1)), int(parts.get("rowspan", 1))
            if column + colspan > columns:
                raise BuildError(f"{where}: row {row_number + 1} has more cells than the {columns} columns")
            line[column] = parts
            for r in range(row_number, row_number + rowspan):
                for c in range(column, column + colspan):
                    if (r, c) != (row_number, column):
                        covered[(r, c)] = True
            column += colspan
        while column < columns and covered.get((row_number, column)):
            column += 1
        if column != columns:
            raise BuildError(
                f"{where}: every table row needs {columns} cells (one per column, merged cells count as several)"
            )
        grid.append(line)
    return grid


# --- borders ---


def _border(cell: Any, side: str, colour: str | None) -> None:
    """One border of a cell (side L, R, T or B): a thin line in `colour`, or none."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    properties = cell._tc.get_or_add_tcPr()
    tag = f"a:ln{side}"
    for old in properties.findall(qn(tag)):
        properties.remove(old)
    line = OxmlElement(tag)
    if colour:
        line.set("w", str(round(BORDER_WIDTH_PT * 12700)))
        fill = OxmlElement("a:solidFill")
        fill.append(OxmlElement("a:srgbClr"))
        fill[0].set("val", colour)
        line.append(fill)
    else:
        line.set("w", "0")
        line.append(OxmlElement("a:noFill"))
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    followers = {qn(t) for t in order[order.index(tag) + 1 :]}
    for child in properties:
        if child.tag in followers or child.tag in (qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill")):
            child.addprevious(line)
            return
    properties.append(line)


def _style_borders(b: Any, table: Any, rows: int, columns: int, style: str) -> None:
    line = b.theme.muted
    for r in range(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            vertical = line if style == "grid" else None
            horizontal = line if style in ("lines", "grid") else None
            _border(cell, "L", vertical)
            _border(cell, "R", vertical)
            _border(cell, "T", horizontal if r > 0 or style == "grid" else None)
            _border(cell, "B", horizontal)


# --- the table ---


def add_table(b: Any, slide: Any, options: Any, box: tuple[float, float, float, float], descr: str) -> None:
    """A native table in `box` (x, y, width, maximum height) built from `options` (a TableBase)."""
    x, y, width, area_h = box
    where = f"slide {b.report.number}"
    columns = list(options.columns)
    ncols = len(columns)
    grid = lay_out_rows(list(options.rows), ncols, where)
    if len(grid) > MAX_TABLE_ROWS:
        b.warn(
            f"{len(grid)} table rows: more than {MAX_TABLE_ROWS} is hard to read, split the table (`split_rows`)"
        )
    if options.col_widths and len(options.col_widths) != ncols:
        raise BuildError(f"{where}: col_widths needs one number per column ({ncols})")
    if options.align and len(options.align) != ncols:
        raise BuildError(f"{where}: align needs one entry per column ({ncols})")
    if options.formats and len(options.formats) != ncols:
        raise BuildError(f"{where}: formats needs one entry per column ({ncols})")
    size = options.size or (24 if len(grid) <= 4 else 20 if len(grid) <= 6 else 16 if len(grid) <= 8 else 14)

    # the text of every cell (numbers formatted per column) and its numeric value for the heat map
    texts: list[list[str | None]] = []
    numbers: list[list[float | None]] = []
    for row in grid:
        texts_row: list[str | None] = []
        numbers_row: list[float | None] = []
        for column, parts in enumerate(row):
            if parts is None:
                texts_row.append(None)
                numbers_row.append(None)
                continue
            value = parts.get("text", "")
            pattern = options.formats[column] if options.formats else ""
            texts_row.append(format_number(value, pattern) if _is_number(value) else str(value))
            numbers_row.append(float(value) if _is_number(value) else None)
        texts.append(texts_row)
        numbers.append(numbers_row)

    # column alignment: given, else numbers right, text left
    aligns = (
        list(options.align)
        if options.align
        else [
            "right" if any(numbers[r][c] is not None for r in range(len(grid))) else "left"
            for c in range(ncols)
        ]
    )
    widths_weight = list(options.col_widths) or [
        max(
            [
                len(columns[c]),
                *(
                    len(texts[r][c] or "")
                    for r in range(len(grid))
                    if (cell := grid[r][c]) is not None and cell.get("colspan", 1) == 1
                ),
            ]
        )
        + 4
        for c in range(ncols)
    ]
    widths = [width * w / sum(widths_weight) for w in widths_weight]
    needed = 0.0
    for row_number, header in enumerate([None, *grid]):
        cells = columns if header is None else [t or "" for t in texts[row_number - 1]]
        lines = max(
            fit.lines_needed(cell, widths[i] - 0.3, size, header is None) for i, cell in enumerate(cells)
        )
        needed += lines * size * 1.2 / 72 + 0.22
    if needed > area_h:
        b.warn("the table is too tall for the slide: remove rows or shorten the cell text")
    frame = slide.shapes.add_table(
        len(grid) + 1, ncols, *(b.m["inches"](v) for v in (x, y, width, min(needed, area_h)))
    )
    frame._element.nvGraphicFramePr.cNvPr.set("descr", descr)
    table = frame.table
    for index, column_width in enumerate(widths):
        table.columns[index].width = b.m["inches"](column_width)

    # heat map: shade numeric cells of the chosen columns between two colours
    shade: dict[tuple[int, int], str] = {}
    if options.heatmap is not None:
        chosen = _heatmap_columns(options.heatmap.columns, columns)
        low, high = _hex(b, options.heatmap.low), _hex(b, options.heatmap.high)
        for c in chosen:  # each column is shaded against its own range
            found = [(r, v) for r in range(len(grid)) if (v := numbers[r][c]) is not None]
            if len(found) < 2:
                continue
            lo, hi = min(v for _, v in found), max(v for _, v in found)
            if hi == lo:
                continue
            for r, value in found:
                shade[(r, c)] = mix(low, high, (value - lo) / (hi - lo))

    header_fill = _hex(b, options.header_fill) if options.header_fill else b.theme.accent
    header_text = readable_on(header_fill)
    for column, name in enumerate(columns):
        cell = table.cell(0, column)
        _fill_cell(b, cell, header_fill)
        _write(b, cell, name, size, header_text, True, False, aligns[column])
    for r, row in enumerate(grid):
        for c, parts in enumerate(row):
            if parts is None:
                continue
            cell = table.cell(r + 1, c)
            if options.style == "striped":
                base = b.theme.surface if r % 2 == 0 else b.theme.background
            else:
                base = b.theme.background
            fill = shade.get((r, c)) or (_hex(b, parts["fill"]) if parts.get("fill") else base)
            _fill_cell(b, cell, fill)
            text_colour = (
                _hex(b, parts["color"])
                if parts.get("color")
                else readable_on(fill)
                if fill != base or options.style != "striped"
                else b.theme.text
            )
            bold = bool(parts.get("bold")) or (options.first_column_bold and c == 0)
            _write(
                b,
                cell,
                texts[r][c] or "",
                float(parts.get("size") or size),
                text_colour,
                bold,
                bool(parts.get("italic")),
                parts.get("align") or aligns[c],
            )
    for r, row in enumerate(grid):  # merges after the content is in place
        for c, parts in enumerate(row):
            if parts and (parts.get("colspan", 1) > 1 or parts.get("rowspan", 1) > 1):
                table.cell(r + 1, c).merge(
                    table.cell(r + int(parts.get("rowspan", 1)), c + int(parts.get("colspan", 1)) - 1)
                )
    _style_borders(b, table, len(grid) + 1, ncols, options.style)


def _hex(b: Any, token: str) -> str:
    from forge.slides.freeform import resolve

    return resolve(b.theme, token)


def _heatmap_columns(chosen: list[Any], columns: list[str]) -> set[int]:
    if not chosen:
        return set(range(len(columns)))
    found: set[int] = set()
    for item in chosen:
        if isinstance(item, int) and 0 <= item < len(columns):
            found.add(item)
        elif isinstance(item, str) and item in columns:
            found.add(columns.index(item))
    return found


def _fill_cell(b: Any, cell: Any, colour: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = b.colour(colour)
    cell.margin_left = cell.margin_right = b.m["inches"](0.12)
    cell.vertical_anchor = b.m["anchor"].MIDDLE


def _write(
    b: Any, cell: Any, text: str, size: float, colour: str, bold: bool, italic: bool, align: str
) -> None:
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = {
        "left": b.m["align"].LEFT,
        "center": b.m["align"].CENTER,
        "right": b.m["align"].RIGHT,
    }[align]
    run = paragraph.add_run()
    run.text = text
    b.style(run, size, colour, bold=bold, italic=italic)
