"""A quick look at a large company template (TODO P2.23, D-331): few files, easy to send back.

`forge template-look template.potx` (or `scripts/dev/template_look.py`) opens the template and writes, in one folder:

- `summary.txt`: about 30 lines: file size, slide size, how many masters, layouts and sample slides; for every master its colour scheme
  (all twelve colours as hex) and its title and body fonts; the fonts the sample slides use and which of them this computer lacks; what
  the sample slides hold (pictures, charts, tables), hidden slides and the layouts used most;
- `palettes.png`: one picture with the colour chips of every master's theme, named;
- `layouts-1.png` ...: contact sheets of one slide per layout, round-robin across the masters so every master shows (at most 40 layouts),
  with the layout names listed in `summary.txt` by number;
- `samples-1.png` ...: contact sheets of the template's own sample slides (at most 40, evenly spread when there are more);
- `template-look.zip`: all of these in one file.

The pictures are made by PowerPoint (Windows); without it the summary and the palettes are still written. Nothing is sent anywhere."""

from __future__ import annotations

import re
import shutil
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from forge.slides.contact_sheet import make_contact_sheets
from forge.slides.installed_fonts import missing_fonts
from forge.slides.open_deck import open_deck
from forge.slides.preview import PreviewUnavailable, available, export_slides, status

MAX_LAYOUTS = 40
MAX_SAMPLES = 40
COLOUR_SLOTS = (
    "dk1",
    "lt1",
    "dk2",
    "lt2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
SHEET_COLUMNS, SHEET_SIZE = 4, 20


@dataclass
class MasterInfo:
    number: int
    layouts: list[str]
    theme_name: str = ""
    colours: dict[str, str] = field(default_factory=dict)
    heading_font: str = ""
    body_font: str = ""


@dataclass
class Look:
    folder: Path
    summary: str = ""
    files: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _theme(master: Any) -> tuple[str, dict[str, str], str, str]:
    from lxml import etree  # type: ignore[import-untyped]
    from pptx.opc.constants import RELATIONSHIP_TYPE

    try:
        root = etree.fromstring(master.part.part_related_by(RELATIONSHIP_TYPE.THEME).blob)
    except (KeyError, ValueError, etree.XMLSyntaxError):
        return "", {}, "", ""
    colours: dict[str, str] = {}
    for slot in COLOUR_SLOTS:
        node = root.find(f".//a:clrScheme/a:{slot}/*", NS)
        if node is not None:
            colours[slot] = (node.get("val") if node.tag.endswith("srgbClr") else node.get("lastClr")) or "?"
    major = root.find(".//a:majorFont/a:latin", NS)
    minor = root.find(".//a:minorFont/a:latin", NS)
    scheme = root.find(".//a:clrScheme", NS)
    return (
        scheme.get("name", "") if scheme is not None else "",
        colours,
        major.get("typeface", "") if major is not None else "",
        minor.get("typeface", "") if minor is not None else "",
    )


def read_masters(prs: Any) -> list[MasterInfo]:
    found = []
    for number, master in enumerate(prs.slide_masters, start=1):
        name, colours, major, minor = _theme(master)
        found.append(
            MasterInfo(number, [layout.name for layout in master.slide_layouts], name, colours, major, minor)
        )
    return found


def _sample_facts(prs: Any) -> tuple[int, int, Counter[str], Counter[str], int, int, int]:
    """(slides, hidden, layouts used, fonts used, with pictures, with charts, with tables)."""
    layouts: Counter[str] = Counter()
    fonts: Counter[str] = Counter()
    hidden = pictures = charts = tables = 0
    for slide in prs.slides:
        layouts[slide.slide_layout.name] += 1
        if slide._element.get("show") == "0":
            hidden += 1
        kinds = set()
        for shape in slide.shapes:
            if shape.shape_type == 13:
                kinds.add("picture")
            if getattr(shape, "has_chart", False) and shape.has_chart:
                kinds.add("chart")
            if getattr(shape, "has_table", False) and shape.has_table:
                kinds.add("table")
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += 1
        pictures += "picture" in kinds
        charts += "chart" in kinds
        tables += "table" in kinds
    return len(prs.slides), hidden, layouts, fonts, pictures, charts, tables


def _inches(value: int) -> str:
    return f"{value / 914400:.2f}"


def summarize(path: Path, prs: Any, masters: list[MasterInfo], layout_order: list[tuple[int, str]]) -> str:
    size_mb = path.stat().st_size / (1024 * 1024)
    slides, hidden, used, fonts, pictures, charts, tables = _sample_facts(prs)
    lines = [
        f"{path.name}: {size_mb:.1f} MB, slide size {_inches(prs.slide_width)} x {_inches(prs.slide_height)} in, "
        f"{len(masters)} master(s), {sum(len(m.layouts) for m in masters)} layouts, {slides} sample slide(s) ({hidden} hidden)",
    ]
    for master in masters:
        palette = " ".join(f"{slot}={master.colours.get(slot, '?')}" for slot in COLOUR_SLOTS)
        lines.append(
            f"master {master.number} ({len(master.layouts)} layouts): theme '{master.theme_name}', fonts {master.heading_font} / {master.body_font}"
        )
        lines.append(f"  colours: {palette}")
    wanted = [m.heading_font for m in masters] + [m.body_font for m in masters] + list(fonts)
    lacking = missing_fonts([f for f in wanted if f and not f.startswith("+")])
    lines.append(
        "fonts used by the samples: "
        + (", ".join(f"{n} ({c})" for n, c in fonts.most_common(8)) or "none set on runs (theme fonts)")
    )
    lines.append(
        "fonts NOT installed on this computer: "
        + (", ".join(lacking) if lacking else "none (or not Windows)")
    )
    lines.append(
        f"samples with pictures {pictures}, charts {charts}, tables {tables}; layouts used most: "
        + ", ".join(f"{n} x{c}" for n, c in used.most_common(6))
    )
    lines.append(
        "layout pictures (layouts-N.png), numbered across the sheets: "
        + "; ".join(f"{i}={master}:{name}" for i, (master, name) in enumerate(layout_order, start=1))
    )
    return "\n".join(lines)


def _layout_deck(template: Path, target: Path) -> list[tuple[int, str]]:
    """A copy of the template that has one slide per layout (round-robin over the masters, at most MAX_LAYOUTS); returns (master, layout)."""
    prs = open_deck(template)
    entries = prs.slides._sldIdLst
    for entry in list(entries):
        prs.part.drop_rel(entry.rId)
        entries.remove(entry)
    per_master = [
        [(m_index, layout) for layout in master.slide_layouts]
        for m_index, master in enumerate(prs.slide_masters, start=1)
    ]
    chosen: list[tuple[int, Any]] = []
    position = 0
    while len(chosen) < MAX_LAYOUTS and any(position < len(group) for group in per_master):
        for group in per_master:
            if position < len(group) and len(chosen) < MAX_LAYOUTS:
                chosen.append(group[position])
        position += 1
    order = []
    for master_number, layout in chosen:
        slide = prs.slides.add_slide(layout)
        order.append((master_number, layout.name))
        for placeholder in slide.placeholders:
            kind = placeholder.placeholder_format.type
            label = (
                f"{layout.name}"
                if kind is not None and "TITLE" in kind.name
                else f"{kind.name.lower() if kind is not None else 'text'} {placeholder.placeholder_format.idx}"
            )
            if getattr(placeholder, "has_text_frame", False) and placeholder.has_text_frame:
                placeholder.text_frame.text = re.sub(r"[^\w .,&/-]", "", label)[:60]
    prs.save(str(target))
    return order


def _sample_deck(template: Path, target: Path) -> int:
    """The template's own slides, at most MAX_SAMPLES (evenly spread when there are more); returns how many."""
    prs = open_deck(template)
    total = len(prs.slides)
    if total == 0:
        return 0
    step = max(1, -(-total // MAX_SAMPLES))
    keep = set(range(1, total + 1, step))
    entries = prs.slides._sldIdLst
    for index, entry in reversed(list(enumerate(list(entries), start=1))):
        if index not in keep:
            prs.part.drop_rel(entry.rId)
            entries.remove(entry)
    prs.save(str(target))
    return len(prs.slides)


def _palettes_picture(masters: list[MasterInfo], target: Path) -> None:
    from PIL import Image, ImageDraw

    chip, gap, label_width, row_height = 64, 8, 260, 110
    width = label_width + len(COLOUR_SLOTS) * (chip + gap) + gap
    image = Image.new("RGB", (width, len(masters) * row_height + gap), (245, 245, 245))
    draw = ImageDraw.Draw(image)
    for row, master in enumerate(masters):
        y = gap + row * row_height
        draw.text((gap, y + 6), f"master {master.number}: {master.theme_name[:30]}", fill=(20, 20, 20))
        draw.text((gap, y + 26), f"{master.heading_font[:26]} / {master.body_font[:26]}", fill=(60, 60, 60))
        for column, slot in enumerate(COLOUR_SLOTS):
            value = master.colours.get(slot, "")
            x = label_width + column * (chip + gap)
            try:
                colour = tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))
            except ValueError:
                colour = (255, 255, 255)
            draw.rectangle((x, y, x + chip, y + chip), fill=colour, outline=(120, 120, 120))
            draw.text((x, y + chip + 4), slot, fill=(40, 40, 40))
            draw.text((x, y + chip + 18), value, fill=(40, 40, 40))
    image.save(target)


def _brightness(picture: Path) -> float:
    """Mean luminance of a rendered slide, 0 (black) to 255 (white)."""
    from PIL import Image, ImageStat

    with Image.open(picture) as image:
        return float(ImageStat.Stat(image.convert("L").resize((64, 36))).mean[0])


def _looks(value: float) -> str:
    return "dark" if value < 90 else "light" if value > 170 else "mid"


def _sheets(deck: Path, folder: Path, prefix: str, look: Look) -> list[float]:
    """Contact sheets of the deck's pictures; returns each slide's brightness (so the summary can say which layouts are dark)."""
    pictures = export_slides(deck, folder / f"_{prefix}")
    pictures.sort(key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)))  # type: ignore[union-attr]
    brightness = [_brightness(p) for p in pictures]
    for sheet in make_contact_sheets(
        pictures, folder, columns=SHEET_COLUMNS, per_sheet=SHEET_SIZE, prefix=prefix
    ):
        look.files.append(sheet.name)
    shutil.rmtree(folder / f"_{prefix}", ignore_errors=True)
    return brightness


def run_look(template: Path, folder: Path) -> Look:
    folder.mkdir(parents=True, exist_ok=True)
    look = Look(folder)
    prs = open_deck(template)
    masters = read_masters(prs)
    order: list[tuple[int, str]] = []
    scratch = folder / "_work"
    scratch.mkdir(exist_ok=True)
    try:
        order = _layout_deck(template, scratch / "layouts.pptx")
        sample_count = _sample_deck(template, scratch / "samples.pptx")
    except (
        Exception
    ) as error:  # a template with something unusual: the summary and palettes are still worth having
        look.notes.append(f"could not prepare the picture decks ({type(error).__name__}: {str(error)[:80]})")
        sample_count = 0
    look.summary = summarize(template, prs, masters, order)
    (folder / "summary.txt").write_text(look.summary + "\n", encoding="utf-8")
    look.files.append("summary.txt")
    try:
        _palettes_picture(masters, folder / "palettes.png")
        look.files.append("palettes.png")
    except ImportError:
        look.notes.append("Pillow is missing: no palettes picture")
    if available() and order:
        try:
            brightness = _sheets(scratch / "layouts.pptx", folder, "layouts", look)
            counts = Counter(_looks(v) for v in brightness)
            look.summary += (
                "\nlayout backgrounds as rendered: "
                + ", ".join(f"{n} {k}" for k, n in counts.most_common())
                + "; by number: "
                + " ".join(f"{i}={_looks(v)}" for i, v in enumerate(brightness, start=1))
            )
            (folder / "summary.txt").write_text(look.summary + "\n", encoding="utf-8")
            if sample_count:
                _sheets(scratch / "samples.pptx", folder, "samples", look)
        except (PreviewUnavailable, OSError) as error:
            look.notes.append(f"PowerPoint pictures were not made: {str(error)[:120]}")
    elif not available():
        look.notes.append(
            "PowerPoint cannot be driven from here (" + status()[1] + "): no layout or sample pictures"
        )
    shutil.rmtree(scratch, ignore_errors=True)
    archive = folder / "template-look.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for name in look.files:
            bundle.write(folder / name, name)
    look.files.append(archive.name)
    return look
