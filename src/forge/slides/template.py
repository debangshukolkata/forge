"""Reading a company template (TODO P2.11, D-288): its slide size, theme colours and fonts, and what each layout is for. A layout is
classified by the placeholders it has (title slide, section header, title and content, two content, comparison, title only,
blank, picture); `describe_template` explains the result to the model, and `template_build.py` uses the same facts to map DebCraft's
layouts onto the template. Only the standard placeholder kinds are interpreted; anything unusual is listed as "other"."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any
from xml.etree import ElementTree as ET

from forge.slides.errors import BuildError
from forge.slides.themes import Theme, mix, readable_on

EMU = 914400
CANVAS_WIDTH = 13.333  # DebCraft's internal canvas (inches); template slides are scaled to it
ASPECT = 16 / 9
ASPECT_TOLERANCE = 0.02
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FOOTER_KINDS = {"DATE", "FOOTER", "SLIDE_NUMBER"}
BODY_KINDS = {"OBJECT", "BODY"}


@dataclass(frozen=True)
class PlaceholderInfo:
    idx: int
    kind: str
    name: str
    x: float
    y: float
    w: float
    h: float


@dataclass(frozen=True)
class LayoutInfo:
    index: int
    name: str
    role: str
    placeholders: list[PlaceholderInfo]


@dataclass
class TemplateInfo:
    width: float
    height: float
    theme: Theme
    heading_font: str
    body_font: str
    layouts: list[LayoutInfo]
    roles: dict[str, int] = field(default_factory=dict)  # role -> index of the first layout with it

    @property
    def widescreen(self) -> bool:
        return abs(self.width / self.height - ASPECT) / ASPECT <= ASPECT_TOLERANCE

    @property
    def scale(self) -> float:
        """Template inches per canvas inch."""
        return self.width / CANVAS_WIDTH


def _placeholder(shape: Any) -> PlaceholderInfo:
    kind = shape.placeholder_format.type
    return PlaceholderInfo(
        idx=shape.placeholder_format.idx,
        kind=kind.name if kind else "?",
        name=shape.name,
        x=(shape.left or 0) / EMU,
        y=(shape.top or 0) / EMU,
        w=(shape.width or 0) / EMU,
        h=(shape.height or 0) / EMU,
    )


def name_is(name: str, wanted: str, cutoff: float = 0.8) -> bool:
    """True when the layout name contains the words `wanted`, spelling mistakes allowed ('Ttile Slide_Black with full wave' is a title slide)."""
    import difflib

    words = re.findall(r"[a-z]+", name.lower())
    size = len(wanted.split())
    return any(
        difflib.SequenceMatcher(None, " ".join(words[i : i + size]), wanted).ratio() >= cutoff
        for i in range(len(words) - size + 1)
    )


def classify_layout(layout: Any) -> tuple[str, list[PlaceholderInfo]]:
    """The role of one layout from its placeholders and name, and its placeholders (footer kinds left out)."""
    holders = [
        _placeholder(p)
        for p in layout.placeholders
        if p.placeholder_format.type is None or p.placeholder_format.type.name not in FOOTER_KINDS
    ]
    kinds = [h.kind for h in holders]
    name = layout.name.lower()
    titles = [k for k in kinds if k in ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE")]
    bodies = [k for k in kinds if k in BODY_KINDS]
    pictures = kinds.count("PICTURE")
    if "CENTER_TITLE" in kinds and "SUBTITLE" in kinds:
        role = "title"
    elif name_is(name, "title slide"):
        role = "title"  # a cover named by its designer (misspellings included), whatever placeholders it has
    elif "vertical" in name or "VERTICAL" in " ".join(kinds):
        role = "other"
    elif pictures:
        role = "picture"
    elif "caption" in name:
        role = "other"
    elif titles and kinds.count("BODY") >= 2 and kinds.count("OBJECT") >= 2:
        role = "comparison"
    elif titles and len(bodies) == 2:
        role = "two_content"
    elif (
        titles
        and len(bodies) == 1
        and ("section" in name or "divider" in name or ("header" in name and "OBJECT" not in kinds))
    ):
        role = "section"
    elif titles and len(bodies) == 1 and len(holders) == 2:
        role = "content"
    elif titles and not bodies and len(holders) == 1:
        role = "title_only"
    elif not holders:
        role = "blank"
    elif titles and "section" in name:
        role = "section"
    else:
        role = "other"
    return role, holders


def _theme_xml(prs: Any) -> ET.Element | None:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        part = prs.slide_master.part.part_related_by(RT.THEME)
        return ET.fromstring(part.blob)
    except Exception:  # a template without a readable theme part
        return None


def _scheme_colours(root: ET.Element | None) -> dict[str, str]:
    colours: dict[str, str] = {}
    if root is None:
        return colours
    scheme = root.find(f".//{NS_A}clrScheme")
    if scheme is None:
        return colours
    for child in scheme:
        name = child.tag.removeprefix(NS_A)
        for inner in child:
            value = inner.get("val") if inner.tag == f"{NS_A}srgbClr" else inner.get("lastClr")
            if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                colours[name] = value.upper()
    return colours


def _fonts(root: ET.Element | None) -> tuple[str, str]:
    heading, body = "Calibri", "Calibri"
    if root is not None:
        major = root.find(f".//{NS_A}majorFont/{NS_A}latin")
        minor = root.find(f".//{NS_A}minorFont/{NS_A}latin")
        heading = (major.get("typeface") if major is not None else None) or heading
        body = (minor.get("typeface") if minor is not None else None) or body
    return heading, body


def _color_map(prs: Any) -> dict[str, str]:
    """The master's colour map: which scheme colour plays background and text (a dark master swaps them)."""
    mapping = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
    try:
        element = prs.slide_master._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap"
        )
        if element is not None:
            mapping.update(
                {key: element.get(key) for key in ("bg1", "tx1", "bg2", "tx2") if element.get(key)}
            )
    except Exception:  # fall back to the usual light background
        pass
    return mapping


def theme_from_template(prs: Any) -> tuple[Theme, str, str]:
    """DebCraft's Theme built from the template's colour scheme and fonts: its background, text and accents."""
    root = _theme_xml(prs)
    scheme = _scheme_colours(root)
    mapping = _color_map(prs)
    background = scheme.get(mapping["bg1"], "FFFFFF")
    text = scheme.get(mapping["tx1"], "1B2430")
    accents = [scheme[k] for k in (f"accent{i}" for i in range(1, 7)) if k in scheme] or ["1F5FBF"]
    while len(accents) < 6:
        accents.append(accents[len(accents) % len(accents)])
    surface = mix(background, text, 0.07)
    muted = mix(text, background, 0.35)
    heading, body = _fonts(root)
    theme = Theme(
        name="template",
        background=background,
        surface=surface,
        text=text,
        muted=muted,
        accent=accents[0],
        accent_text=readable_on(accents[0]),
        palette=tuple(accents[:6]),
        title_font=heading,
        body_font=body,
        dark=sum(int(background[i : i + 2], 16) for i in (0, 2, 4)) < 3 * 110,
    )
    return theme, heading, body


def analyze_template(prs: Any) -> TemplateInfo:
    """Everything DebCraft needs to know about a template presentation (python-pptx object)."""
    theme, heading, body = theme_from_template(prs)
    layouts: list[LayoutInfo] = []
    roles: dict[str, int] = {}
    try:
        slide_layouts = list(prs.slide_layouts)
    except KeyError as error:  # a master lists a layout whose part is missing: seen in a LibreOffice regression file (the reference-deck sweep, D-399)
        raise BuildError(
            "The presentation is damaged: its slide master lists a layout that is not in the file. Open it in PowerPoint and save it again (or use another file)."
        ) from error
    for index, layout in enumerate(slide_layouts):
        role, holders = classify_layout(layout)
        layouts.append(LayoutInfo(index, layout.name, role, holders))
        roles.setdefault(role, index)
    return TemplateInfo(
        width=prs.slide_width / EMU,
        height=prs.slide_height / EMU,
        theme=theme,
        heading_font=heading,
        body_font=body,
        layouts=layouts,
        roles=roles,
    )


ROLE_USE = {
    "title": "the opening slide (title and subtitle)",
    "section": "a divider between parts",
    "content": "a title with one body area: bullets, or a chart, table or picture placed in the body area",
    "two_content": "a title with two side-by-side areas: compare two things",
    "comparison": "two headed columns for a side-by-side comparison",
    "title_only": "a title with a free area: charts, tables, pictures, diagrams, designed slides",
    "blank": "no placeholders: a fully designed slide",
    "picture": "a picture with a caption",
    "other": "a special layout: check it with a preview before using it",
}


def content_box(info: TemplateInfo) -> tuple[float, float, float, float]:
    """The free area for charts, tables and pictures, in canvas inches: the body placeholder of a content layout, else the space
    under the title placeholder of a title-only layout."""
    for role in ("content", "title_only"):
        index = info.roles.get(role)
        if index is None:
            continue
        layout = info.layouts[index]
        body = next((p for p in layout.placeholders if p.kind in BODY_KINDS), None)
        title = next((p for p in layout.placeholders if p.kind in ("TITLE", "CENTER_TITLE")), None)
        s = info.scale
        if body is not None:
            return body.x / s, body.y / s, body.w / s, body.h / s
        if title is not None:
            top = title.y + title.h + 0.15 * s
            return title.x / s, top / s, title.w / s, max(info.height - top - 0.6 * s, 1.0 * s) / s
    margin = 0.8
    return margin, 1.95, CANVAS_WIDTH - 2 * margin, 4.8


def describe_template(info: TemplateInfo, name: str = "the template") -> str:
    """The analysis as text for the model: size, colours, fonts, and one line per layout with its role and placeholders."""
    t = info.theme
    lines = [
        f"{name}: {info.width:.2f} x {info.height:.2f} in"
        + (
            ""
            if info.widescreen
            else " (NOT widescreen 16:9: build_presentation with this template is not supported; use it with edit_presentation)"
        ),
        f"Theme colours: background {t.background}, text {t.text}, accent {t.accent}, chart colours {', '.join(t.palette)}; dark background: {'yes' if t.dark else 'no'}",
        f"Fonts: headings {info.heading_font}, text {info.body_font}",
        f"Layouts ({len(info.layouts)}), with what each is for:",
    ]
    for layout in info.layouts:
        holders = (
            ", ".join(f"{p.idx}:{p.kind.lower()} ({p.w:.1f}x{p.h:.1f} in)" for p in layout.placeholders)
            or "none"
        )
        lines.append(
            f'  {layout.index}. "{layout.name}" [{layout.role}] {ROLE_USE.get(layout.role, "")}. Placeholders: {holders}'
        )
    x, y, w, h = content_box(info)
    lines.append(
        f"Free area for charts, tables and pictures: {w:.1f} x {h:.1f} in at ({x:.1f}, {y:.1f}) on a 13.3 x 7.5 canvas."
    )
    missing = [r for r in ("title", "content", "title_only") if r not in info.roles]
    if missing:
        lines.append(
            "Missing layout roles (DebCraft falls back to a title-only or blank layout): "
            + ", ".join(missing)
        )
    return "\n".join(lines)
