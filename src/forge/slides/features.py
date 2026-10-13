"""Slide features (TODO P2.8, D-316): hidden slides, transitions, basic entrance animations, sections and comments.

All of them are small pieces of XML that python-pptx does not model, so they are written here directly:

- hidden slide: `show="0"` on the slide element;
- transition: a `p:transition` element after `p:clrMapOvr` (fade, push, wipe, split, cover or none; speed; optional auto-advance);
- entrance animation: one effect per call in the slide's main sequence (appear, fade or wipe), on click, with the previous effect or
  after it;
- sections: the PowerPoint 2010 `p14:sectionLst` extension of the presentation, rebuilt from "name starts at slide N" pairs;
- comments: the classic comment parts (`commentAuthors.xml`, `comments/commentN.xml`) which every PowerPoint version opens; modern
  comments (`ppt/comments/modernComment_*.xml`) are read, not written."""

from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

from lxml import etree  # type: ignore[import-untyped]

from forge.slides.open_deck import DeckError

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
P = "{" + NS_P + "}"
TRANSITIONS = {"fade": "fade", "push": "push", "wipe": "wipe", "split": "split", "cover": "cover", "none": ""}
SPEEDS = {"slow": "slow", "medium": "med", "fast": "fast"}
EFFECTS = ("appear", "fade", "wipe")
STARTS = ("on_click", "with_previous", "after_previous")
EFFECT_MS = 500
COMMENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.comments+xml"
AUTHORS_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"


def _xp(node: Any, expression: str) -> Any:
    return etree.XPath(expression, namespaces={"p": NS_P, "a": NS_A})(node)


def _xml(text: str) -> Any:
    from pptx.oxml import (
        parse_xml,
    )  # the parser python-pptx uses, so the new nodes answer xpath with the usual prefixes

    return parse_xml(text)


# --- hidden slides ---


def set_hidden(slide: Any, hidden: bool) -> None:
    if hidden:
        slide._element.set("show", "0")
    elif "show" in slide._element.attrib:
        del slide._element.attrib["show"]


def is_hidden(slide: Any) -> bool:
    return bool(slide._element.get("show") == "0")


def copy_features(source: Any, target: Any) -> None:
    """A duplicated slide keeps being hidden, its transition and its animations (shape ids are copied as they are)."""
    from copy import deepcopy

    if is_hidden(source):
        set_hidden(target, True)
    for tag in ("transition", "timing"):
        node = source._element.find(f"{P}{tag}")
        if node is not None:
            _drop(target, tag)
            _insert_after_overrides(target, deepcopy(node))


# --- transitions ---


def _drop(slide: Any, tag: str) -> None:
    for node in slide._element.findall(f"{P}{tag}"):
        slide._element.remove(node)


def _insert_after_overrides(slide: Any, element: Any) -> None:
    """A slide's children are in a fixed order: cSld, clrMapOvr, transition, timing, extLst."""
    root = slide._element
    anchor = root.find(f"{P}clrMapOvr")
    if anchor is None:
        anchor = root.find(f"{P}cSld")
    if element.tag == f"{P}timing":
        later = root.find(f"{P}transition")
        anchor = later if later is not None else anchor
    anchor.addnext(element)


def set_transition(slide: Any, kind: str, speed: str = "medium", advance_after: float | None = None) -> None:
    if kind not in TRANSITIONS:
        raise DeckError(f"unknown transition {kind!r}; choose one of: {', '.join(TRANSITIONS)}")
    if speed not in SPEEDS:
        raise DeckError(f"unknown speed {speed!r}; choose one of: {', '.join(SPEEDS)}")
    _drop(slide, "transition")
    if kind == "none":
        return
    attributes = f'spd="{SPEEDS[speed]}"'
    if advance_after is not None:
        attributes += f' advClick="1" advTm="{round(advance_after * 1000)}"'
    _insert_after_overrides(
        slide, _xml(f'<p:transition xmlns:p="{NS_P}" {attributes}><p:{TRANSITIONS[kind]}/></p:transition>')
    )


def transition_of(slide: Any) -> str:
    node = slide._element.find(f"{P}transition")
    if node is None:
        return ""
    names = [child.tag.split("}")[1] for child in node if child.tag.startswith(P)]
    return names[0] if names else "unknown"


# --- entrance animations ---


def _main_sequence(slide: Any, create: bool) -> Any:
    nodes = _xp(slide._element, "./p:timing//p:cTn[@nodeType='mainSeq']/p:childTnLst")
    if nodes or not create:
        return nodes[0] if nodes else None
    _drop(slide, "timing")
    timing = _xml(
        f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst/></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
        "</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
    )
    _insert_after_overrides(slide, timing)
    return _xp(slide._element, "./p:timing//p:cTn[@nodeType='mainSeq']/p:childTnLst")[0]


def _next_id(slide: Any) -> int:
    ids = [int(i) for i in _xp(slide._element, "./p:timing//p:cTn/@id")]
    return max([*ids, 2]) + 1


def _effect_xml(first_id: int, shape_id: int, effect: str, node_type: str, delay: int) -> str:
    target = f'<p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>'
    visible = (
        f'<p:set><p:cBhvr><p:cTn id="{first_id + 1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
        f"{target}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
        '<p:to><p:strVal val="visible"/></p:to></p:set>'
    )
    preset, subtype, filter_name = {
        "appear": (1, 0, ""),
        "fade": (10, 0, "fade"),
        "wipe": (22, 8, "wipe(left)"),
    }[effect]
    motion = (
        f'<p:animEffect transition="in" filter="{filter_name}"><p:cBhvr><p:cTn id="{first_id + 2}" dur="{EFFECT_MS}"/>{target}</p:cBhvr></p:animEffect>'
        if filter_name
        else ""
    )
    return (
        f'<p:par xmlns:p="{NS_P}"><p:cTn id="{first_id}" presetID="{preset}" presetClass="entr" presetSubtype="{subtype}" fill="hold" '
        f'nodeType="{node_type}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{visible}{motion}</p:childTnLst></p:cTn></p:par>'
    )


def _group(first_id: int, inner: str, delay: str) -> str:
    return (
        f'<p:par xmlns:p="{NS_P}"><p:cTn id="{first_id}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
        f"<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>"
    )


def add_animation(slide: Any, shape: Any, effect: str = "fade", start: str = "on_click") -> None:
    """Appends one entrance effect for `shape` to the slide's animations."""
    if effect not in EFFECTS:
        raise DeckError(f"unknown effect {effect!r}; choose one of: {', '.join(EFFECTS)}")
    if start not in STARTS:
        raise DeckError(f"unknown start {start!r}; choose one of: {', '.join(STARTS)}")
    sequence = _main_sequence(slide, create=True)
    first = _next_id(slide)
    clicks = list(sequence)
    if start == "on_click" or not clicks:
        outer_id = first
        effect_xml = _effect_xml(first + 2, shape.shape_id, effect, "clickEffect", 0)
        inner = _group(first + 1, effect_xml.replace(f' xmlns:p="{NS_P}"', "", 1), "0")
        sequence.append(_xml(_group(outer_id, inner.replace(f' xmlns:p="{NS_P}"', "", 1), "indefinite")))
        return
    click = clicks[-1]
    inner_list = click.find(f"{P}cTn/{P}childTnLst")
    if start == "with_previous":
        last = inner_list[-1]
        effect_xml = _effect_xml(first, shape.shape_id, effect, "withEffect", 0)
        last.find(f"{P}cTn/{P}childTnLst").append(_xml(effect_xml))
        return
    delay = EFFECT_MS * len(inner_list)
    effect_xml = _effect_xml(first + 1, shape.shape_id, effect, "afterEffect", 0)
    inner_list.append(_xml(_group(first, effect_xml.replace(f' xmlns:p="{NS_P}"', "", 1), str(delay))))


def clear_animations(slide: Any) -> int:
    count = len(_xp(slide._element, "./p:timing//p:cTn[@presetClass]"))
    _drop(slide, "timing")
    return count


def animations_of(slide: Any) -> list[tuple[str, str, str]]:
    """(shape name, effect, start) for each entrance effect, in order."""
    names = {s.shape_id: s.name for s in slide.shapes}
    presets = {1: "appear", 10: "fade", 22: "wipe"}
    starts = {"clickEffect": "on_click", "withEffect": "with_previous", "afterEffect": "after_previous"}
    found = []
    for node in _xp(slide._element, "./p:timing//p:cTn[@presetClass='entr']"):
        target = _xp(node, ".//p:spTgt/@spid")
        shape = names.get(int(target[0]), "?") if target else "?"
        found.append(
            (
                shape,
                presets.get(int(node.get("presetID", "0")), "other"),
                starts.get(node.get("nodeType", ""), "on_click"),
            )
        )
    return found


# --- sections ---


@dataclass(frozen=True)
class Section:
    name: str
    first_slide: int  # 1-based


def _section_extension(prs: Any, create: bool) -> Any:
    root = prs.part._element
    extension_list = root.find(f"{P}extLst")
    if extension_list is None:
        if not create:
            return None
        extension_list = etree.SubElement(root, f"{P}extLst")
    for extension in extension_list.findall(f"{P}ext"):
        if extension.get("uri") == SECTION_URI:
            return extension
    if not create:
        return None
    return etree.SubElement(extension_list, f"{P}ext", uri=SECTION_URI)


def set_sections(prs: Any, sections: list[Section]) -> None:
    """Replaces all sections. The first one must start at slide 1 and the starts must rise; an empty list removes the sections."""
    total = len(prs.slides)
    if not sections:
        extension = _section_extension(prs, create=False)
        if extension is not None:
            extension.getparent().remove(extension)
        return
    if sections[0].first_slide != 1:
        raise DeckError("the first section must start at slide 1")
    starts = [s.first_slide for s in sections]
    if starts != sorted(set(starts)) or starts[-1] > total:
        raise DeckError(f"section starts must rise and stay within the {total} slides: got {starts}")
    ids = [entry.get("id") for entry in prs.slides._sldIdLst]
    extension = _section_extension(prs, create=True)
    for child in list(extension):
        extension.remove(child)
    listing = etree.SubElement(extension, f"{{{NS_P14}}}sectionLst", nsmap={"p14": NS_P14})
    for index, section in enumerate(sections):
        end = sections[index + 1].first_slide - 1 if index + 1 < len(sections) else total
        element = etree.SubElement(
            listing, f"{{{NS_P14}}}section", name=section.name, id="{" + str(uuid.uuid4()).upper() + "}"
        )
        slide_ids = etree.SubElement(element, f"{{{NS_P14}}}sldIdLst")
        for slide_id in ids[section.first_slide - 1 : end]:
            etree.SubElement(slide_ids, f"{{{NS_P14}}}sldId", id=slide_id)


def read_sections(prs: Any) -> list[Section]:
    extension = _section_extension(prs, create=False)
    if extension is None:
        return []
    position = {entry.get("id"): number for number, entry in enumerate(prs.slides._sldIdLst, start=1)}
    found = []
    for element in extension.iter(f"{{{NS_P14}}}section"):
        members = [
            position[i.get("id")] for i in element.iter(f"{{{NS_P14}}}sldId") if i.get("id") in position
        ]
        if members:
            found.append(Section(element.get("name", ""), min(members)))
    return found


def keep_sections_valid(prs: Any) -> None:
    """After slides were deleted, added or moved: every slide belongs to a section in the order of the deck. A slide without a
    section joins the one before it; a section without slides disappears."""
    extension = _section_extension(prs, create=False)
    if extension is None:
        return
    owner: dict[str, str] = {}
    for element in extension.iter(f"{{{NS_P14}}}section"):
        for member in element.iter(f"{{{NS_P14}}}sldId"):
            owner[member.get("id")] = element.get("name", "")
    runs: list[tuple[str, list[str]]] = []
    previous = next(iter(owner.values()), "Default Section")
    for entry in prs.slides._sldIdLst:
        name = owner.get(entry.get("id"), previous)
        previous = name
        if runs and runs[-1][0] == name:
            runs[-1][1].append(entry.get("id"))
        else:
            runs.append((name, [entry.get("id")]))
    for child in list(extension):
        extension.remove(child)
    listing = etree.SubElement(extension, f"{{{NS_P14}}}sectionLst", nsmap={"p14": NS_P14})
    for name, members in runs:
        element = etree.SubElement(
            listing, f"{{{NS_P14}}}section", name=name, id="{" + str(uuid.uuid4()).upper() + "}"
        )
        slide_ids = etree.SubElement(element, f"{{{NS_P14}}}sldIdLst")
        for member in members:
            etree.SubElement(slide_ids, f"{{{NS_P14}}}sldId", id=member)


# --- comments ---


@dataclass(frozen=True)
class Comment:
    slide: int
    author: str
    text: str
    when: str


def _parts(prs: Any) -> list[Any]:
    seen: dict[str, Any] = {}
    pending = [prs.part]
    while pending:
        part = pending.pop()
        for relation in part.rels.values():
            if relation.is_external or str(relation.target_part.partname) in seen:
                continue
            seen[str(relation.target_part.partname)] = relation.target_part
            pending.append(relation.target_part)
    return list(seen.values())


def read_comments(prs: Any) -> list[Comment]:
    """Classic and modern comments of every slide."""
    found: list[Comment] = []
    authors: dict[str, str] = {}
    for part in _parts(prs):
        name = str(part.partname)
        if name.endswith("/commentAuthors.xml") or (
            name.endswith("/authors.xml") and b"authorLst" in part.blob[:600]
        ):
            for node in etree.fromstring(part.blob):
                authors[node.get("id", "")] = node.get("name", "")
    for number, slide in enumerate(prs.slides, start=1):
        for relation in slide.part.rels.values():
            if relation.is_external or not relation.reltype.endswith("/comments"):
                continue
            for node in etree.fromstring(relation.target_part.blob):
                text = "".join(node.xpath(".//*[local-name()='text' or local-name()='txBody']//text()"))
                author = authors.get(node.get("authorId", ""), node.get("authorId", ""))
                found.append(Comment(number, author, text.strip(), node.get("dt", node.get("created", ""))))
    return found


def add_comment(prs: Any, slide_number: int, text: str, author: str = "DebCraft") -> None:
    """Adds a classic comment to the slide (a note for the reader of the deck, shown in the Comments pane)."""
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI

    if not 1 <= slide_number <= len(prs.slides):
        raise DeckError(f"slide {slide_number} does not exist: the deck has {len(prs.slides)} slides")
    slide = prs.slides[slide_number - 1]
    authors_part = next((p for p in _parts(prs) if str(p.partname) == "/ppt/commentAuthors.xml"), None)
    initials = "".join(word[0] for word in author.split()[:2]).upper() or "D"
    if authors_part is None:
        blob = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:cmAuthorLst xmlns:p="{NS_P}"/>'.encode()
        authors_part = Part(PackURI("/ppt/commentAuthors.xml"), AUTHORS_TYPE, prs.part.package, blob)
        prs.part.relate_to(
            authors_part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors"
        )
    authors = etree.fromstring(authors_part.blob)
    match = next((a for a in authors if a.get("name") == author), None)
    if match is None:
        match = etree.SubElement(
            authors,
            f"{P}cmAuthor",
            id=str(len(authors)),
            name=author,
            initials=initials,
            lastIdx="0",
            clrIdx=str(len(authors)),
        )
    index = int(match.get("lastIdx", "0")) + 1
    match.set("lastIdx", str(index))
    authors_part._blob = etree.tostring(authors, xml_declaration=True, encoding="UTF-8", standalone=True)
    comment = (
        f'<p:cm authorId="{match.get("id")}" dt="{datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%S.000")}" idx="{index}">'
        '<p:pos x="10" y="10"/><p:text>' + _escape(text) + "</p:text></p:cm>"
    )
    existing = next(
        (
            r.target_part
            for r in slide.part.rels.values()
            if not r.is_external and r.reltype.endswith("/comments")
        ),
        None,
    )
    if existing is None:
        number = (
            len(
                {
                    str(p.partname)
                    for p in _parts(prs)
                    if re.search(r"/comments/comment\d+\.xml$", str(p.partname))
                }
            )
            + 1
        )
        blob = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:cmLst xmlns:p="{NS_P}">{comment}</p:cmLst>'.encode()
        part = Part(PackURI(f"/ppt/comments/comment{number}.xml"), COMMENT_TYPE, prs.part.package, blob)
        slide.part.relate_to(
            part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
        )
        return
    listing = etree.fromstring(existing.blob)
    listing.append(_xml(comment.replace("<p:cm ", f'<p:cm xmlns:p="{NS_P}" ', 1)))
    existing._blob = etree.tostring(listing, xml_declaration=True, encoding="UTF-8", standalone=True)


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def apply_slide_features(
    prs: Any, slide: Any, spec: Any, deck_transition: str | None, sections: list[Section]
) -> None:
    """The outline's per-slide features for the slide just built (`sections` collects the section starts for the end)."""
    if spec.hidden:
        set_hidden(slide, True)
    kind = spec.transition or deck_transition
    if kind and kind != "none":
        set_transition(slide, kind)
    if spec.section:
        sections.append(Section(spec.section, len(prs.slides)))
    if spec.animate:
        title = slide.shapes.title
        floor = title.top + title.height if title is not None else 0
        content = [s for s in slide.shapes if not s.is_placeholder and s.top is not None and s.top >= floor]
        for position, shape in enumerate(content):
            add_animation(slide, shape, spec.animate, "on_click" if position == 0 else "with_previous")
