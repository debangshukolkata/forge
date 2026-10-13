"""Turns an outline into a real, editable .pptx with python-pptx (D-240). Every slide uses PowerPoint's own
"Title Only" layout, so each slide has a real title (outline view, navigation and screen readers work), text
sits in real text boxes, charts are native charts and tables are native tables: nothing is a picture of text.

Text is sized to its box before it is written (`fit.py`); what still cannot fit becomes a warning that names
the slide, so the model can split or shorten it."""

from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any

from forge.errors import ForgeError
from forge.slides import charts, fit, lint
from forge.slides.errors import BuildError
from forge.slides.freeform import build_freeform
from forge.slides.spec import (
    BulletGroup,
    BulletsSlide,
    ChartSlide,
    ClosingSlide,
    Column,
    Deck,
    FreeformSlide,
    ImageSlide,
    QuoteSlide,
    SectionSlide,
    StatsSlide,
    TableSlide,
    TitleSlide,
    TwoColumnSlide,
)
from forge.slides.themes import THEMES, Brand, Theme, apply_brand

WIDTH, HEIGHT = 13.333, 7.5  # inches, 16:9
MARGIN = 0.8
CONTENT_WIDTH = WIDTH - 2 * MARGIN
TITLE_TOP, TITLE_HEIGHT = 0.5, 1.15
BODY_TOP, BODY_BOTTOM = 1.95, 6.75
MIN_BODY_PT = 16  # smaller text is not readable from the back of a room
MAX_BULLETS, MAX_BULLET_WORDS, MAX_TABLE_ROWS = 6, 18, 10
BULLET, SUB_BULLET = chr(0x2022), chr(0x2013)


class SlidesUnavailable(ForgeError):
    """python-pptx is not installed."""


@dataclass
class SlideReport:
    number: int
    layout: str
    title: str
    warnings: list[str] = field(default_factory=list)


@dataclass
class BuiltDeck:
    data: bytes
    slides: list[SlideReport]
    deck_warnings: list[str] = field(
        default_factory=list
    )  # about the deck as a whole (rhythm, text-only, ...)


ReadImage = Callable[[str], bytes]


def _modules() -> Any:
    try:
        import pptx
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.enum.dml import MSO_LINE
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Emu, Inches, Pt
    except ImportError as error:
        raise SlidesUnavailable(
            "python-pptx is not installed here. Install it with: pip install python-pptx (or forge[slides])."
        ) from error
    return {
        "pptx": pptx,
        "chart_data": CategoryChartData,
        "rgb": RGBColor,
        "chart_type": XL_CHART_TYPE,
        "legend": XL_LEGEND_POSITION,
        "shape": MSO_SHAPE,
        "connector": MSO_CONNECTOR,
        "line_style": MSO_LINE,
        "anchor": MSO_ANCHOR,
        "align": PP_ALIGN,
        "qn": qn,
        "emu": Emu,
        "inches": Inches,
        "pt": Pt,
    }


class _Builder:
    def __init__(self, deck: Deck, read_image: ReadImage) -> None:
        self.m = _modules()
        self.deck, self.read_image = deck, read_image
        self.theme: Theme = THEMES[deck.theme]
        self.brand_problems: list[str] = []
        if deck.brand is not None:
            kit = {
                key: value.lstrip("#").upper() if key not in ("title_font", "body_font") else value
                for key, value in deck.brand.model_dump(exclude_none=True).items()
            }
            self.theme, self.brand_problems = apply_brand(self.theme, Brand(**kit))
        self.prs = self.m["pptx"].Presentation()
        self.prs.slide_width, self.prs.slide_height = self.m["inches"](WIDTH), self.m["inches"](HEIGHT)
        self.reports: list[SlideReport] = []
        self.report: SlideReport = SlideReport(0, "", "")

    # --- small helpers ---

    def colour(self, hex_value: str) -> Any:
        return self.m["rgb"].from_string(hex_value)

    def rectangle(
        self, slide: Any, x: float, y: float, w: float, h: float, fill: str, rounded: bool = False
    ) -> Any:
        kind = self.m["shape"].ROUNDED_RECTANGLE if rounded else self.m["shape"].RECTANGLE
        shape = slide.shapes.add_shape(kind, *(self.m["inches"](v) for v in (x, y, w, h)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.colour(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = 0.06
        return shape

    def style(
        self,
        run: Any,
        size: float,
        colour: str,
        bold: bool = False,
        italic: bool = False,
        heading: bool = False,
        font: str | None = None,
    ) -> None:
        run.font.name = font or (self.theme.title_font if heading else self.theme.body_font)
        run.font.size = self.m["pt"](size)
        run.font.bold, run.font.italic = bold, italic
        run.font.color.rgb = self.colour(colour)

    def textbox(self, slide: Any, x: float, y: float, w: float, h: float, anchor: str = "top") -> Any:
        box = slide.shapes.add_textbox(*(self.m["inches"](v) for v in (x, y, w, h)))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = self.m["anchor"].MIDDLE if anchor == "middle" else self.m["anchor"].TOP
        return frame

    def paragraph(self, frame: Any, first: bool, text: str, size: float, colour: str, **options: Any) -> Any:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        runs = options.get("runs")  # rich text: a list of richtext.RunSpec instead of one plain run
        pieces = (
            [(text, options.get("bold", False), options.get("italic", False), colour, None)]
            if runs is None
            else [
                (
                    r.text,
                    options.get("bold", False) or r.bold,
                    options.get("italic", False) or r.italic,
                    options["colour_for"](r.colour) if r.colour else colour,
                    r.link,
                )
                for r in runs
            ]
        )
        for piece_text, bold, italic, piece_colour, link in pieces:
            run = paragraph.add_run()
            run.text = piece_text
            self.style(
                run, size, piece_colour, bold, italic, options.get("heading", False), options.get("font")
            )
            if options.get("letter_spacing") is not None:
                run._r.get_or_add_rPr().set("spc", str(round(options["letter_spacing"] * 100)))
            if link:
                run.hyperlink.address = link
        if options.get("line_spacing"):
            paragraph.line_spacing = options["line_spacing"]
        if options.get("align"):
            paragraph.alignment = {
                "left": self.m["align"].LEFT,
                "center": self.m["align"].CENTER,
                "right": self.m["align"].RIGHT,
            }[options["align"]]
        paragraph.space_after = self.m["pt"](options.get("after", 0))
        if options.get("bullet"):
            self.bullet(paragraph, options.get("level", 0))
        return paragraph

    def bullet(self, paragraph: Any, level: int) -> None:
        indent = 0.3
        paragraph.level = level  # also recorded as a level, so reading the deck back shows the indentation
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(int(self.m["inches"](indent * (level + 1)))))
        properties.set("indent", str(-int(self.m["inches"](indent))))
        marker = properties.makeelement(
            self.m["qn"]("a:buChar"), {"char": BULLET if level == 0 else SUB_BULLET}
        )
        properties.append(marker)

    def title(self, slide: Any, text: str, size: float = 34, colour: str | None = None, **frame: Any) -> Any:
        placeholder = slide.shapes.title
        x, y, w, h = (
            frame.get("x", MARGIN),
            frame.get("y", TITLE_TOP),
            frame.get("w", CONTENT_WIDTH),
            frame.get("h", TITLE_HEIGHT),
        )
        placeholder.left, placeholder.top, placeholder.width, placeholder.height = (
            self.m["inches"](v) for v in (x, y, w, h)
        )
        frame_ = placeholder.text_frame
        frame_.word_wrap = True
        frame_.margin_left = frame_.margin_right = frame_.margin_top = frame_.margin_bottom = 0
        anchor = frame.get("anchor") or ("middle" if frame.get("middle") else "bottom")
        frame_.vertical_anchor = {
            "top": self.m["anchor"].TOP,
            "middle": self.m["anchor"].MIDDLE,
            "bottom": self.m["anchor"].BOTTOM,
        }[anchor]
        chosen, fits = fit.fit_size([(text, 1.0, True)], w, h, size, 22)
        if not fits:
            self.warn(f"the title is too long for its space: shorten it ({len(text.split())} words)")
        self.paragraph(
            frame_,
            True,
            text,
            chosen,
            colour or self.theme.text,
            bold=True,
            heading=True,
            align=frame.get("align", "left"),  # the layout default for a title is centred
        )
        return placeholder

    def warn(self, message: str) -> None:
        self.report.warnings.append(message)

    def new_slide(self, layout: str, title: str, notes: str, background: str | None = None) -> Any:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])  # "Title Only"
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self.colour(background or self.theme.background)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self.report = SlideReport(len(self.prs.slides), layout, title)
        self.reports.append(self.report)
        return slide

    def content_slide(self, layout: str, title: str, notes: str) -> Any:
        slide = self.new_slide(layout, title, notes)
        self.title(slide, title)
        self.add_footer(slide)
        return slide

    def add_footer(self, slide: Any) -> None:
        footer = self.deck.footer or self.deck.title
        small = self.textbox(slide, MARGIN, 7.0, 8, 0.3)
        self.paragraph(small, True, footer, 11, self.theme.muted)
        number = self.textbox(slide, WIDTH - MARGIN - 1, 7.0, 1, 0.3)
        self.paragraph(number, True, str(len(self.prs.slides)), 11, self.theme.muted, align="right")

    # --- bullets shared by the bullets and two-column layouts ---

    def bullet_items(self, bullets: list[str | BulletGroup]) -> list[tuple[str, int]]:
        items: list[tuple[str, int]] = []
        for entry in bullets:
            if isinstance(entry, str):
                items.append((entry, 0))
            else:
                items.append((entry.text, 0))
                items.extend((sub, 1) for sub in entry.sub)
        return items

    def write_bullets(
        self, slide: Any, items: list[tuple[str, int]], x: float, y: float, w: float, h: float, start: float
    ) -> None:
        texts = [(text, 1.0 if level == 0 else 0.85, False) for text, level in items]
        short = len(items) <= 4  # a short list gets bigger text and more air, so the slide is not top-heavy
        start, gap = (start + 4, 0.9) if short else (start, 0.5)
        size, fits = fit.fit_size(texts, w - 0.3, h, start, MIN_BODY_PT, gap_pt=start * gap)
        if not fits:
            self.warn(
                f"too much text for one slide at {MIN_BODY_PT} pt: split it into two slides or cut words"
            )
        top = [text for text, level in items if level == 0]
        if len(top) > MAX_BULLETS:
            self.warn(f"{len(top)} bullets: keep to {MAX_BULLETS} or fewer")
        for text in top:
            if len(text.split()) > MAX_BULLET_WORDS:
                self.warn(
                    f"a bullet has {len(text.split())} words: bullets are short phrases ({text[:40]}...)"
                )
                break
        frame = self.textbox(slide, x, y, w, h)
        for index, (text, level) in enumerate(items):
            self.paragraph(
                frame,
                index == 0,
                text,
                size if level == 0 else size * 0.85,
                self.theme.text if level == 0 else self.theme.muted,
                bullet=True,
                level=level,
                after=size * gap,
            )

    # --- the layouts ---

    def title_slide(self, slide_spec: TitleSlide | ClosingSlide, closing: bool) -> None:
        slide = self.new_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        self.title(slide, slide_spec.title, 46, x=MARGIN, y=2.45, w=CONTENT_WIDTH * 0.8, h=1.7, middle=False)
        subtitle = slide_spec.subtitle or (self.deck.subtitle if not closing else "")
        if subtitle:
            frame = self.textbox(slide, MARGIN, 4.35, CONTENT_WIDTH * 0.8, 1.0)
            self.paragraph(frame, True, subtitle, 22, self.theme.muted)
        if self.deck.author and not closing:
            frame = self.textbox(slide, MARGIN, 6.3, 8, 0.4)
            self.paragraph(frame, True, self.deck.author, 16, self.theme.muted)

    def section_slide(self, slide_spec: SectionSlide) -> None:
        slide = self.new_slide(slide_spec.layout, slide_spec.title, slide_spec.notes, self.theme.accent)
        self.title(
            slide, slide_spec.title, 44, self.theme.accent_text, x=MARGIN, y=2.6, w=CONTENT_WIDTH * 0.8, h=1.6
        )
        if slide_spec.subtitle:
            frame = self.textbox(slide, MARGIN, 4.4, CONTENT_WIDTH * 0.8, 1.0)
            self.paragraph(frame, True, slide_spec.subtitle, 22, self.theme.accent_text)

    def bullets_slide(self, slide_spec: BulletsSlide) -> None:
        slide = self.content_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        self.write_bullets(
            slide,
            self.bullet_items(slide_spec.bullets),
            MARGIN,
            BODY_TOP,
            CONTENT_WIDTH,
            BODY_BOTTOM - BODY_TOP,
            28,
        )

    def two_column_slide(self, slide_spec: TwoColumnSlide) -> None:
        slide = self.content_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        gap = 0.6
        width = (CONTENT_WIDTH - gap) / 2
        for index, column in enumerate((slide_spec.left, slide_spec.right)):
            self.column(slide, column, MARGIN + index * (width + gap), width)

    def column(self, slide: Any, column: Column, x: float, width: float) -> None:
        top = BODY_TOP
        if column.heading:
            frame = self.textbox(slide, x, top, width, 0.5)
            self.paragraph(frame, True, column.heading, 22, self.theme.accent, bold=True, heading=True)
            top += 0.65
        self.write_bullets(slide, self.bullet_items(column.bullets), x, top, width, BODY_BOTTOM - top, 24)

    def add_picture(self, slide: Any, data: bytes, x: float, y: float, w: float, h: float, alt: str) -> Any:
        shape = slide.shapes.add_picture(io.BytesIO(data), *(self.m["inches"](v) for v in (x, y, w, h)))
        shape._element.nvPicPr.cNvPr.set("descr", alt)
        return shape

    def image_slide(self, slide_spec: ImageSlide) -> None:
        from PIL import Image

        slide = self.content_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        data = self.read_image(slide_spec.image)
        with Image.open(io.BytesIO(data)) as picture:
            ratio = picture.width / picture.height
        caption_height = 0.5 if slide_spec.caption else 0.0
        area_w, area_h = CONTENT_WIDTH, BODY_BOTTOM - BODY_TOP - caption_height
        w, h = (area_w, area_w / ratio) if area_w / ratio <= area_h else (area_h * ratio, area_h)
        x, y = MARGIN + (area_w - w) / 2, BODY_TOP + (area_h - h) / 2
        self.add_picture(slide, data, x, y, w, h, slide_spec.alt or slide_spec.caption or slide_spec.title)
        if not (slide_spec.alt or slide_spec.caption):
            self.warn("the picture has no alt text or caption: say what it shows")
        if slide_spec.caption:
            frame = self.textbox(slide, MARGIN, BODY_BOTTOM - 0.35, CONTENT_WIDTH, 0.4)
            self.paragraph(frame, True, slide_spec.caption, 14, self.theme.muted, italic=True, align="center")

    def add_table(
        self,
        slide: Any,
        columns: list[str],
        rows: list[list[str | int | float]],
        box: tuple[float, float, float, float],
        descr: str,
        size: float | None = None,
    ) -> None:
        """A native table in `box` (x, y, width, maximum height). Row text is sized by the number of rows."""
        x, y, width, area_h = box
        if any(len(row) != len(columns) for row in rows):
            raise BuildError(
                f"slide {self.report.number}: every table row needs {len(columns)} cells (one per column)"
            )
        if len(rows) > MAX_TABLE_ROWS:
            self.warn(f"{len(rows)} table rows: more than {MAX_TABLE_ROWS} is hard to read, split the table")
        size = size or (24 if len(rows) <= 4 else 20 if len(rows) <= 6 else 16 if len(rows) <= 8 else 14)
        every_row: list[list[str | int | float]] = [list(columns), *rows]
        grid = [[str(cell) for cell in cells] for cells in every_row]
        weights = [max(len(row[i]) for row in grid) + 4 for i in range(len(columns))]
        widths = [width * wt / sum(weights) for wt in weights]
        needed = 0.0
        for row_number, cells in enumerate(grid):
            lines = max(
                fit.lines_needed(cell, widths[i] - 0.3, size, row_number == 0) for i, cell in enumerate(cells)
            )
            needed += lines * size * 1.2 / 72 + 0.22
        if needed > area_h:
            self.warn("the table is too tall for the slide: remove rows or shorten the cell text")
        frame = slide.shapes.add_table(
            len(rows) + 1,
            len(columns),
            *(self.m["inches"](v) for v in (x, y, width, min(needed, area_h))),
        )
        frame._element.nvGraphicFramePr.cNvPr.set("descr", descr)
        table = frame.table
        for index, column_width in enumerate(widths):
            table.columns[index].width = self.m["inches"](column_width)
        for row_index, cells in enumerate(grid):
            for column_index, value in enumerate(cells):
                cell = table.cell(row_index, column_index)
                header = row_index == 0
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.colour(
                    self.theme.accent
                    if header
                    else self.theme.surface
                    if row_index % 2 == 0
                    else self.theme.background
                )
                cell.margin_left = cell.margin_right = self.m["inches"](0.12)
                cell.vertical_anchor = self.m["anchor"].MIDDLE
                paragraph = cell.text_frame.paragraphs[0]
                run = paragraph.add_run()
                run.text = value
                self.style(run, size, self.theme.accent_text if header else self.theme.text, bold=header)

    def table_slide(self, slide_spec: TableSlide) -> None:
        slide = self.content_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        caption_height = 0.5 if slide_spec.caption else 0.0
        box = (MARGIN, BODY_TOP, CONTENT_WIDTH, BODY_BOTTOM - BODY_TOP - caption_height)
        self.add_table(
            slide, slide_spec.columns, slide_spec.rows, box, slide_spec.caption or slide_spec.title
        )
        if slide_spec.caption:
            note = self.textbox(slide, MARGIN, BODY_BOTTOM - 0.35, CONTENT_WIDTH, 0.4)
            self.paragraph(note, True, slide_spec.caption, 14, self.theme.muted, italic=True)

    def add_chart(
        self,
        slide: Any,
        options: Any,
        box: tuple[float, float, float, float],
        descr: str,
        font_size: float = 16,
    ) -> None:
        """A native chart in `box` (x, y, width, height): see charts.py for everything it can do."""
        charts.add_chart(self, slide, options, box, descr, font_size)

    def chart_slide(self, slide_spec: ChartSlide) -> None:
        slide = self.content_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        caption_height = 0.5 if slide_spec.caption else 0.0
        box = (MARGIN, BODY_TOP, CONTENT_WIDTH, BODY_BOTTOM - BODY_TOP - caption_height)
        self.add_chart(slide, slide_spec, box, slide_spec.caption or slide_spec.title)
        if slide_spec.caption:
            note = self.textbox(slide, MARGIN, BODY_BOTTOM - 0.35, CONTENT_WIDTH, 0.4)
            self.paragraph(note, True, slide_spec.caption, 14, self.theme.muted, italic=True)

    def quote_slide(self, slide_spec: QuoteSlide) -> None:
        slide = self.new_slide(slide_spec.layout, slide_spec.quote[:60], slide_spec.notes)
        mark = self.textbox(slide, MARGIN, 1.0, 2, 1.6)
        self.paragraph(mark, True, "“", 120, self.theme.accent, heading=True)
        self.title(
            slide, slide_spec.quote, 36, x=MARGIN + 0.1, y=2.3, w=CONTENT_WIDTH * 0.85, h=3.0, middle=True
        )
        if slide_spec.by:
            frame = self.textbox(slide, MARGIN + 0.1, 5.6, CONTENT_WIDTH * 0.85, 0.5)
            self.paragraph(frame, True, f"— {slide_spec.by}", 20, self.theme.muted)

    def stats_slide(self, slide_spec: StatsSlide) -> None:
        slide = self.content_slide(slide_spec.layout, slide_spec.title, slide_spec.notes)
        items = slide_spec.items
        if not 1 <= len(items) <= 4:
            raise BuildError(
                f"slide {self.report.number}: a stats slide shows 1 to 4 numbers, not {len(items)}"
            )
        gap = 0.4
        width = (CONTENT_WIDTH - gap * (len(items) - 1)) / len(items)
        for index, stat in enumerate(items):
            x = MARGIN + index * (width + gap)
            self.rectangle(slide, x, 2.4, width, 3.4, self.theme.surface, rounded=True)
            size, fits = fit.fit_size([(stat.value, 1.0, True)], width - 0.4, 1.3, 60, 32)
            if not fits:
                self.warn(f"'{stat.value}' is too wide for its card: use a shorter figure")
            value = self.textbox(slide, x + 0.2, 2.8, width - 0.4, 1.4, "middle")
            self.paragraph(
                value, True, stat.value, size, self.theme.accent, bold=True, heading=True, align="center"
            )
            label = self.textbox(slide, x + 0.25, 4.35, width - 0.5, 1.2)
            self.paragraph(label, True, stat.label, 20, self.theme.text, align="center")


def build_deck(deck: Deck, read_image: ReadImage) -> BuiltDeck:
    """The finished file and one report per slide. `read_image` returns the bytes of a workspace picture."""
    builder = _Builder(deck, read_image)
    for index, slide_spec in enumerate(deck.slides, start=1):
        if isinstance(slide_spec, TitleSlide):
            builder.title_slide(slide_spec, closing=False)
        elif isinstance(slide_spec, ClosingSlide):
            builder.title_slide(slide_spec, closing=True)
        elif isinstance(slide_spec, SectionSlide):
            builder.section_slide(slide_spec)
        elif isinstance(slide_spec, BulletsSlide):
            builder.bullets_slide(slide_spec)
        elif isinstance(slide_spec, TwoColumnSlide):
            builder.two_column_slide(slide_spec)
        elif isinstance(slide_spec, ImageSlide):
            builder.image_slide(slide_spec)
        elif isinstance(slide_spec, TableSlide):
            builder.table_slide(slide_spec)
        elif isinstance(slide_spec, ChartSlide):
            builder.chart_slide(slide_spec)
        elif isinstance(slide_spec, QuoteSlide):
            builder.quote_slide(slide_spec)
        elif isinstance(slide_spec, FreeformSlide):
            build_freeform(builder, slide_spec)
        else:
            builder.stats_slide(slide_spec)
        del index
    properties = builder.prs.core_properties
    properties.title, properties.author = deck.title, deck.author or ""
    output = io.BytesIO()
    builder.prs.save(output)
    for number, message in lint.slide_title_warnings(deck).items():
        builder.reports[number - 1].warnings.append(message)
    notes = [f"brand colours: {problem}" for problem in builder.brand_problems]
    if deck.brand is not None:
        notes += [
            w
            for name in (deck.brand.title_font, deck.brand.body_font)
            if name and (w := lint.font_warning(name))
        ]
    return BuiltDeck(output.getvalue(), builder.reports, [*notes, *lint.deck_warnings(deck)])
