"""Bringing slides from one deck into another, merging decks and splitting one (TODO P2.14, D-320).

`import_slides` copies slides of a source deck to the end (or a position) of the target:
- design "keep" (default): the source's slide master, layouts and theme are cloned into the target (once; a master identical to one the
  target already has is reused), so the slides look as they did. Only when both decks have the same slide size.
- design "match": each slide goes on the target's own layout of the same name (else the one with the most placeholders in common), so
  the slides take the target's look. Shapes are scaled when the slide sizes differ.
`merge_decks` builds on the first deck and imports all slides of the others (a section per source deck); `split_deck` keeps chosen slides."""

from __future__ import annotations

import hashlib
from typing import Any

from forge.slides.clone import clone_part, copy_slide
from forge.slides.features import copy_features, keep_sections_valid
from forge.slides.open_deck import DeckError

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SIZE_TOLERANCE = 0.01
MASTER_RELATION = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
FIRST_MASTER_ID = 2147483648  # master and layout ids start here and must differ across the whole deck


def _same_size(source: Any, target: Any) -> bool:
    return bool(
        abs(source.slide_width - target.slide_width) <= target.slide_width * SIZE_TOLERANCE
        and abs(source.slide_height - target.slide_height) <= target.slide_height * SIZE_TOLERANCE
    )


def _signature(master: Any) -> str:
    """What makes two masters the same: their XML, their theme and the names of their layouts."""
    from pptx.opc.constants import RELATIONSHIP_TYPE

    theme = master.part.part_related_by(RELATIONSHIP_TYPE.THEME).blob
    names = "|".join(layout.name for layout in master.slide_layouts)
    return hashlib.sha256(master.part.blob + theme + names.encode()).hexdigest()


def _layout_by_name(master: Any, name: str) -> Any:
    for layout in master.slide_layouts:
        if layout.name == name:
            return layout
    raise DeckError(f"the layout {name!r} is missing in the imported master")


def _renumber(target: Any, master: Any) -> None:
    """Gives the imported master and its layouts fresh ids (they are numbers that must differ across the deck)."""
    used: list[int] = []
    for other in target.slide_masters:
        if other.part is master.part:
            continue
        used += [int(i) for i in other._element.xpath("./p:sldLayoutIdLst/p:sldLayoutId/@id")]
    for entry in target.part._element.xpath("./p:sldMasterIdLst/p:sldMasterId"):
        if target.part.related_part(entry.get(f"{{{NS_R}}}id")) is not master.part:
            used.append(int(entry.get("id")))
    fresh = max([*used, FIRST_MASTER_ID - 1]) + 1
    for entry in target.part._element.xpath("./p:sldMasterIdLst/p:sldMasterId"):
        if target.part.related_part(entry.get(f"{{{NS_R}}}id")) is master.part:
            entry.set("id", str(fresh))
    for layout_entry in master._element.xpath("./p:sldLayoutIdLst/p:sldLayoutId"):
        fresh += 1
        layout_entry.set("id", str(fresh))


def _import_master(target: Any, source_master: Any, memo: dict[str, Any]) -> Any:
    key = "master:" + str(source_master.part.partname)
    if key in memo:
        return memo[key]
    signature = _signature(source_master)
    for existing in target.slide_masters:
        if _signature(existing) == signature:
            memo[key] = existing
            return existing
    from pptx.oxml import parse_xml

    cloned = clone_part(source_master.part, target.part.package, memo)
    relation_id = target.part.relate_to(cloned, MASTER_RELATION)
    listing = target.part._element.find(f"{{{NS_P}}}sldMasterIdLst")
    listing.append(
        parse_xml(
            f'<p:sldMasterId xmlns:p="{NS_P}" xmlns:r="{NS_R}" id="{FIRST_MASTER_ID}" r:id="{relation_id}"/>'
        )
    )
    master = cloned.slide_master
    _renumber(target, master)
    memo[key] = master
    return master


def _match_layout(target: Any, source_layout: Any) -> Any:
    for layout in target.slide_layouts:
        if layout.name == source_layout.name:
            return layout
    wanted = {p.placeholder_format.type for p in source_layout.placeholders}
    best, best_score = None, -1
    for layout in target.slide_layouts:
        have = {p.placeholder_format.type for p in layout.placeholders}
        score = len(wanted & have) * 2 - len(have ^ wanted)
        if score > best_score:
            best, best_score = layout, score
    if best is None:
        raise DeckError("the deck has no layouts to put the slides on")
    return best


def import_slides(
    target: Any,
    source: Any,
    numbers: list[int] | None = None,
    at: int | None = None,
    design: str = "keep",
    memo: dict[str, Any] | None = None,
) -> list[str]:
    """Copies slides `numbers` (default all) of `source` into `target` at position `at` (default the end). Returns notes for the user."""
    memo = {} if memo is None else memo
    total = len(source.slides)
    chosen = numbers or list(range(1, total + 1))
    bad = [n for n in chosen if not 1 <= n <= total]
    if bad:
        raise DeckError(f"the source deck has {total} slides: there is no slide {bad[0]}")
    if design not in ("keep", "match"):
        raise DeckError("design must be 'keep' or 'match'")
    notes: list[str] = []
    same = _same_size(source, target)
    if design == "keep" and not same:
        design = "match"
        notes.append(
            "the slide sizes differ, so the slides were put on this deck's layouts (design 'match') and scaled"
        )
    scale = (
        (1.0, 1.0)
        if same
        else (target.slide_width / source.slide_width, target.slide_height / source.slide_height)
    )
    first_new = len(target.slides) + 1
    for number in chosen:
        original = source.slides[number - 1]
        if design == "keep":
            master = _import_master(target, original.slide_layout.slide_master, memo)
            layout = _layout_by_name(master, original.slide_layout.name)
        else:
            layout = _match_layout(target, original.slide_layout)
        try:
            copy = copy_slide(original, target, layout, memo, scale)
        except DeckError as error:
            raise DeckError(f"slide {number} of the source: {error}") from error
        copy_features(original, copy)
    entries = target.slides._sldIdLst
    if at is not None and 1 <= at < first_new:
        moved = list(entries)[first_new - 1 :]
        for offset, entry in enumerate(moved):
            entries.remove(entry)
            entries.insert(at - 1 + offset, entry)
    keep_sections_valid(target)
    return notes


def merge_decks(decks: list[tuple[str, Any]], design: str = "keep") -> tuple[Any, list[str]]:
    """The first deck with every slide of the others added, a section per deck (named by its label). Returns (deck, notes)."""
    from forge.slides.features import Section, read_sections, set_sections

    if len(decks) < 2:
        raise DeckError("merging needs at least two decks")
    label, base = decks[0]
    starts = [Section(label, 1)] if not read_sections(base) else []
    notes: list[str] = []
    for label, source in decks[1:]:
        begin = len(base.slides) + 1
        notes += import_slides(base, source, design=design)
        if starts:
            starts.append(Section(label, begin))
    if starts:
        set_sections(base, starts)
    return base, notes


def split_deck(deck: Any, keep: list[int]) -> Any:
    """Removes every slide not in `keep` (the deck is changed; open a fresh copy for each part)."""
    total = len(deck.slides)
    bad = [n for n in keep if not 1 <= n <= total]
    if bad or not keep:
        raise DeckError(
            f"choose slides between 1 and {total}" + (f" (there is no slide {bad[0]})" if bad else "")
        )
    entries = deck.slides._sldIdLst
    for index, entry in reversed(list(enumerate(list(entries), start=1))):
        if index not in keep:
            deck.part.drop_rel(entry.rId)
            entries.remove(entry)
    keep_sections_valid(deck)
    return deck
