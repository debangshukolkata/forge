"""A real slide master, layouts and theme for a generated deck (TODO P2.7, D-315).

python-pptx starts from a 4:3 Office template. Without this step a deck looks right only because every slide is drawn by hand: change
the theme in PowerPoint (Design > Variants) and nothing follows, and a slide added by hand starts from a 4:3 layout with Office fonts.
`install_master` rewrites the template so it matches the deck:

- the theme's colour scheme (text, background, muted, surface, the chart palette, links) and font scheme (title and body font);
- the master: background, title and body text styles, and the title, body, date, footer and slide-number placeholders at the same
  places the builder draws them;
- four layouts (Title Slide, Title and Content, Title Only, Blank) in 16:9; the other seven 4:3 layouts are removed.

`add_footer_placeholders` puts a real footer and a real slide-number field on a slide (they follow the master, and the number counts
by itself when slides move) instead of two text boxes."""

from __future__ import annotations

import re
import uuid
from typing import Any

from lxml import etree  # type: ignore[import-untyped]

from forge.slides.themes import Theme

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
A = "{" + NS["a"] + "}"
P = "{" + NS["p"] + "}"
EMU = 914400
KEPT_LAYOUTS = ("Title Slide", "Title and Content", "Title Only", "Blank")
LAYOUT_TITLE_ONLY = "Title Only"
TITLE_POINTS, BODY_POINTS = 34, 20
FOOTER_POINTS = 11
FOOTER_TOP, FOOTER_HEIGHT = 7.0, 0.3
NUMBER_WIDTH = 1.0
DATE_WIDTH = 2.3


def _emu(inches: float) -> str:
    return str(round(inches * EMU))


def _set_frame(shape: Any, x: float, y: float, w: float, h: float) -> None:
    """Sets (or adds) the explicit position of a placeholder element."""
    properties = shape.find(f"{P}spPr")
    transform = properties.find(f"{A}xfrm")
    if transform is None:
        transform = etree.SubElement(properties, f"{A}xfrm")
        properties.insert(0, transform)
        etree.SubElement(transform, f"{A}off")
        etree.SubElement(transform, f"{A}ext")
    transform.find(f"{A}off").set("x", _emu(x))
    transform.find(f"{A}off").set("y", _emu(y))
    transform.find(f"{A}ext").set("cx", _emu(w))
    transform.find(f"{A}ext").set("cy", _emu(h))


def _placeholder(tree: Any, kind: str) -> Any:
    for shape in tree.iterfind(f".//{P}sp"):
        holder = shape.find(f".//{P}ph")
        if holder is not None and holder.get("type") == kind:
            return shape
    return None


def _colour_scheme(theme: Theme) -> str:
    accents = list(theme.palette) or [theme.accent]
    while len(accents) < 6:  # a theme needs six accents; the missing ones are the palette again
        accents.append(accents[len(accents) % len(theme.palette)] if theme.palette else theme.accent)
    slots = {
        "dk1": theme.text,
        "lt1": theme.background,
        "dk2": theme.muted,
        "lt2": theme.surface,
        **{f"accent{i + 1}": accents[i] for i in range(6)},
        "hlink": theme.accent,
        "folHlink": theme.muted,
    }
    inner = "".join(
        f'<a:{name}><a:srgbClr val="{value.lstrip("#").upper()}"/></a:{name}>'
        for name, value in slots.items()
    )
    return f'<a:clrScheme xmlns:a="{NS["a"]}" name="{_xml_name(theme.name)}">{inner}</a:clrScheme>'


def _xml_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9 _-]", "", text) or "DebCraft"


def _theme_part(master: Any) -> Any:
    from pptx.opc.constants import RELATIONSHIP_TYPE

    return master.part.part_related_by(RELATIONSHIP_TYPE.THEME)


def _write_theme(master: Any, theme: Theme) -> None:
    part = _theme_part(master)
    root = etree.fromstring(part.blob)
    root.set("name", f"DebCraft {_xml_name(theme.name)}")
    elements = root.find(f"{A}themeElements")
    old = elements.find(f"{A}clrScheme")
    elements.replace(old, etree.fromstring(_colour_scheme(theme)))
    fonts = elements.find(f"{A}fontScheme")
    fonts.set("name", _xml_name(theme.title_font))
    for tag, face in (("majorFont", theme.title_font), ("minorFont", theme.body_font)):
        fonts.find(f"{A}{tag}/{A}latin").set("typeface", face)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _style_master(
    master: Any, theme: Theme, width: float, margin: float, title_top: float, title_height: float
) -> None:
    root = master._element
    content = width - 2 * margin
    background = root.find(f"{P}cSld/{P}bg")
    if background is not None:
        background.getparent().remove(background)
    colours = etree.fromstring(
        f'<p:bg xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}"><p:bgPr><a:solidFill><a:schemeClr val="bg1"/></a:solidFill>'
        "<a:effectLst/></p:bgPr></p:bg>"
    )
    root.find(f"{P}cSld").insert(0, colours)
    # Text styles: titles bold and left aligned, body at a size readable from the back of a room.
    title = root.find(f"{P}txStyles/{P}titleStyle/{A}lvl1pPr")
    title.set("algn", "l")
    title.find(f"{A}defRPr").set("sz", str(TITLE_POINTS * 100))
    title.find(f"{A}defRPr").set("b", "1")
    body = root.find(f"{P}txStyles/{P}bodyStyle")
    for level, points in enumerate(
        (BODY_POINTS, BODY_POINTS - 2, BODY_POINTS - 4, BODY_POINTS - 4, BODY_POINTS - 4), start=1
    ):
        style = body.find(f"{A}lvl{level}pPr")
        if style is not None:
            style.find(f"{A}defRPr").set("sz", str(points * 100))
    # Placeholders where the builder draws things.
    bottom = FOOTER_TOP
    positions = {
        "title": (margin, title_top, content, title_height),
        "body": (
            margin,
            title_top + title_height + 0.3,
            content,
            bottom - 0.25 - (title_top + title_height + 0.3),
        ),
        "dt": (width - margin - NUMBER_WIDTH - 0.2 - DATE_WIDTH, FOOTER_TOP, DATE_WIDTH, FOOTER_HEIGHT),
        "ftr": (margin, FOOTER_TOP, 8.0, FOOTER_HEIGHT),
        "sldNum": (width - margin - NUMBER_WIDTH, FOOTER_TOP, NUMBER_WIDTH, FOOTER_HEIGHT),
    }
    for kind, frame in positions.items():
        shape = _placeholder(root, kind)
        if shape is not None:
            _set_frame(shape, *frame)
    title_shape = _placeholder(root, "title")
    if title_shape is not None:
        title_shape.find(f"{P}txBody/{A}bodyPr").set("anchor", "b")
        for side in ("lIns", "tIns", "rIns", "bIns"):
            title_shape.find(f"{P}txBody/{A}bodyPr").set(side, "0")
    for kind, alignment in (("dt", "r"), ("ftr", "l"), ("sldNum", "r")):
        shape = _placeholder(root, kind)
        if shape is None:
            continue
        level = shape.find(f".//{A}lvl1pPr")
        level.set("algn", alignment)
        run = level.find(f"{A}defRPr")
        run.set("sz", str(FOOTER_POINTS * 100))
        for child in list(run):
            run.remove(child)
        fill = etree.SubElement(run, f"{A}solidFill")
        etree.SubElement(fill, f"{A}schemeClr").set("val", "tx2")
        body_properties = shape.find(f"{P}txBody/{A}bodyPr")
        body_properties.set("lIns", "0")
        body_properties.set("rIns", "0")


def _style_layouts(prs: Any, margin: float, width: float) -> None:
    for layout in list(prs.slide_layouts):
        if layout.name not in KEPT_LAYOUTS:
            prs.slide_layouts.remove(layout)
    content = width - 2 * margin
    title_slide = next(layout for layout in prs.slide_layouts if layout.name == "Title Slide")
    tree = title_slide._element
    centre = _placeholder(tree, "ctrTitle")
    subtitle = _placeholder(tree, "subTitle")
    _set_frame(centre, margin, 2.2, content, 1.7)
    centre.find(f"{P}txBody/{A}bodyPr").set("anchor", "b")
    _set_frame(subtitle, margin, 4.1, content, 1.2)
    for level in subtitle.iterfind(f".//{A}lvl1pPr"):
        level.set("algn", "l")
    for level in subtitle.iterfind(f".//{A}lvl2pPr"):
        level.set("algn", "l")
    for layout in prs.slide_layouts:  # layouts follow the master's places for these
        for kind in ("dt", "ftr", "sldNum"):
            shape = _placeholder(layout._element, kind)
            if shape is not None:
                transform = shape.find(f"{P}spPr/{A}xfrm")
                if transform is not None:
                    transform.getparent().remove(transform)


def install_master(
    prs: Any, theme: Theme, width: float, margin: float, title_top: float, title_height: float
) -> None:
    """Rewrites the template's theme, master and layouts for `theme` (see the module text). Call it on a fresh presentation."""
    master = prs.slide_master
    _write_theme(master, theme)
    _style_master(master, theme, width, margin, title_top, title_height)
    _style_layouts(prs, margin, width)


def layout_named(prs: Any, name: str) -> Any:
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    return prs.slide_layouts[0]


def _layout_index(layout: Any, kind: str) -> str | None:
    shape = _placeholder(layout._element, kind)
    holder = shape.find(f".//{P}ph") if shape is not None else None
    return holder.get("idx") if holder is not None else None


def _next_id(slide: Any) -> int:
    return max([int(i) for i in slide.shapes._spTree.xpath("//p:cNvPr/@id")] + [1]) + 1


def add_footer_placeholders(slide: Any, text: str, number: int, language: str = "en-US") -> None:
    """A real footer and a real slide-number field, linked to the layout's placeholders (idx taken from the layout)."""
    tree = slide.shapes._spTree
    for kind, name, paragraph in (
        (
            "ftr",
            "Footer Placeholder",
            f'<a:p><a:r><a:rPr lang="{language}"/><a:t></a:t></a:r></a:p>',
        ),
        (
            "sldNum",
            "Slide Number Placeholder",
            f'<a:p><a:fld id="{{{str(uuid.uuid4()).upper()}}}" type="slidenum"><a:rPr lang="{language}"/><a:t>{number}</a:t></a:fld></a:p>',
        ),
    ):
        index = _layout_index(slide.slide_layout, kind)
        if index is None:
            continue
        identifier = _next_id(slide)
        element = etree.fromstring(
            f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}"><p:nvSpPr><p:cNvPr id="{identifier}" name="{name} {identifier - 1}"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="{kind}" sz="quarter" idx="{index}"/></p:nvPr></p:nvSpPr>'
            f"<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>{paragraph}</p:txBody></p:sp>"
        )
        if kind == "ftr":
            element.find(f".//{A}t").text = text
        tree.append(element)
