"""Changes to an existing deck (D-241): replace text, set a shape's text, notes, delete / move / duplicate
slides, add a slide from the deck's own layouts (which is how a company template is used), update chart data
and table cells.

Formatting is kept: new text takes the look of the first run and paragraph it replaces. Every operation either
works or stops the whole edit with a message naming the operation, so a half-edited deck is never written."""

from __future__ import annotations

from collections.abc import Callable
from copy import deepcopy
from typing import Annotated, Any, Literal

from pydantic import BaseModel, ConfigDict, Field, TypeAdapter, ValidationError

from forge.slides.edit_find import find_shape, slide_at, walk
from forge.slides.edit_more import MORE_OPERATIONS, MoreOperation, apply_more
from forge.slides.edit_text import set_frame_text
from forge.slides.open_deck import DeckError
from forge.slides.spec import Series

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
UNSUPPORTED_COPY = ("chart", "oleObject", "package", "video", "audio", "media", "diagramData")


class _Op(BaseModel):
    model_config = ConfigDict(extra="forbid")


class ReplaceText(_Op):
    op: Literal["replace_text"]
    find: str
    replace: str
    slide: int | None = None  # only this slide; default every slide
    notes: bool = False  # also change speaker notes


class SetText(_Op):
    op: Literal["set_text"]
    slide: int
    shape: int | str  # the id or the exact name shown by read_presentation
    text: str  # one paragraph per line; two leading spaces per bullet level


class SetNotes(_Op):
    op: Literal["set_notes"]
    slide: int
    text: str


class DeleteSlide(_Op):
    op: Literal["delete_slide"]
    slide: int


class MoveSlide(_Op):
    op: Literal["move_slide"]
    slide: int
    to: int


class DuplicateSlide(_Op):
    op: Literal["duplicate_slide"]
    slide: int
    at: int | None = None  # position of the copy; default right after the original


class AddSlide(_Op):
    op: Literal["add_slide"]
    layout: str  # a layout name listed by read_presentation
    title: str = ""
    body: list[str] = Field(
        default_factory=list
    )  # the main text placeholder; two leading spaces per bullet level
    placeholders: dict[str, str] = Field(
        default_factory=dict
    )  # other placeholders by number, e.g. {"2": "..."}
    notes: str = ""
    at: int | None = None  # position; default the end


class UpdateChart(_Op):
    op: Literal["update_chart"]
    slide: int
    shape: int | str
    categories: list[str]
    series: list[Series]


class UpdateTable(_Op):
    op: Literal["update_table"]
    slide: int
    shape: int | str
    rows: list[list[str]]  # every row including the header, same size as the table now


Operation = Annotated[
    ReplaceText
    | SetText
    | SetNotes
    | DeleteSlide
    | MoveSlide
    | DuplicateSlide
    | AddSlide
    | UpdateChart
    | UpdateTable
    | MoreOperation,
    Field(discriminator="op"),
]
OPERATIONS = (
    "replace_text, set_text, set_notes, delete_slide, move_slide, duplicate_slide, add_slide, "
    "update_chart, update_table, " + MORE_OPERATIONS
)


def parse_operations(raw: list[dict[str, Any]]) -> list[Operation]:
    try:
        return TypeAdapter(list[Operation]).validate_python(raw)
    except ValidationError as error:
        problems = []
        for item in error.errors()[:6]:
            where = ".".join(str(p) for p in item["loc"] if isinstance(p, int | str))
            problems.append(f"- operation {where}: {item['msg']}")
        raise DeckError(
            "The operations have problems:\n" + "\n".join(problems) + f"\nOperations are: {OPERATIONS}."
        ) from error


# --- finding things ---


# --- text that keeps its look ---


def _replace_in_paragraph(paragraph: Any, find: str, replace: str) -> int:
    runs = paragraph.runs
    count = "".join(r.text for r in runs).count(find)
    if not count:
        return 0
    for run in runs:
        if find in run.text:
            run.text = run.text.replace(find, replace)
    joined = "".join(r.text for r in runs)
    if find in joined:  # the text was split across runs (a word with a different font in the middle)
        runs[0].text = joined.replace(find, replace)
        for run in runs[1:]:
            run.text = ""
    return count


def _frames(slide: Any, include_notes: bool) -> Any:
    for shape in walk(slide.shapes):
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
    if include_notes and slide.has_notes_slide:
        yield slide.notes_slide.notes_text_frame


# --- slide structure ---


def _slide_list(deck: Any) -> Any:
    return deck.slides._sldIdLst


def _delete(deck: Any, number: int) -> None:
    slide_at(deck, number)
    entries = _slide_list(deck)
    entry = list(entries)[number - 1]
    deck.part.drop_rel(entry.rId)
    entries.remove(entry)


def _move(deck: Any, number: int, to: int) -> None:
    slide_at(deck, number)
    slide_at(deck, to)
    entries = _slide_list(deck)
    entry = list(entries)[number - 1]
    entries.remove(entry)
    entries.insert(to - 1, entry)


def _duplicate(deck: Any, number: int, at: int | None) -> None:
    source = slide_at(deck, number)
    copy = deck.slides.add_slide(source.slide_layout)
    for shape in list(copy.shapes):
        shape._element.getparent().remove(shape._element)
    tree = copy.shapes._spTree
    for element in source.shapes._spTree:
        if not element.tag.endswith(("}nvGrpSpPr", "}grpSpPr", "}extLst")):
            tree.append(deepcopy(element))
    for node in tree.iter():
        for attribute in ("embed", "link", "id", "pict"):
            old = node.get(f"{{{NS_R}}}{attribute}")
            if old is None:
                continue
            relation = source.part.rels[old]
            kind = relation.reltype.rsplit("/", 1)[-1]
            if kind in UNSUPPORTED_COPY:
                raise DeckError(
                    f"slide {number} holds a {kind}, which cannot be copied safely: "
                    "add a new slide and build it again"
                )
            if relation.is_external:
                node.set(
                    f"{{{NS_R}}}{attribute}",
                    copy.part.relate_to(relation.target_ref, relation.reltype, is_external=True),
                )
            else:
                node.set(
                    f"{{{NS_R}}}{attribute}", copy.part.relate_to(relation.target_part, relation.reltype)
                )
    background = source._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if background is not None:
        copy._element.cSld.insert(0, deepcopy(background))
    if source.has_notes_slide:
        copy.notes_slide.notes_text_frame.text = source.notes_slide.notes_text_frame.text
    _move(deck, len(deck.slides), at if at is not None else number + 1)


def _layout(deck: Any, name: str) -> Any:
    layouts = list(deck.slide_layouts)
    exact = [layout for layout in layouts if layout.name.lower() == name.lower()]
    near = [layout for layout in layouts if name.lower() in layout.name.lower()]
    chosen = exact or near
    if not chosen:
        raise DeckError(
            f"the deck has no layout called {name!r}. Layouts: "
            + ", ".join(f'"{layout.name}"' for layout in layouts)
        )
    return chosen[0]


def _add(deck: Any, operation: AddSlide) -> None:
    slide = deck.slides.add_slide(_layout(deck, operation.layout))
    title = slide.shapes.title
    others = [
        p
        for p in slide.placeholders
        if title is None or p.placeholder_format.idx != title.placeholder_format.idx
        if p.placeholder_format.type is None
        or p.placeholder_format.type.name not in {"DATE", "FOOTER", "SLIDE_NUMBER"}
    ]
    if operation.title:
        if title is None:
            raise DeckError(f'the layout "{operation.layout}" has no title placeholder')
        set_frame_text(title.text_frame, [operation.title])
    if operation.body:
        if not others:
            raise DeckError(
                f'the layout "{operation.layout}" has no text placeholder for the body: pick another layout'
            )
        set_frame_text(others[0].text_frame, operation.body)
    by_number = {p.placeholder_format.idx: p for p in others}
    for key, text in operation.placeholders.items():
        holder = by_number.get(int(key)) if key.isdigit() else None
        if holder is None:
            raise DeckError(
                f"the layout has no placeholder {key}: it has " + ", ".join(str(i) for i in by_number)
            )
        set_frame_text(holder.text_frame, [text])
    for holder in list(slide.placeholders):  # an untouched placeholder would show "Click to add text"
        if getattr(holder, "has_text_frame", False) and holder.has_text_frame and not holder.text_frame.text:
            holder._element.getparent().remove(holder._element)
    if operation.notes:
        slide.notes_slide.notes_text_frame.text = operation.notes
    if operation.at is not None:
        _move(deck, len(deck.slides), operation.at)


def _update_chart(deck: Any, operation: UpdateChart) -> None:
    from pptx.chart.data import CategoryChartData

    shape = find_shape(slide_at(deck, operation.slide), operation.shape)
    if not getattr(shape, "has_chart", False) or not shape.has_chart:
        raise DeckError(f'"{shape.name}" is not a chart')
    for series in operation.series:
        if len(series.values) != len(operation.categories):
            raise DeckError(
                f"series '{series.name}' has {len(series.values)} values "
                f"but there are {len(operation.categories)} categories"
            )
    factory: Any = CategoryChartData
    data = factory()
    data.categories = operation.categories
    for series in operation.series:
        data.add_series(series.name, series.values)
    shape.chart.replace_data(data)


def _update_table(deck: Any, operation: UpdateTable) -> None:
    shape = find_shape(slide_at(deck, operation.slide), operation.shape)
    if not getattr(shape, "has_table", False) or not shape.has_table:
        raise DeckError(f'"{shape.name}" is not a table')
    table = shape.table
    rows, columns = len(table.rows), len(table.columns)
    if len(operation.rows) != rows or any(len(r) != columns for r in operation.rows):
        raise DeckError(f"the table has {rows} rows and {columns} columns: give exactly that many cells")
    for row_index, values in enumerate(operation.rows):
        for column_index, value in enumerate(values):
            set_frame_text(table.cell(row_index, column_index).text_frame, [value])


def apply_operations(
    deck: Any, operations: list[Operation], read_image: Callable[[str], bytes] | None = None
) -> list[str]:
    """Runs the operations in order. Returns what was done; raises DeckError naming the failing operation."""
    done: list[str] = []
    for number, operation in enumerate(operations, start=1):
        try:
            done.append(_apply(deck, operation, read_image))
        except DeckError as error:
            raise DeckError(
                f"operation {number} ({operation.op}) failed: {error}. Nothing was saved."
            ) from error
    return done


def _apply(deck: Any, operation: Operation, read_image: Callable[[str], bytes] | None) -> str:
    if isinstance(operation, ReplaceText):
        slides = [slide_at(deck, operation.slide)] if operation.slide else list(deck.slides)
        count = sum(
            _replace_in_paragraph(paragraph, operation.find, operation.replace)
            for slide in slides
            for frame in _frames(slide, operation.notes)
            for paragraph in frame.paragraphs
        )
        if not count:
            raise DeckError(f"{operation.find!r} was not found")
        return f"replaced {count} occurrence(s) of {operation.find!r}"
    if isinstance(operation, SetText):
        shape = find_shape(slide_at(deck, operation.slide), operation.shape)
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            raise DeckError(f'"{shape.name}" has no text (it is a {shape.shape_type})')
        set_frame_text(shape.text_frame, operation.text.split("\n"))
        return f'set the text of "{shape.name}" on slide {operation.slide}'
    if isinstance(operation, SetNotes):
        slide_at(deck, operation.slide).notes_slide.notes_text_frame.text = operation.text
        return f"set the notes of slide {operation.slide}"
    if isinstance(operation, DeleteSlide):
        _delete(deck, operation.slide)
        return f"deleted slide {operation.slide}"
    if isinstance(operation, MoveSlide):
        _move(deck, operation.slide, operation.to)
        return f"moved slide {operation.slide} to position {operation.to}"
    if isinstance(operation, DuplicateSlide):
        _duplicate(deck, operation.slide, operation.at)
        return f"duplicated slide {operation.slide}"
    if isinstance(operation, AddSlide):
        _add(deck, operation)
        return f'added a slide from the layout "{operation.layout}"'
    if isinstance(operation, UpdateChart):
        _update_chart(deck, operation)
        return f"updated the chart on slide {operation.slide}"
    if isinstance(operation, UpdateTable):
        _update_table(deck, operation)
        return f"updated the table on slide {operation.slide}"
    return apply_more(deck, operation, read_image)
