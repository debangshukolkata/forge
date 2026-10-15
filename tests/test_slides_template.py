"""Building decks on a company template (P2.10) and reading one (P2.11)."""

from __future__ import annotations

import io
import json

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.errors import BuildError
from forge.slides.spec import parse_deck
from forge.slides.template import analyze_template, content_box, describe_template
from forge.slides.validate import validate_package
from tests.slide_templates import make_template

SLIDES = [
    {"layout": "title", "title": "Q3 review"},
    {"layout": "section", "title": "Results", "subtitle": "What changed"},
    {"layout": "bullets", "title": "We shipped three things", "bullets": ["Search", "Billing", "Mobile"]},
    {
        "layout": "two_column",
        "title": "Two options compared",
        "left": {"heading": "Build", "bullets": ["More control"]},
        "right": {"heading": "Buy", "bullets": ["Faster"]},
    },
    {
        "layout": "chart",
        "title": "Users doubled",
        "kind": "column",
        "categories": ["Jan", "Feb"],
        "series": [{"name": "Users", "values": [1, 2]}],
    },
    {"layout": "closing", "title": "Thank you"},
]


def build(template: bytes, **extra: object) -> tuple[object, Presentation]:
    deck = parse_deck(
        json.dumps({"title": "Deck", "subtitle": "Sub", "template": "t.potx", "slides": SLIDES, **extra})
    )
    built = build_deck(deck, lambda name: template)
    return built, Presentation(io.BytesIO(built.data))


def test_analysis_finds_roles_theme_and_fonts() -> None:
    info = analyze_template(Presentation(io.BytesIO(make_template())))
    assert info.widescreen and abs(info.scale - 1.0) < 0.01
    assert {"title", "content", "two_content", "title_only", "blank"} <= set(info.roles)
    assert info.theme.accent == "0B6E4F" and info.heading_font == "Georgia" and info.body_font == "Arial"
    text = describe_template(info, "company")
    assert "Georgia" in text and "[title_only]" in text and "Free area" in text


def test_dark_template_gives_a_dark_theme() -> None:
    info = analyze_template(Presentation(io.BytesIO(make_template(dark=True))))
    assert info.theme.dark and info.theme.background == "0E1621"


def test_four_by_three_template_is_flagged_and_refused() -> None:
    data = make_template(10, 7.5)
    assert not analyze_template(Presentation(io.BytesIO(data))).widescreen
    with pytest.raises(BuildError, match="16:9"):
        build(data)


def test_build_uses_template_layouts_and_drops_samples() -> None:
    built, prs = build(make_template())
    assert not validate_package(built.data)
    assert len(prs.slides) == len(SLIDES)  # the template's sample slide is gone
    names = [s.slide_layout.name for s in prs.slides]
    assert names[0] == "Title Slide" and names[2] == "Title and Content" and names[3] == "Two Content"
    assert all("SAMPLE" not in (s.shapes.title.text_frame.text if s.shapes.title else "") for s in prs.slides)
    assert prs.slides[2].shapes.title.text_frame.text == "We shipped three things"


def test_no_empty_placeholders_are_left() -> None:
    _, prs = build(make_template())
    for slide in prs.slides:
        for shape in slide.placeholders:
            assert (
                shape.has_text_frame is False
                or shape.text_frame.text.strip()
                or shape.placeholder_format.type.name in ("TITLE", "CENTER_TITLE")
            )


def test_chart_takes_the_templates_accent_and_sits_in_the_free_area() -> None:
    template = make_template()
    _, prs = build(template)
    chart_slide = prs.slides[4]
    chart = next(s for s in chart_slide.shapes if s.has_chart)
    info = analyze_template(Presentation(io.BytesIO(template)))
    x, y, _w, _h = content_box(info)
    assert chart.top / 914400 >= y - 0.1 and chart.left / 914400 >= x - 0.1
    assert "0B6E4F" in chart.chart._chartSpace.xml


def test_other_size_template_scales_the_canvas() -> None:
    built, prs = build(make_template(20, 11.25))
    assert prs.slide_width == (20 * 914400)
    chart = next(s for s in prs.slides[4].shapes if s.has_chart)
    assert chart.left + chart.width <= prs.slide_width and chart.top + chart.height <= prs.slide_height
    assert not validate_package(built.data)


def test_missing_layouts_fall_back_to_title_only() -> None:
    data = make_template(strip_layouts=("Two Content",))
    built, prs = build(data)
    assert len(prs.slides) == len(SLIDES) and not validate_package(built.data)


def test_a_deck_whose_master_lists_a_missing_layout_is_reported_as_damaged() -> None:
    """From the reference-deck sweep (D-399): python-pptx raises a bare KeyError while listing the layouts; the reader says what is wrong."""
    from pptx import Presentation

    from forge.slides.errors import BuildError
    from forge.slides.template import analyze_template

    real = Presentation()

    class Broken:
        @property
        def slide_layouts(self) -> list[object]:
            raise KeyError("no relationship with key 'rId1'")

        def __getattr__(self, name: str) -> object:  # everything else comes from a normal deck
            return getattr(real, name)

    with pytest.raises(BuildError, match="damaged"):
        analyze_template(Broken())
