"""The `cards` and `steps` layouts (D-294): the better layout as easy to ask for as a bullet list."""

from __future__ import annotations

import io
import json

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.errors import BuildError
from forge.slides.spec import parse_deck
from forge.slides.validate import validate_package
from tests.slide_templates import make_corporate_template


def _build(slides: list[dict], template: bytes | None = None):  # type: ignore[no-untyped-def]
    outline = {"title": "Deck", "slides": slides, **({"template": "t.pptx"} if template else {})}
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: template or b"")
    return built, Presentation(io.BytesIO(built.data))


def _texts(slide) -> list[str]:  # type: ignore[no-untyped-def]
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


CARDS = {
    "layout": "cards",
    "title": "Three reasons customers stay with us",
    "items": [
        {"heading": "Speed", "text": "Answers in under five minutes."},
        {"heading": "Care", "text": "A named contact for every account."},
        {"heading": "Value", "text": "No hidden fees, ever."},
    ],
}
STEPS = {
    "layout": "steps",
    "title": "The rollout takes four weeks",
    "steps": [
        {"heading": "Plan", "text": "Agree scope"},
        {"heading": "Build"},
        {"heading": "Pilot", "text": "Ten users"},
        {"heading": "Launch"},
    ],
}


def test_cards_draw_one_card_per_item_inside_the_slide() -> None:
    built, prs = _build([CARDS])
    assert not built.slides[0].warnings and not validate_package(built.data)
    texts = _texts(prs.slides[0])
    assert {"Speed", "Care", "Value", "No hidden fees, ever."} <= set(texts)
    cards = [s for s in prs.slides[0].shapes if not s.has_text_frame or not s.text_frame.text.strip()]
    assert len(cards) >= 3
    for shape in prs.slides[0].shapes:
        assert shape.left + shape.width <= prs.slide_width and shape.top + shape.height <= prs.slide_height


def test_numbered_cards_get_badges() -> None:
    _, prs = _build([{**CARDS, "numbered": True}])
    assert {"1", "2", "3"} <= set(_texts(prs.slides[0]))


def test_steps_are_numbered_in_order() -> None:
    built, prs = _build([STEPS])
    texts = _texts(prs.slides[0])
    assert [t for t in texts if t in "1234"][-4:] == ["1", "2", "3", "4"]  # the first "1" is the footer slide number
    assert "Pilot" in texts and "Ten users" in texts
    assert not validate_package(built.data)


@pytest.mark.parametrize(
    "slide",
    [
        {**CARDS, "items": CARDS["items"][:1]},
        {**CARDS, "items": CARDS["items"] * 2},
        {**STEPS, "steps": STEPS["steps"][:1]},
        {**STEPS, "steps": STEPS["steps"] * 2},
    ],
)
def test_wrong_item_counts_are_refused_with_the_range(slide: dict) -> None:
    with pytest.raises(BuildError, match=r"2 to [45]"):
        _build([slide])


def test_too_long_text_in_a_card_is_reported() -> None:
    long = {
        **CARDS,
        "items": [
            {"heading": "A heading that is far too long to fit on one card at all", "text": "word " * 120}
        ]
        * 3,
    }
    built, _ = _build([long])
    assert any("too long" in w for w in built.slides[0].warnings)


def test_cards_use_the_templates_own_cards_sample() -> None:
    built, prs = _build([CARDS], make_corporate_template())
    assert built.slides[0].layout.startswith("cards (template sample 2")
    texts = _texts(prs.slides[0])
    assert "Speed" in texts and "No hidden fees, ever." in texts
    assert "Acme Corp. Confidential" in texts and not [t for t in texts if "Pillar" in t]


def test_steps_without_a_template_sample_are_drawn_in_the_template_theme() -> None:
    built, prs = _build([STEPS], make_corporate_template())
    assert "Acme Corp. Confidential" in _texts(prs.slides[0])  # the template's shared decoration comes along
    assert not validate_package(built.data)
