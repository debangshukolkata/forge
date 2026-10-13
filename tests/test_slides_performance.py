"""P3.3 (D-324): large decks stay fast, pictures are made slide-sized, big files are flagged, limits grow with the deck."""

from __future__ import annotations

import io
import json
import time
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from PIL import Image
from pptx import Presentation

from forge.slides.accessibility import check_deck
from forge.slides.build import build_deck
from forge.slides.geometry import Geometry
from forge.slides.images import MAX_EDGE, prepare_image
from forge.slides.preview import timeout_for
from forge.slides.review import review_deck
from forge.slides.spec import parse_deck


def _jpeg(width: int, height: int, orientation: int | None = None) -> bytes:
    import random

    random.seed(1)
    picture = Image.frombytes(
        "RGB", (width, height), bytes(random.getrandbits(8) for _ in range(width * height * 3))
    )
    out = io.BytesIO()
    exif = Image.Exif()
    if orientation:
        exif[0x0112] = orientation
    picture.save(out, "JPEG", quality=95, exif=exif)
    return out.getvalue()


def test_a_large_photo_is_shrunk_to_slide_size_and_stays_a_jpeg() -> None:
    big = _jpeg(3600, 2400)
    smaller = prepare_image(big)
    with Image.open(io.BytesIO(smaller)) as picture:
        assert picture.format == "JPEG" and max(picture.size) == MAX_EDGE
    assert len(smaller) < len(big)


def test_a_photo_turned_by_its_exif_flag_is_stood_up() -> None:
    sideways = _jpeg(400, 200, orientation=6)  # "rotate 90 degrees clockwise to view"
    with Image.open(io.BytesIO(prepare_image(sideways))) as picture:
        assert picture.size == (200, 400)


def test_small_images_and_unknown_formats_are_returned_byte_for_byte() -> None:
    small = _jpeg(300, 200)
    assert prepare_image(small) == small
    out = io.BytesIO()
    Image.new("RGB", (50, 50), (10, 20, 30)).save(out, "GIF")
    assert prepare_image(out.getvalue()) == out.getvalue()
    assert prepare_image(b"not an image") == b"not an image"


def test_a_diagram_keeps_its_transparency_and_a_big_photo_png_becomes_a_jpeg() -> None:
    out = io.BytesIO()
    Image.new("RGBA", (3000, 1500), (255, 0, 0, 0)).save(out, "PNG")
    with Image.open(io.BytesIO(prepare_image(out.getvalue()))) as picture:
        assert picture.format == "PNG" and picture.mode == "RGBA" and picture.width == MAX_EDGE
    photo = io.BytesIO()
    Image.open(io.BytesIO(_jpeg(2000, 1500))).save(photo, "PNG")
    assert len(photo.getvalue()) > 1_500_000
    with Image.open(io.BytesIO(prepare_image(photo.getvalue()))) as picture:
        assert picture.format == "JPEG"


def test_the_build_puts_the_shrunk_picture_in_the_file() -> None:
    outline = {
        "title": "Pictures",
        "slides": [
            {"layout": "image", "title": "The big photo", "image": "big.jpg", "alt": "A noisy picture"}
        ],
    }
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: _jpeg(3600, 2400))
    prs = Presentation(io.BytesIO(built.data))
    picture = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    assert max(picture.image.size) == MAX_EDGE


def test_a_deck_over_the_mail_limit_gets_a_warning(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("forge.slides.build.MAIL_LIMIT_BYTES", 1000)
    built = build_deck(
        parse_deck(json.dumps({"title": "T", "slides": [{"layout": "title", "title": "T"}]})), lambda n: b""
    )
    assert any("mail systems" in w for w in built.deck_warnings)


def test_the_time_allowance_grows_with_the_deck(tmp_path: Path) -> None:
    outline = {
        "title": "T",
        "slides": [{"layout": "title", "title": "T"}]
        + [{"layout": "bullets", "title": f"S{i}", "bullets": ["a"]} for i in range(30)],
    }
    path = tmp_path / "d.pptx"
    path.write_bytes(build_deck(parse_deck(json.dumps(outline)), lambda n: b"").data)
    assert timeout_for(path, 180) == 180 + 2 * 31
    assert timeout_for(tmp_path / "missing.pptx", 180) == 180


def test_a_big_deck_builds_and_is_checked_in_reasonable_time() -> None:
    """The first version took 86 seconds to build 150 slides: reading inherited footer positions through python-pptx was quadratic."""
    slides = [{"layout": "title", "title": "Big deck"}]
    for i in range(90):
        slides.append(
            {
                "layout": "bullets",
                "title": f"Point number {i} is made here",
                "bullets": ["alpha", "beta"],
                "notes": "Some notes for the speaker to say.",
            }
        )
    start = time.time()
    built = build_deck(parse_deck(json.dumps({"title": "Big", "slides": slides})), lambda n: b"")
    prs = Presentation(io.BytesIO(built.data))
    review_deck(prs)
    check_deck(prs)
    assert time.time() - start < 40  # about 5 s on the build machine; the margin is for a busy one


def test_the_geometry_table_answers_like_python_pptx_for_every_shape() -> None:
    slides = [
        {"layout": "title", "title": "G"},
        {"layout": "bullets", "title": "Point", "bullets": ["a"]},
        {"layout": "chart", "title": "Chart", "categories": ["x"], "series": [{"name": "s", "values": [1]}]},
    ]
    prs = Presentation(
        io.BytesIO(build_deck(parse_deck(json.dumps({"title": "G", "slides": slides})), lambda n: b"").data)
    )
    geometry = Geometry()
    for slide in prs.slides:
        for shape in slide.shapes:
            assert geometry.box(shape) == (shape.left, shape.top, shape.width, shape.height), shape.name
