"""D-388: the brand kit's logo appears in the header of Word documents and in the PDFs reportlab builds."""

from __future__ import annotations

import io

import pytest

from forge.brand.kit import BrandKit
from forge.documents.build import build_document
from forge.documents.spec import DocumentError, DocumentSpec, Paragraph
from forge.pdfs.build import PdfSpec, build_pdf


def _png() -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (120, 40), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


def _load(path: str) -> bytes:
    if path == "assets/logo.png":
        return _png()
    raise FileNotFoundError(path)


def test_the_word_header_carries_the_logo() -> None:
    from docx import Document

    spec = DocumentSpec(
        title="Report",
        header="Acme",
        blocks=[Paragraph(text="Hello")],
        brand=BrandKit(logo="assets/logo.png"),
    )
    document = Document(io.BytesIO(build_document(spec, _load)))
    header = document.sections[0].header
    assert len(header.paragraphs[0]._p.xpath(".//pic:pic")) == 1
    assert "Acme" in header.paragraphs[0].text


def test_a_missing_logo_is_a_clear_error() -> None:
    spec = DocumentSpec(
        title="Report", blocks=[Paragraph(text="Hello")], brand=BrandKit(logo="assets/none.png")
    )
    with pytest.raises(DocumentError, match="logo"):
        build_document(spec, _load)


def test_the_pdf_draws_the_logo() -> None:
    from forge.pdfs.build import PdfError

    spec = PdfSpec(
        title="Report",
        blocks=[{"type": "paragraph", "text": "Hello"}],
        brand=BrandKit(logo="assets/logo.png"),
    )
    data = build_pdf(spec, _load)
    assert b"/Subtype /Image" in data or b"/Image" in data
    with pytest.raises(PdfError, match="logo"):
        build_pdf(spec.model_copy(update={"brand": BrandKit(logo="assets/none.png")}), _load)


def test_a_deck_puts_the_logo_at_the_top_right_of_every_slide() -> None:
    from pptx import Presentation

    from forge.slides.build import build_deck
    from forge.slides.spec import parse_deck

    outline = {
        "title": "Q3",
        "brand": {"logo": "assets/logo.png"},
        "slides": [
            {"layout": "title", "title": "Q3", "subtitle": "What shipped"},
            {"layout": "bullets", "title": "We shipped three things", "bullets": ["SSO", "Search"]},
        ],
    }
    built = build_deck(parse_deck(outline), _load)
    presentation = Presentation(io.BytesIO(built.data))
    for slide in presentation.slides:
        logos = [shape for shape in slide.shapes if shape.name == "Brand logo"]
        assert len(logos) == 1
        assert logos[0].top < 914400 // 2 and logos[0].left + logos[0].width <= presentation.slide_width
