"""P2.9 (D-314): accessibility checks and repairs, and chart colours a colour-blind viewer can tell apart."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches

from forge.slides.accessibility import check_deck, improve_deck
from forge.slides.build import build_deck
from forge.slides.colorvision import (
    _worst_pair,
    accessible_palette,
    palette_conflicts,
    simulate,
)
from forge.slides.spec import parse_deck
from forge.slides.themes import THEMES
from forge.slides.validate import validate_package


def _levels(problems: list) -> str:  # type: ignore[type-arg]
    return " | ".join(p.render() for p in problems)


def test_a_deck_made_elsewhere_gets_every_kind_of_finding() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank: no title
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Some words"
    slide.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(3), Inches(2), Inches(1))
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1))
    table.table.first_row = False
    text = _levels(check_deck(prs))
    assert "the document has no title" in text
    assert "has no title: screen reader users" in text
    assert "has no alt text" in text and "no header row" in text
    assert "no language set" in text


def test_building_a_deck_repairs_what_needs_no_guessing_and_the_repair_is_valid() -> None:
    outline = {
        "title": "Quarterly review",
        "language": "de-DE",
        "slides": [
            {"layout": "title", "title": "Quarterly review"},
            {
                "layout": "chart",
                "title": "Users doubled since January",
                "categories": ["Jan", "Feb"],
                "series": [{"name": "Users", "values": [1, 2]}],
            },
            {
                "layout": "table",
                "title": "Numbers by region at a glance",
                "columns": ["Region", "Revenue"],
                "rows": [["North", 1]],
            },
        ],
    }
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: b"")
    assert not validate_package(built.data)
    prs = Presentation(io.BytesIO(built.data))
    assert prs.core_properties.title == "Quarterly review"
    languages = {
        r._r.rPr.get("lang")
        for s in prs.slides
        for sh in s.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
        for r in p.runs
    }
    assert languages == {"de-DE"}
    problems = check_deck(prs)
    assert [p for p in problems if p.level == "error"] == []
    for slide in list(prs.slides)[1:]:  # content slides: the title is read first, the footer last
        shapes = list(slide.shapes)
        assert shapes[0] == slide.shapes.title
        assert shapes[-1].text_frame.text.strip().isdigit() or "Quarterly" in shapes[-1].text_frame.text


def test_the_reading_order_is_repaired_title_first_footer_last() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text_frame.text = "The point of the slide"
    footer = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(4), Inches(0.3))
    footer.text_frame.text = "Company footer"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    body.text_frame.text = "The content"
    spTree = slide.shapes._spTree
    element = slide.shapes.title._element
    spTree.remove(element)  # make it the last shape
    spTree.append(element)
    assert any("title is read after" in p.message for p in check_deck(prs))
    improve_deck(prs)
    order = [s.text_frame.text for s in slide.shapes]
    assert order == ["The point of the slide", "The content", "Company footer"]
    assert not any(("is read after" in p.message or "read before" in p.message) for p in check_deck(prs))


def test_alt_text_for_charts_and_tables_comes_from_the_slide_title_and_header_rows_are_set() -> None:
    outline = {
        "title": "T",
        "slides": [
            {
                "layout": "table",
                "title": "Revenue by region",
                "columns": ["Region", "Revenue"],
                "rows": [["North", 1]],
            }
        ],
    }
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: b"")
    prs = Presentation(io.BytesIO(built.data))
    table = next(s for s in prs.slides[0].shapes if getattr(s, "has_table", False) and s.has_table)
    assert table.table.first_row
    assert "Revenue by region" in table._element.xpath("(.//p:cNvPr)[1]/@descr")[0]


def test_vague_link_text_is_flagged() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text_frame.text = "Links"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "click here"
    run.hyperlink.address = "https://example.com"
    assert any("says nothing about where it goes" in p.message for p in check_deck(prs))


def test_colour_vision_simulation_and_conflicts() -> None:
    red, green = simulate("#C0392B", "deuteranopia"), simulate("#2E8B57", "deuteranopia")
    assert max(abs(a - b) for a, b in zip(red, green, strict=True)) < 0.35  # red and green come together
    assert palette_conflicts(["#C0392B", "#2E8B57"])
    assert palette_conflicts(["#000000", "#FFFFFF"]) == []


@pytest.mark.parametrize("name", sorted(THEMES))
def test_two_and_three_series_use_colours_a_colour_blind_viewer_can_tell_apart(name: str) -> None:
    theme = THEMES[name]
    background = "#" + theme.background.lstrip("#")
    for count in (2, 3):
        colours = accessible_palette(theme.palette, count, background)
        assert (
            len(colours) == count and colours[0].lstrip("#").upper() == theme.palette[0].lstrip("#").upper()
        )
        assert _worst_pair(colours)[0] >= 17.5, (name, count, colours)


def test_more_series_than_the_palette_repeat_and_one_colour_is_the_accent() -> None:
    palette = THEMES["clean"].palette
    assert accessible_palette(palette, 1) == ["#" + palette[0].lstrip("#").upper()]
    assert len(accessible_palette(palette, 9)) == 9


async def test_the_check_tool_reports_and_a_clean_deck_passes(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.presentation import CheckAccessibility
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    context = ToolContext(workspace=workspace)
    bad = Presentation()
    bad.slides.add_slide(bad.slide_layouts[6])
    buffer = io.BytesIO()
    bad.save(buffer)
    workspace.write_bytes("docs/bad.pptx", buffer.getvalue(), "x")
    built = build_deck(
        parse_deck(json.dumps({"title": "Good", "slides": [{"layout": "title", "title": "Good deck"}]})),
        lambda n: b"",
    )
    workspace.write_bytes("docs/good.pptx", built.data, "x")
    tool = CheckAccessibility()
    found = await tool.run(tool.Args(path="docs/bad.pptx"), context)
    assert found.ok and "error(s)" in found.content and "slide 1: has no title" in found.content
    clean = await tool.run(tool.Args(path="docs/good.pptx"), context)
    assert clean.ok and "No accessibility problems found." in clean.content


def _png() -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (20, 10), (10, 80, 200)).save(out, "PNG")
    return out.getvalue()
