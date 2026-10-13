"""D-335: a picture goes into a sample's frame (a circle-cropped picture, a shape filled with a picture, an empty placeholder), covering it."""

from __future__ import annotations

import io

import pytest

pytest.importorskip("pptx")
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from forge.slides.sample_objects import _cover_crop, picture_frames, put_picture
from tests.slide_templates import make_template

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _png(width: int, height: int, colour: tuple[int, int, int]) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(out, "PNG")
    return out.getvalue()


def _slide() -> tuple[Presentation, object]:
    prs = Presentation(io.BytesIO(make_template(sample_slide=False)))
    return prs, prs.slides.add_slide(prs.slide_layouts[5])


def test_the_crop_covers_the_frame_without_stretching() -> None:
    assert _cover_crop(200, 100, 100, 100) == (
        0.25,
        0.0,
        0.25,
        0.0,
    )  # a wide picture in a square: cut the sides
    assert _cover_crop(100, 200, 100, 100) == (0.0, 0.25, 0.0, 0.25)  # a tall one: cut top and bottom
    assert _cover_crop(100, 100, 200, 100) == (0.0, 0.25, 0.0, 0.25)
    assert _cover_crop(160, 90, 16, 9) == (0.0, 0.0, 0.0, 0.0)


def test_a_circle_cropped_picture_keeps_its_shape_and_gets_the_new_picture_covering_it() -> None:
    _prs, slide = _slide()
    frame = slide.shapes.add_picture(
        io.BytesIO(_png(50, 50, (200, 0, 0))), Inches(1), Inches(1), Inches(2), Inches(2)
    )  # type: ignore[attr-defined]
    frame.auto_shape_type = MSO_SHAPE.OVAL
    assert picture_frames(slide) == [frame] or [f.shape_id for f in picture_frames(slide)] == [frame.shape_id]
    put_picture(slide, frame, _png(400, 200, (0, 0, 200)))
    picture = next(s for s in slide.shapes if s.shape_type == 13)  # type: ignore[attr-defined]
    assert picture.auto_shape_type == MSO_SHAPE.OVAL and (picture.width, picture.height) == (
        Inches(2),
        Inches(2),
    )
    assert (
        abs(picture.crop_left - 0.25) < 1e-6
        and abs(picture.crop_right - 0.25) < 1e-6
        and picture.image.size == (400, 200)
    )


def test_a_shape_filled_with_a_picture_gets_a_source_rectangle_crop() -> None:
    _prs, slide = _slide()
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(2), Inches(3))  # type: ignore[attr-defined]
    _, relation = slide.part.get_or_add_image_part(io.BytesIO(_png(30, 30, (0, 200, 0))))  # type: ignore[attr-defined]
    properties = oval._element.find(f"{{{NS_P}}}spPr")
    fill = etree.fromstring(
        f'<a:blipFill xmlns:a="{NS_A}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><a:blip r:embed="{relation}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
    )
    properties.insert(2, fill)
    assert [f.shape_id for f in picture_frames(slide)] == [oval.shape_id]
    put_picture(slide, oval, _png(300, 200, (9, 9, 9)))
    blip_fill = oval._element.find(f"{{{NS_P}}}spPr").find(f"{{{NS_A}}}blipFill")
    source = blip_fill.find(f"{{{NS_A}}}srcRect")
    assert (
        source is not None
        and int(source.get("l")) > 0
        and int(source.get("r")) > 0
        and source.get("t") == "0"
    )
    assert (
        blip_fill.find(f"{{{NS_A}}}blip").get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        != relation
    )


def test_frames_are_listed_in_reading_order_and_plain_shapes_are_not_frames() -> None:
    _prs, slide = _slide()
    right = slide.shapes.add_picture(
        io.BytesIO(_png(10, 10, (1, 1, 1))), Inches(6), Inches(1), Inches(1), Inches(1)
    )  # type: ignore[attr-defined]
    left = slide.shapes.add_picture(
        io.BytesIO(_png(10, 10, (1, 1, 1))), Inches(1), Inches(1), Inches(1), Inches(1)
    )  # type: ignore[attr-defined]
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(1), Inches(1))  # type: ignore[attr-defined]
    assert [f.shape_id for f in picture_frames(slide)] == [left.shape_id, right.shape_id]
