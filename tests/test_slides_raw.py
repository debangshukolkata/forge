"""P2.15 (D-321): unpack a deck to XML, edit, repack with validation; hostile packages are refused."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides import rawpackage
from forge.slides.build import build_deck
from forge.slides.open_deck import DeckError
from forge.slides.spec import parse_deck


def _deck_bytes() -> bytes:
    outline = {
        "title": "Raw",
        "slides": [
            {"layout": "title", "title": "Raw"},
            {"layout": "bullets", "title": "Point to edit", "bullets": ["alpha", "beta"]},
        ],
    }
    return build_deck(parse_deck(json.dumps(outline)), lambda name: b"").data


def _zip(entries: dict[str, bytes], method: int = zipfile.ZIP_DEFLATED) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", method) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return out.getvalue()


def _replace_part(data: bytes, name: str, change) -> bytes:  # type: ignore[no-untyped-def]
    source = zipfile.ZipFile(io.BytesIO(data))
    return _zip(
        {
            i.filename: (change(source.read(i.filename)) if i.filename == name else source.read(i.filename))
            for i in source.infolist()
        }
    )


def test_unpack_pretty_prints_xml_and_leaves_binary_parts() -> None:
    unpacked = rawpackage.unpack(_deck_bytes())
    slide = unpacked.xml["ppt/slides/slide2.xml"].decode()
    assert slide.count("\n") > 20 and "Point to edit" in slide
    assert "[Content_Types].xml" in unpacked.xml and "ppt/presentation.xml" in unpacked.xml
    assert all(not name.endswith((".xml", ".rels")) for name in unpacked.binary)


def test_an_edit_in_the_xml_survives_packing_and_the_deck_opens() -> None:
    original = _deck_bytes()
    unpacked = rawpackage.unpack(original)
    files = dict(unpacked.xml)
    files["ppt/slides/slide2.xml"] = files["ppt/slides/slide2.xml"].replace(
        b"Point to edit", b"Point edited by hand"
    )
    data, notes = rawpackage.pack(original, files, set(unpacked.xml))
    assert notes == []
    prs = Presentation(io.BytesIO(data))
    assert prs.slides[1].shapes.title.text_frame.text == "Point edited by hand"


def test_packing_refuses_a_deck_that_powerpoint_would_repair() -> None:
    original = _deck_bytes()
    unpacked = rawpackage.unpack(original)
    files = dict(unpacked.xml)
    files["ppt/slides/_rels/slide2.xml.rels"] = files["ppt/slides/_rels/slide2.xml.rels"].replace(
        b"slideLayout", b"slideLayoutMissing"
    )
    with pytest.raises(DeckError, match="structural faults"):
        rawpackage.pack(original, files, set(unpacked.xml))


def test_new_external_links_are_flagged_and_a_part_still_in_use_cannot_be_removed() -> None:
    original = _deck_bytes()
    unpacked = rawpackage.unpack(original)
    files = dict(unpacked.xml)
    rels = files["ppt/slides/_rels/slide2.xml.rels"].decode()
    files["ppt/slides/_rels/slide2.xml.rels"] = rels.replace(
        "</Relationships>",
        '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="https://example.org/x" TargetMode="External"/></Relationships>',
    ).encode()
    _, notes = rawpackage.pack(original, files, set(unpacked.xml))
    assert any("external links" in n and "https://example.org/x" in n for n in notes)
    del files["docProps/app.xml"]  # still referenced: the validator must refuse
    with pytest.raises(DeckError, match="not in the file"):
        rawpackage.pack(original, files, set(unpacked.xml))


@pytest.mark.parametrize(
    ("entries", "message"),
    [
        ({"../evil.xml": b"x"}, "unsafe entry name"),
        ({"/abs.xml": b"x"}, "unsafe entry name"),
        ({"C:evil.xml": b"x"}, "unsafe entry name"),
        ({"ppt/a.xml": b"x" * 200_000_000}, "zip bomb"),
    ],
)
def test_hostile_archives_are_refused_before_anything_is_read(
    entries: dict[str, bytes], message: str
) -> None:
    with pytest.raises(DeckError, match=message):
        rawpackage.check_archive(_zip(entries))


def test_too_many_entries_and_not_a_zip() -> None:
    many = {f"p{i}.txt": b"" for i in range(rawpackage.MAX_ENTRIES + 1)}
    with pytest.raises(DeckError, match="entries"):
        rawpackage.check_archive(_zip(many))
    with pytest.raises(DeckError, match="not a valid"):
        rawpackage.check_archive(b"not a zip")


def test_xml_with_a_doctype_or_entities_and_macros_are_refused() -> None:
    bad = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><x>&a;</x>'
    with pytest.raises(DeckError, match="DOCTYPE"):
        rawpackage.unpack(_replace_part(_deck_bytes(), "docProps/app.xml", lambda _: bad))
    with_macro = _zip({**{i: b"" for i in ("[Content_Types].xml",)}, "ppt/vbaProject.bin": b"x"})
    with pytest.raises(DeckError, match="macros"):
        rawpackage.unpack(with_macro)
    original = _deck_bytes()
    unpacked = rawpackage.unpack(original)
    files = {**unpacked.xml, "ppt/vbaProject.bin": b"x"}
    with pytest.raises(DeckError, match="macros"):
        rawpackage.pack(original, files, set(unpacked.xml))
    files = {**unpacked.xml, "ppt/slides/slide9.xml": bad}
    with pytest.raises(DeckError, match="DOCTYPE"):
        rawpackage.pack(original, files, set(unpacked.xml))


async def test_the_tools_unpack_edit_and_pack(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.presentation_raw import PackPresentation, UnpackPresentation
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    context = ToolContext(workspace=workspace)
    workspace.write_bytes("docs/deck.pptx", _deck_bytes(), "x")
    unpacked = await UnpackPresentation().run(UnpackPresentation.Args(path="docs/deck.pptx"), context)
    assert unpacked.ok and "work/deck_xml/ppt/slides/slide2.xml" in unpacked.content
    slide = "work/deck_xml/ppt/slides/slide2.xml"
    text, _ = workspace.read_text(slide)
    workspace.write_text(slide, text.replace("Point to edit", "Edited through XML"), "edit")
    packed = await PackPresentation().run(
        PackPresentation.Args(folder="work/deck_xml", output="docs/edited.pptx"), context
    )
    assert packed.ok, packed.content
    assert (
        Presentation(str(workspace.path_of("docs/edited.pptx"))).slides[1].shapes.title.text_frame.text
        == "Edited through XML"
    )
    assert not workspace.path_of("work/deck_xml/ppt/slides/slide2.xml").exists()  # the folder was removed
    refused = await PackPresentation().run(
        PackPresentation.Args(folder="docs", output="docs/x.pptx"), context
    )
    assert not refused.ok and "not made by unpack_presentation" in refused.content


@pytest.mark.parametrize(
    "name",
    ["ppt/a.xml:stream", "ppt/CON.xml", "ppt/nul", "ppt/x.xml.", "ppt/a<b>.xml", "ppt/with\x01ctrl.xml"],
)
def test_names_that_windows_treats_specially_are_refused(name: str) -> None:
    with pytest.raises(DeckError, match="unsafe entry name"):
        rawpackage.check_archive(_zip({name: b"x"}))


def test_the_normal_part_names_of_a_deck_pass_the_name_check() -> None:
    rawpackage.check_package(_deck_bytes())


def test_every_deck_the_tools_open_is_checked_first(tmp_path: Path) -> None:
    """P3.4 (D-323): open_deck refuses bombs, unsafe names and DOCTYPE XML before python-pptx reads a byte."""
    from forge.slides.open_deck import open_deck

    bad = _replace_part(
        _deck_bytes(),
        "docProps/app.xml",
        lambda _: b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><x/>',
    )
    for name, data, message in (
        ("doctype.pptx", bad, "DOCTYPE"),
        ("slip.pptx", _zip({"../evil.xml": b"x"}), "unsafe entry name"),
        ("bomb.pptx", _zip({"ppt/a.xml": b"x" * 150_000_000}), "zip bomb"),
    ):
        path = tmp_path / name
        path.write_bytes(data)
        with pytest.raises(DeckError, match=message):
            open_deck(path)
    good = tmp_path / "good.pptx"
    good.write_bytes(_deck_bytes())
    assert len(open_deck(good).slides) == 2


def test_powerpoint_is_started_with_macros_force_disabled_and_the_setting_restored() -> None:
    from forge.slides import convert, measure, preview

    for module in (convert, measure, preview):
        script = module.SCRIPT
        assert "$app.AutomationSecurity = 3" in script and "$app.AutomationSecurity = $oldSecurity" in script
        assert script.index("AutomationSecurity = 3") < script.index("Presentations.Open")
