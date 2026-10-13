"""D-336: the design of an existing deck, read as a summary (P2.24)."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.design_reader import describe_design, read_design
from forge.slides.spec import parse_deck
from tests.slide_templates import make_infographic_template


def _design():  # type: ignore[no-untyped-def]
    return read_design(Presentation(io.BytesIO(make_infographic_template())))


def test_size_theme_dark_and_light_titles_and_the_repeated_frame_are_found() -> None:
    design = _design()
    assert (design.width, round(design.height, 2)) == (10.0, 5.62) and design.slides == 10
    assert design.colours["accent1"] == "0B6E4F" and (design.heading_font, design.body_font) == (
        "Georgia",
        "Arial",
    )
    assert (design.dark_slides, design.light_slides) == (6, 4)
    assert design.title["left"] == 0.5 and design.title["width"] == 9.0
    assert design.margins[0] <= 0.5  # the smallest start of the text, not the footer's 4 in
    names = {r["name"] for r in design.repeated}
    assert "Company logo" in names and any(r["text"] == "Acme confidential" for r in design.repeated)
    assert all(r["share"] >= 0.5 for r in design.repeated)
    assert design.layouts[0][0] == "Title Only"


def test_the_brand_kit_follows_the_side_the_deck_stays_on() -> None:
    brand = _design().brand
    assert brand["accent"] == "#0B6E4F" and brand["title_font"] == "Georgia" and brand["body_font"] == "Arial"
    assert brand["background"] == "#0A0A0A" and brand["text"] == "#FFFFFF"  # mostly dark slides


def test_a_deck_made_by_debcraft_can_be_read_too_and_the_summary_names_the_frame() -> None:
    outline = {
        "title": "T",
        "footer": "Acme internal",
        "slides": [{"layout": "title", "title": "T"}]
        + [{"layout": "bullets", "title": f"Point {i}", "bullets": ["a"]} for i in range(5)],
    }
    prs = Presentation(io.BytesIO(build_deck(parse_deck(json.dumps(outline)), lambda n: b"").data))
    design = read_design(prs)
    text = describe_design(design, "built.pptx")
    assert (
        "Design of built.pptx: 6 slides" in text
        and "theme colours:" in text
        and "brand kit to use in an outline" in text
    )
    assert any(r["text"] == "Acme internal" for r in design.repeated)
    assert json.dumps(design.as_dict())  # the same data as plain values for programs


def test_a_tiny_deck_has_no_repeated_shapes_claimed() -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text_frame.text = "One"
    assert read_design(prs).repeated == []


async def test_the_tool_describes_and_advises(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.presentation import DescribeDesign
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    workspace.write_bytes("docs/style.pptx", make_infographic_template(), "x")
    result = await DescribeDesign().run(
        DescribeDesign.Args(path="docs/style.pptx"), ToolContext(workspace=workspace)
    )
    assert (
        result.ok
        and "Design of style.pptx" in result.content
        and "build on it as a template" in result.content
    )
    missing = await DescribeDesign().run(
        DescribeDesign.Args(path="docs/none.pptx"), ToolContext(workspace=workspace)
    )
    assert not missing.ok
