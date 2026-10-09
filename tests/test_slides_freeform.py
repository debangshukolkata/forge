"""Free-form slides (D-242): the model places every element; the builder checks what a designer would see."""

from __future__ import annotations

import io
import json
from typing import Any

import pytest

pptx = pytest.importorskip("pptx")

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from forge.slides.build import BuildError, build_deck  # noqa: E402
from forge.slides.freeform import Placed, contrast, resolve  # noqa: E402
from forge.slides.read import describe_deck  # noqa: E402
from forge.slides.spec import SpecError, parse_deck  # noqa: E402
from forge.slides.themes import THEMES  # noqa: E402


def _png(width: int = 200, height: int = 100) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 40, 40)).save(buffer, "PNG")
    return buffer.getvalue()


def _build(elements: list[dict[str, Any]], theme: str = "clean", **slide: Any) -> Any:
    outline = {
        "title": "Deck",
        "theme": theme,
        "slides": [{"layout": "freeform", "title": "A free slide", "elements": elements, **slide}],
    }
    return build_deck(parse_deck(json.dumps(outline)), lambda name: _png(400, 100))


def _warnings(elements: list[dict[str, Any]], **slide: Any) -> list[str]:
    return _build(elements, **slide).slides[0].warnings


def test_a_designed_slide_builds_with_every_kind_of_element_and_no_warnings() -> None:
    built = _build(
        [
            {"type": "shape", "shape": "rect", "x": 0, "y": 0, "w": 0.35, "h": 7.5, "fill": "accent"},
            {
                "type": "text",
                "x": 1.0,
                "y": 1.8,
                "w": 6,
                "h": 1.5,
                "text": "Big idea",
                "size": 54,
                "bold": True,
                "font": "title",
            },
            {
                "type": "text",
                "x": 1.0,
                "y": 3.6,
                "w": 5,
                "h": 1.5,
                "text": "First\nSecond",
                "bullets": True,
                "fill": "surface",
                "rounded": True,
            },
            {
                "type": "shape",
                "shape": "oval",
                "x": 8.0,
                "y": 1.9,
                "w": 2.2,
                "h": 2.2,
                "fill": "series2",
                "text": "42%",
                "size": 36,
                "bold": True,
                "color": "text",
            },
            {"type": "line", "x1": 7.4, "y1": 2.7, "x2": 7.9, "y2": 2.7, "arrow": True, "dash": True},
            {
                "type": "image",
                "x": 8.0,
                "y": 4.2,
                "w": 4.0,
                "h": 2.0,
                "path": "a.png",
                "alt": "A red box",
                "fit": "cover",
            },
            {
                "type": "chart",
                "x": 6.2,
                "y": 3.9,
                "w": 1.7,
                "h": 2.4,
                "kind": "pie",
                "categories": ["a", "b"],
                "series": [{"name": "s", "values": [1, 2]}],
            },
            {
                "type": "table",
                "x": 1.0,
                "y": 5.6,
                "w": 4.0,
                "h": 1.2,
                "columns": ["a", "b"],
                "rows": [["1", "2"]],
                "size": 14,
            },
        ],
        notes="Say it slowly.",
    )
    assert built.slides[0].warnings == []
    deck = Presentation(io.BytesIO(built.data))
    slide = deck.slides[0]
    assert slide.shapes.title.text_frame.text == "A free slide"  # a real title, even on a free slide
    assert slide.notes_slide.notes_text_frame.text == "Say it slowly."
    kinds = {shape.shape_type for shape in slide.shapes}
    assert any(shape.has_chart for shape in slide.shapes if hasattr(shape, "has_chart"))
    assert any(shape.has_table for shape in slide.shapes if hasattr(shape, "has_table")) and kinds
    assert "[id " in describe_deck(deck)  # reading a free slide back works like any other


def test_cover_fit_crops_the_picture_to_fill_its_box_and_alt_text_is_kept() -> None:
    built = _build(
        [{"type": "image", "x": 1, "y": 1, "w": 4, "h": 4, "path": "wide.png", "alt": "Wide", "fit": "cover"}]
    )
    picture = next(s for s in Presentation(io.BytesIO(built.data)).slides[0].shapes if s.shape_type == 13)
    assert (
        picture.crop_left > 0 and picture.crop_right > 0 and picture.crop_top == 0
    )  # a wide picture loses its sides
    assert (picture.width, picture.height) == (picture.width, picture.width)  # the box stays square
    assert picture._element.nvPicPr.cNvPr.get("descr") == "Wide"


def test_the_title_can_be_hidden_but_stays_for_navigation() -> None:
    built = _build(
        [{"type": "text", "x": 1, "y": 1, "w": 6, "h": 1, "text": "Hello"}],
        title_visible=False,
        background="accent",
    )
    slide = Presentation(io.BytesIO(built.data)).slides[0]
    assert (
        slide.shapes.title.text_frame.text == "A free slide" and slide.shapes.title.top < 0
    )  # parked above the slide


def test_the_checks_name_each_problem_by_element() -> None:
    problems = _warnings(
        [
            {"type": "text", "x": 0.1, "y": 2, "w": 6, "h": 1, "text": "Hugging the edge"},
            {
                "type": "text",
                "x": 3,
                "y": 2.4,
                "w": 6,
                "h": 1.5,
                "text": "Collides",
                "fill": "accent",
                "color": "#777777",
            },
            {"type": "text", "x": 9, "y": 6, "w": 6, "h": 2, "text": "Off the slide"},
            {"type": "image", "x": 1, "y": 5, "w": 2, "h": 1, "path": "x.png"},
            {
                "type": "text",
                "x": 1,
                "y": 4.2,
                "w": 1.5,
                "h": 0.4,
                "text": "A very long sentence that cannot fit in this small box at all",
                "size": 20,
            },
        ]
    )
    joined = "\n".join(problems)
    assert "element 1 (text: Hugging the edge): text sits only 0.1 in from the slide edge" in joined
    assert "element 2 (text: Collides) overlaps element 1" in joined
    assert "element 2 (text: Collides): low contrast" in joined
    assert "element 3 (text: Off the slide): extends" in joined and "beyond the slide" in joined
    assert "element 4 (image: x.png): the picture has no alt text" in joined
    assert "element 5" in joined and "does not fit its box even at 12 pt" in joined


def test_a_card_around_text_is_not_an_overlap_and_a_filled_background_sets_the_contrast() -> None:
    assert (
        _warnings(
            [
                {"type": "shape", "shape": "rect", "x": 1, "y": 2, "w": 5, "h": 3, "fill": "text"},
                {
                    "type": "text",
                    "x": 1.3,
                    "y": 2.3,
                    "w": 4.4,
                    "h": 1,
                    "text": "Light on dark card",
                    "color": "background",
                    "size": 24,
                },
            ]
        )
        == []
    )
    low = _warnings(
        [
            {"type": "shape", "shape": "rect", "x": 1, "y": 2, "w": 5, "h": 3, "fill": "text"},
            {
                "type": "text",
                "x": 1.3,
                "y": 2.3,
                "w": 4.4,
                "h": 1,
                "text": "Dark on dark card",
                "color": "text",
                "size": 24,
            },
        ]
    )
    assert any("low contrast" in w for w in low)  # judged against the card behind it, not the slide


def test_text_shrinks_to_fit_before_it_is_reported() -> None:
    elements = [
        {
            "type": "text",
            "x": 1,
            "y": 2,
            "w": 3,
            "h": 1.2,
            "text": "Eleven words in a small box that has to shrink to fit",
            "size": 28,
        }
    ]
    built = _build(elements)
    assert built.slides[0].warnings == []
    box = next(
        s
        for s in Presentation(io.BytesIO(built.data)).slides[0].shapes
        if s.has_text_frame and "Eleven" in s.text_frame.text
    )
    assert box.text_frame.paragraphs[0].runs[0].font.size.pt < 28


def test_bad_input_is_explained_in_words() -> None:
    with pytest.raises(SpecError) as raised:
        parse_deck(
            json.dumps(
                {
                    "title": "T",
                    "slides": [
                        {
                            "layout": "freeform",
                            "title": "x",
                            "elements": [
                                {
                                    "type": "text",
                                    "x": 1,
                                    "y": 1,
                                    "w": 2,
                                    "h": 1,
                                    "text": "a",
                                    "color": "purple",
                                }
                            ],
                        }
                    ],
                }
            )
        )
    assert "color" in str(raised.value) or "pattern" in str(raised.value)
    with pytest.raises(SpecError):
        parse_deck(
            json.dumps({"title": "T", "slides": [{"layout": "freeform", "title": "x", "elements": []}]})
        )
    with pytest.raises(BuildError, match="pie chart takes exactly one series"):
        _build(
            [
                {
                    "type": "chart",
                    "x": 1,
                    "y": 1,
                    "w": 4,
                    "h": 3,
                    "kind": "pie",
                    "categories": ["a"],
                    "series": [{"name": "a", "values": [1]}, {"name": "b", "values": [2]}],
                }
            ]
        )


def test_colours_resolve_from_the_theme_and_contrast_is_measured() -> None:
    theme = THEMES["warm"]
    assert resolve(theme, "accent") == theme.accent and resolve(theme, "#abcdef") == "ABCDEF"
    assert resolve(theme, "series2") == theme.palette[1]
    assert contrast("000000", "FFFFFF") == pytest.approx(21.0)
    assert contrast("777777", "777777") == pytest.approx(1.0)
    box = Placed(1, "a", 1, 1, 4, 4)
    assert box.contains(Placed(2, "b", 2, 2, 1, 1)) and not box.contains(Placed(3, "c", 4, 4, 3, 3))
    assert box.overlap_area(Placed(3, "c", 4, 4, 3, 3)) == pytest.approx(1.0)
