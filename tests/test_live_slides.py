"""Live acceptance of presentation making (D-240): the real model loads the skill, writes an outline, builds
a deck, has PowerPoint draw it and looks at the slides. Run with: pytest -m live tests/test_live_slides.py"""

from __future__ import annotations

import pytest

from forge.protocol.events import EventType
from tests.test_live_agent import events_of, make_host, run_turn

pptx = pytest.importorskip("pptx")
pytestmark = pytest.mark.live
__all__ = ["make_host"]  # the fixture comes from test_live_agent


def tool_names(host) -> list[str]:  # type: ignore[no-untyped-def]
    return [e.payload["name"] for e in events_of(host, EventType.TOOL_CALL_STARTED)]


async def test_model_makes_a_deck_looks_at_it_and_delivers_a_pptx(make_host) -> None:  # type: ignore[no-untyped-def]
    host = make_host("auto")
    await run_turn(
        host,
        "Make a short PowerPoint presentation (about 6 slides) for a team meeting on why we should add "
        "automated tests to our claims service: the problem, what it costs us, the plan, and the next steps. "
        "Save it as docs/testing-plan.pptx. Use only facts I gave you: do not invent numbers.",
    )
    names = tool_names(host)
    assert "load_skill" in names and "build_presentation" in names, names
    deck = host.workspace.path_of("docs/testing-plan.pptx")
    assert deck.is_file()
    from pptx import Presentation

    slides = Presentation(str(deck)).slides
    assert 4 <= len(slides) <= 12
    assert all(slide.shapes.title is not None and slide.shapes.title.text_frame.text for slide in slides)
    from forge.slides import preview

    if preview.available():  # the look-and-fix loop only exists where PowerPoint does
        assert "preview_presentation" in names and "view_image" in names, names

    # Keep the result where a person can look at it (test-artifacts/ is not committed).
    import shutil

    from tests.conftest import REPO_ROOT

    keep = REPO_ROOT / "test-artifacts" / "slides-live"
    shutil.rmtree(keep, ignore_errors=True)
    keep.mkdir(parents=True)
    shutil.copy(deck, keep / deck.name)
    previews = host.workspace.forge_dir / "reports" / "slides"
    if previews.is_dir():
        shutil.copytree(previews, keep / "preview")
