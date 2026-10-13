"""D-332, D-333: a template like the company one (infographic samples, charts, a table, misspelt covers, sticky notes) is used as it is."""

from __future__ import annotations

import io
import json

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.edit_find import walk
from forge.slides.spec import parse_deck
from forge.slides.template_samples import analyze_samples
from forge.slides.validate import validate_package
from tests.slide_templates import make_infographic_template


def _build(outline: dict) -> tuple[Presentation, list]:  # type: ignore[type-arg]
    template = make_infographic_template()
    deck = parse_deck(json.dumps({"title": "T", "template": "t.pptx", **outline}))
    built = build_deck(deck, lambda name: template)
    assert not validate_package(built.data), validate_package(built.data)
    return Presentation(io.BytesIO(built.data)), built.slides


def _texts(slide: object) -> list[str]:
    return [
        s.text_frame.text
        for s in walk(slide.shapes)
        if getattr(s, "has_text_frame", False) and s.has_text_frame
    ]  # type: ignore[attr-defined]


def test_samples_are_recognised_by_what_they_hold() -> None:
    samples = analyze_samples(Presentation(io.BytesIO(make_infographic_template())))
    kinds = {s.number: s.kind for s in samples}
    assert kinds[1] == "title" and kinds[2] == "section" and kinds[4] == "chart" and kinds[5] == "chart"
    assert kinds[6] == "table" and kinds[10] == "thanks"
    assert (
        samples[3].chart_families == ["donut"]
        and samples[4].chart_families == ["column"]
        and samples[5].table_sizes == [(4, 3)]
    )


def test_the_dark_or_light_side_of_the_template_is_read_from_the_file() -> None:
    samples = analyze_samples(Presentation(io.BytesIO(make_infographic_template())))
    assert [s.dark for s in samples] == [True, True, True, False, True, False, False, True, False, True]


def test_a_deck_built_on_it_uses_the_samples_charts_and_table_and_leaves_no_sticky_note() -> None:
    prs, reports = _build(
        {
            "slides": [
                {"layout": "title", "title": "Quarterly review", "subtitle": "Finance"},
                {"layout": "section", "title": "Results"},
                {
                    "layout": "chart",
                    "title": "Share by region",
                    "kind": "donut",
                    "categories": ["North", "South", "East", "West"],
                    "series": [{"name": "Share", "values": [40, 30, 20, 10]}],
                },
                {
                    "layout": "chart",
                    "title": "Sales by quarter",
                    "kind": "column",
                    "categories": ["Q1", "Q2", "Q3"],
                    "series": [
                        {"name": "2025", "values": [1, 2, 3]},
                        {"name": "2026", "values": [2, 3, 5]},
                        {"name": "Plan", "values": [2, 3, 4]},
                    ],
                },
                {
                    "layout": "table",
                    "title": "Numbers by region",
                    "columns": ["Region", "Revenue", "Margin", "Notes"],
                    "rows": [
                        ["North", 120, "12%", "ok"],
                        ["South", 95, "9%", "low"],
                        ["East", 70, "11%", ""],
                        ["West", 60, "8%", "new"],
                        ["Online", 40, "15%", "up"],
                    ],
                },
                {"layout": "closing", "title": "Thank you"},
            ]
        }
    )
    assert prs.slide_width == (10 * 914400) and len(prs.slides) == 6
    labels = [r.layout for r in reports]
    assert any("template sample 8" in label for label in labels) and any(
        "template sample 5" in label for label in labels
    )
    assert any("template sample 6" in label for label in labels)
    donut = next(
        s for s in walk(prs.slides[2].shapes) if getattr(s, "has_chart", False) and s.has_chart
    ).chart
    assert donut.chart_type.name == "DOUGHNUT" and list(donut.plots[0].categories) == [
        "North",
        "South",
        "East",
        "West",
    ]
    assert list(donut.plots[0].series[0].values) == [40, 30, 20, 10]
    column = next(
        s for s in walk(prs.slides[3].shapes) if getattr(s, "has_chart", False) and s.has_chart
    ).chart
    assert [s.name for s in column.plots[0].series] == ["2025", "2026", "Plan"]
    table = next(s for s in walk(prs.slides[4].shapes) if getattr(s, "has_table", False) and s.has_table)
    assert (len(table.table.rows), len(table.table.columns)) == (6, 4)
    assert (
        table.table.cell(0, 3).text == "Notes"
        and table.table.cell(5, 0).text == "Online"
        and table.table.cell(1, 1).text == "120"
    )
    assert table.width == int(8.4 * 914400)
    for slide in prs.slides:  # nothing of the sticky note, everything of the company's frame
        assert not any("Delete before" in t or "Changing image" in t for t in _texts(slide))
        assert all(shape.name != "Picture 12" for shape in slide.shapes)
    assert all(any(shape.name == "Company logo" for shape in slide.shapes) for slide in list(prs.slides)[:5])
    assert all("Acme confidential" in " ".join(_texts(slide)) for slide in list(prs.slides)[:5])


def test_a_chart_kind_the_template_has_no_sample_for_is_drawn_in_the_template_theme() -> None:
    prs, reports = _build(
        {
            "slides": [
                {"layout": "title", "title": "T"},
                {
                    "layout": "chart",
                    "title": "Trend",
                    "kind": "line",
                    "categories": ["a", "b"],
                    "series": [{"name": "s", "values": [1, 2]}],
                },
            ]
        }
    )
    chart = next(
        s for s in walk(prs.slides[1].shapes) if getattr(s, "has_chart", False) and s.has_chart
    ).chart
    assert "LINE" in chart.chart_type.name and not any(
        "template sample" in r.layout and "chart" in r.layout for r in reports
    )


def test_a_table_with_too_many_rows_falls_back_to_the_quiet_native_table() -> None:
    rows = [[f"R{i}", i] for i in range(20)]
    _, reports = _build(
        {
            "slides": [
                {"layout": "title", "title": "T"},
                {"layout": "table", "title": "Long table", "columns": ["Name", "Value"], "rows": rows},
            ]
        }
    )
    assert not any("table (template sample" in r.layout for r in reports)


def _chart_label_numbers(reports: list) -> list[str]:  # type: ignore[type-arg]
    return [r.layout for r in reports if "chart (template sample" in r.layout]


def test_the_deck_stays_on_the_side_the_first_slide_chose() -> None:
    """The first slide is the dark cover, so the column chart comes from the dark sample (5) and the donut from the dark twin (8)."""
    outline = {
        "slides": [
            {"layout": "title", "title": "T"},
            {
                "layout": "chart",
                "title": "Donut",
                "kind": "donut",
                "categories": ["a", "b"],
                "series": [{"name": "s", "values": [1, 2]}],
            },
            {
                "layout": "chart",
                "title": "Columns",
                "kind": "column",
                "categories": ["a", "b"],
                "series": [{"name": "s", "values": [1, 2]}],
            },
        ]
    }
    _, reports = _build(outline)
    assert any("template sample 8" in label for label in _chart_label_numbers(reports))
    assert any("template sample 5" in label for label in _chart_label_numbers(reports))


def test_the_mood_can_be_asked_for() -> None:
    outline = {
        "mood": "light",
        "slides": [
            {"layout": "title", "title": "T"},
            {
                "layout": "chart",
                "title": "Donut",
                "kind": "donut",
                "categories": ["a", "b"],
                "series": [{"name": "s", "values": [1, 2]}],
            },
            {
                "layout": "chart",
                "title": "Columns",
                "kind": "column",
                "categories": ["a", "b"],
                "series": [{"name": "s", "values": [1, 2]}],
            },
        ],
    }
    _, reports = _build(outline)
    assert any("template sample 4" in label for label in _chart_label_numbers(reports))
    assert any("template sample 9" in label for label in _chart_label_numbers(reports))


def test_text_longer_than_the_samples_is_made_smaller_to_fit_and_a_huge_text_is_reported() -> None:
    """Probe on the real template: the long-text slides overflowed the sample's boxes (D-332)."""
    samples = analyze_samples(Presentation(io.BytesIO(make_infographic_template())))
    caption = next(s for s in samples[2].slots if s.role == "text" and len(s.text) > 30)
    long = (
        "A caption that says a good deal more than the sample's own few words do, in order to need a smaller size. "
        * 2
    )
    huge = "word " * 400
    for text, warned in ((long, False), (huge, True)):
        prs, reports = _build(
            {
                "slides": [
                    {"layout": "title", "title": "T"},
                    {
                        "layout": "template_slide",
                        "source": 3,
                        "fill": {"title": "Facts", f"slot:{caption.id}": text},
                    },
                ]
            }
        )
        shape = next(
            s
            for s in walk(prs.slides[1].shapes)
            if getattr(s, "has_text_frame", False) and s.text_frame.text.startswith(text[:20])
        )
        size = shape.text_frame.paragraphs[0].runs[0].font.size.pt
        assert size < 14 and size >= 8.4
        assert any("too long for its box" in w for w in reports[1].warnings) is warned


def test_a_picture_goes_into_the_samples_round_frame_covering_it() -> None:
    """A portrait or photo frame of the real template is a picture cropped to a circle (D-335)."""
    from PIL import Image
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation(io.BytesIO(make_infographic_template()))
    slide = prs.slides.add_slide(next(layout for layout in prs.slide_layouts if layout.name == "Title Only"))
    slide.shapes.title.text_frame.text = "Meet the speaker"
    old = io.BytesIO()
    Image.new("RGB", (40, 40), (200, 0, 0)).save(old, "PNG")
    frame = slide.shapes.add_picture(old, Inches(1), Inches(1.5), Inches(2.5), Inches(2.5))
    frame.auto_shape_type = MSO_SHAPE.OVAL
    out = io.BytesIO()
    prs.save(out)
    template = out.getvalue()
    new = io.BytesIO()
    Image.new("RGB", (600, 300), (0, 0, 200)).save(new, "PNG")
    deck = parse_deck(
        json.dumps(
            {
                "title": "T",
                "template": "t.pptx",
                "slides": [
                    {"layout": "title", "title": "T"},
                    {
                        "layout": "template_slide",
                        "source": 11,
                        "fill": {"title": "Ada", "picture": "ada.png"},
                    },
                ],
            }
        )
    )
    built = build_deck(deck, lambda name: new.getvalue() if name == "ada.png" else template)
    result = Presentation(io.BytesIO(built.data))
    picture = next(s for s in walk(result.slides[1].shapes) if s.shape_type == 13)
    assert picture.auto_shape_type == MSO_SHAPE.OVAL and picture.image.size == (600, 300)
    assert abs(picture.crop_left - 0.25) < 1e-6 and abs(picture.crop_right - 0.25) < 1e-6
    assert not validate_package(built.data)
