"""Input files for the deck benchmark tasks, made on demand (so no binaries are committed). `make_seed(task_id, destination)`."""

from __future__ import annotations

import csv
import io
import json
import random
from pathlib import Path

POLICY = """# Remote Work and Equipment Policy (version 3.2)

## 1. Purpose and scope
This policy sets out the Company's requirements regarding remote working arrangements ("Arrangements") for all employees
("Employees") of the Company and its subsidiaries, notwithstanding any contrary provisions in individual contracts, save where
mandatory local law provides otherwise.

## 2. Eligibility
Employees who have completed their probationary period of three (3) months and whose role, as determined by their line manager
in consultation with Human Resources, can be performed remotely, shall be eligible. Roles requiring physical presence (warehouse,
reception, laboratory) are excluded.

## 3. Working pattern
Employees on a hybrid Arrangement must attend the office on at least two (2) days per week, which shall be agreed with the team
lead. Core hours are 10:00 to 15:00 local time, during which Employees shall be reachable. Fully remote working requires written
approval from a Director and is reviewed every six (6) months.

## 4. Equipment
The Company shall provide a laptop, a monitor and a headset. A one-off home office allowance of 400 EUR may be claimed within the
first 12 months through the expenses system with receipts. Equipment remains Company property and must be returned within 14 days
of termination of employment.

## 5. Security
Employees must use the Company VPN on any network that is not the office network, must not use personal devices to store customer
data, must lock their screen when away, and must report a lost or stolen device to the IT service desk within 1 hour.

## 6. Working abroad
Working from another country for more than 10 working days per calendar year requires prior approval from HR because of tax and
social security consequences. Working abroad is never permitted from countries on the Company's restricted list (Appendix A).

## 7. Health and safety
Employees are responsible for maintaining a safe workspace. The Company will reimburse an ergonomic assessment once per year.

## 8. Breach
Failure to comply may lead to withdrawal of the Arrangement and, in serious cases, disciplinary action.
"""


def _sales(dest: Path) -> None:
    random.seed(7)
    regions = {"North": 1.0, "South": 0.8, "East": 1.3, "West": 0.6}
    products = {"Router": 120, "Switch": 90, "Firewall": 150}
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["month", "region", "product", "units", "revenue_eur"])
    for year in (2024, 2025):
        for month in range(1, 13):
            for region, factor in regions.items():
                for product, price in products.items():
                    growth = 1 + (0.02 * (month + 12 * (year - 2024)))
                    if region == "East" and year == 2025:
                        growth *= 1.35  # the story: East took off in 2025
                    if region == "West" and year == 2025:
                        growth *= 0.7  # and West fell
                    units = int(60 * factor * growth * random.uniform(0.9, 1.1))
                    if product == "Firewall" and month in (11, 12) and year == 2025:
                        units *= 2
                    writer.writerow([f"{year}-{month:02d}", region, product, units, units * price])
    text = buffer.getvalue().splitlines()
    text[40] = (
        text[40].rsplit(",", 2)[0] + ",,"
    )  # a row with missing units and revenue: the data quality issue
    text.append("2025-12,south,Router,55,6600")  # a region spelt in lower case
    (dest / "sales.csv").write_text("\n".join(text) + "\n", encoding="utf-8")


def _template(dest: Path) -> None:
    from tests.slide_templates import make_template

    (dest / "company-template.pptx").write_bytes(make_template())


def _board_update(dest: Path) -> None:
    from forge.slides.build import build_deck
    from forge.slides.spec import parse_deck

    outline = {
        "title": "Board update",
        "slides": [
            {"layout": "title", "title": "Board update FY24", "subtitle": "Finance and operations"},
            {
                "layout": "bullets",
                "title": "FY24 in numbers",
                "bullets": ["Revenue up 12%", "Margin 6.1%", "Headcount 410"],
                "notes": "Open with the headline.",
            },
            {
                "layout": "chart",
                "title": "FY24 revenue by quarter (millions)",
                "kind": "column",
                "categories": ["Q1", "Q2", "Q3", "Q4"],
                "series": [{"name": "Revenue", "values": [3.8, 4.1, 4.5, 4.9]}],
            },
            {
                "layout": "bullets",
                "title": "FY24 priorities delivered",
                "bullets": ["New warehouse opened", "ERP migration done"],
            },
            {"layout": "closing", "title": "Thank you"},
        ],
    }
    (dest / "board_update.pptx").write_bytes(
        build_deck(parse_deck(json.dumps(outline)), lambda name: b"").data
    )


UGLY_CONTENT = [
    (
        "Customer onboarding",
        [
            "Our onboarding process has historically consisted of a number of manual steps that new customers must complete before "
            "they can start using the product, including account creation, data import, user invitations and a kick-off call",
            "In the last year we have been working to automate as many of these steps as we possibly can so that the average time "
            "to first value falls below five days",
            "Results so far: median time to first value fell from 19 days to 8 days; 64 percent of customers now import their data "
            "without help; kick-off calls dropped by half",
        ],
    ),
    (
        "What customers say",
        [
            "Customers told us in 212 survey responses that the data import was the most confusing step, followed by inviting "
            "colleagues, and that they would like a checklist inside the product that shows what is left to do",
            "Satisfaction with onboarding rose from 3.4 to 4.1 out of 5, driven mainly by the new import wizard",
        ],
    ),
    (
        "Next steps",
        [
            "Ship the in-product checklist in Q2, add sample data for empty accounts in Q2, automate user invitations from the "
            "directory in Q3, and measure time to first value weekly on the executive dashboard so that we can see the effect of "
            "each change as it lands",
        ],
    ),
]


def _ugly(dest: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    colours = [("FFFF00", "FF00FF"), ("00FFFF", "FF0000"), ("00FF00", "FF8800")]
    for index, (title, paragraphs) in enumerate(UGLY_CONTENT):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        back, front = colours[index % 3]
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(back)
        title_box = slide.shapes.add_textbox(Inches(0.2), Inches(0.1), Inches(12.9), Inches(0.9))
        title_box.text_frame.text = title.upper()
        title_run = title_box.text_frame.paragraphs[0].runs[0]
        title_run.font.size, title_run.font.name = Pt(40), "Comic Sans MS"
        title_run.font.color.rgb = RGBColor.from_string(front)
        body = slide.shapes.add_textbox(Inches(0.2), Inches(1.2), Inches(12.9), Inches(6))
        body.text_frame.word_wrap = True
        for number, paragraph in enumerate(paragraphs):
            target = body.text_frame.paragraphs[0] if number == 0 else body.text_frame.add_paragraph()
            run = target.add_run()
            run.text = paragraph
            run.font.size, run.font.name = Pt(13), "Papyrus"
            run.font.color.rgb = RGBColor.from_string(front)
    prs.save(str(dest / "ugly.pptx"))


def make_seed(task_id: str, dest: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    if task_id.startswith("d04"):
        _sales(dest)
    elif task_id.startswith("d12"):
        (dest / "policy.md").write_text(POLICY, encoding="utf-8")
    elif task_id.startswith("d06"):
        _template(dest)
    elif task_id.startswith("d07"):
        _board_update(dest)
    elif task_id.startswith("d08"):
        _ugly(dest)
