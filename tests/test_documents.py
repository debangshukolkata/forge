"""D-342: Word documents are built from an outline, read, edited and checked (W1)."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from PIL import Image  # noqa: E402

from forge.documents.build import build_document  # noqa: E402
from forge.documents.check import check_document, describe_problems  # noqa: E402
from forge.documents.edit import Operation, edit_document, replace_in_paragraph  # noqa: E402
from forge.documents.read import read_document  # noqa: E402
from forge.documents.spec import DocumentError, parse_document  # noqa: E402


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


OUTLINE = {
    "title": "Quarterly report",
    "author": "Finance",
    "toc": True,
    "header": "Confidential",
    "footer": "Acme",
    "blocks": [
        {"type": "heading", "level": 1, "text": "Summary"},
        {
            "type": "paragraph",
            "runs": [
                {"text": "Sales rose "},
                {"text": "12%", "bold": True},
                {"text": ". See "},
                {"text": "the site", "link": "https://example.com"},
            ],
        },
        {"type": "bullets", "items": ["one", "two"]},
        {"type": "numbered", "items": ["first", "second"]},
        {"type": "heading", "level": 2, "text": "Detail"},
        {
            "type": "table",
            "header": ["Region", "Q1"],
            "rows": [["North", "100"], ["South", "80"]],
            "caption": "Table 1",
        },
        {"type": "image", "path": "p.png", "alt": "A red box", "caption": "Figure 1"},
        {"type": "quote", "text": "Good work."},
    ],
}


def _built(outline: dict | None = None) -> bytes:
    return build_document(parse_document(json.dumps(outline or OUTLINE)), lambda name: _png())


def test_build_has_structure_fields_links_alt_text_and_header_row() -> None:
    data = _built()
    info = read_document(data)
    assert info.title == "Quarterly report" and info.language == "en-GB"
    assert "# Quarterly report" in info.markdown and "## Detail" in info.markdown
    assert "- one" in info.markdown and "1. first" in info.markdown and "| North | 100 |" in info.markdown
    assert info.pictures == 1 and info.pictures_without_alt == 0
    assert describe_problems(check_document(data)).startswith("No structure")
    built = docx.Document(io.BytesIO(data))
    xml = built.element.xml
    assert "TOC" in xml and "w:hyperlink" in xml and "w:tblHeader" in xml
    assert "PAGE" in built.sections[0].footer._element.xml


def test_bad_links_and_outlines_are_clear_errors() -> None:
    with pytest.raises(DocumentError):
        parse_document("{no")
    bad = {
        "title": "x",
        "blocks": [{"type": "paragraph", "runs": [{"text": "a", "link": "javascript:alert(1)"}]}],
    }
    with pytest.raises(DocumentError, match="only https"):
        _built(bad)
    with pytest.raises(DocumentError):
        parse_document(
            json.dumps({"title": "x", "blocks": [{"type": "image", "path": "p.png"}]})
        )  # alt text is required


def test_check_finds_skipped_heading_levels_faked_headings_and_missing_title() -> None:
    document = docx.Document()
    document.add_heading("Top", level=1)
    document.add_heading("Deep", level=3)
    fake = document.add_paragraph()
    fake.add_run("Looks like a heading").bold = True
    out = io.BytesIO()
    document.save(out)
    text = describe_problems(check_document(out.getvalue()))
    assert "jumps from level 1 to 3" in text and "bold text" in text and "no title" in text


def test_replace_text_works_across_runs_and_keeps_the_formatting() -> None:
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Hello Wo")
    paragraph.add_run("rld and Wor").bold = True
    paragraph.add_run("ld")
    assert replace_in_paragraph(paragraph, "World", "There") == 2
    assert paragraph.text == "Hello There and There"

    data = _built()
    edited, _ = edit_document(
        data,
        [
            Operation(op="replace_text", find="Sales rose", replace="Revenue grew"),
            Operation(op="append_paragraph", text="Appendix note"),
            Operation(op="delete_paragraph", index=3),
        ],
    )
    markdown = read_document(edited).markdown
    assert (
        "Revenue grew" in markdown
        and "Appendix note" in markdown
        and "- one" not in markdown.split("\n")[2:3]
    )
    with pytest.raises(DocumentError, match="not in the document"):
        edit_document(data, [Operation(op="replace_text", find="zzzzqq", replace="x")])
    with pytest.raises(DocumentError, match="index must be"):
        edit_document(data, [Operation(op="delete_paragraph", index=999)])


def test_a_template_keeps_its_header_and_drops_its_sample_text() -> None:
    base = docx.Document()
    base.add_paragraph("SAMPLE TEXT FROM THE TEMPLATE")
    base.sections[0].header.paragraphs[0].text = "Company header"
    out = io.BytesIO()
    base.save(out)
    data = build_document(
        parse_document(json.dumps(OUTLINE | {"header": ""})), lambda n: _png(), out.getvalue()
    )
    built = docx.Document(io.BytesIO(data))
    assert "SAMPLE TEXT" not in read_document(data).markdown
    assert built.sections[0].header.paragraphs[0].text == "Company header"


async def test_the_tools(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.document import BuildDocument, CheckDocument, EditDocument, ReadDocument
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    workspace.write_text("docs/outline.json", json.dumps(OUTLINE), "x")
    workspace.write_bytes("p.png", _png(), "x")
    context = ToolContext(workspace=workspace)
    built = await BuildDocument().run(
        BuildDocument.Args(outline="docs/outline.json", output="docs/r.docx"), context
    )
    assert built.ok and "Wrote docs/r.docx" in built.content
    read = await ReadDocument().run(ReadDocument.Args(path="docs/r.docx"), context)
    assert read.ok and "Summary" in read.content
    edited = await EditDocument().run(
        EditDocument.Args(
            path="docs/r.docx", operations=[Operation(op="replace_text", find="Detail", replace="Details")]
        ),
        context,
    )
    assert edited.ok
    assert (await CheckDocument().run(CheckDocument.Args(path="docs/r.docx"), context)).ok
    assert not (await ReadDocument().run(ReadDocument.Args(path="docs/none.docx"), context)).ok
    from forge.documents.preview import libreoffice, word_available
    from forge.tools.document_preview import PreviewFile

    if word_available() or libreoffice():
        shown = await PreviewFile().run(
            PreviewFile.Args(path="docs/r.docx", pdf_output="docs/r.pdf"), context
        )
        assert shown.ok and "page(s)" in shown.content and workspace.path_of("docs/r.pdf").is_file()


TRACKED = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:t xml:space="preserve">Keep </w:t></w:r>'
    '<w:ins w:id="1" w:author="A"><w:r><w:t>added </w:t></w:r></w:ins>'
    '<w:del w:id="2" w:author="A"><w:r><w:delText>removed </w:delText></w:r></w:del><w:r><w:t>end</w:t></w:r></w:p>'
)


def _with_tracked_changes() -> bytes:
    from docx.oxml import parse_xml

    document = docx.Document()
    document.add_paragraph("placeholder")
    document.element.body.insert(0, parse_xml(TRACKED))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_tracked_changes_are_counted_accepted_and_rejected() -> None:
    data = _with_tracked_changes()
    info = read_document(data)
    assert info.insertions == 1 and info.deletions == 1
    accepted, done = edit_document(data, [Operation(op="accept_changes")])
    text = read_document(accepted).markdown
    assert text.startswith("Keep added end") and "removed" not in text and "accepted 2" in done[0]
    rejected, _ = edit_document(data, [Operation(op="reject_changes")])
    text = read_document(rejected).markdown
    assert text.startswith("Keep removed end") and "added" not in text
    after = read_document(accepted)
    assert after.insertions == 0 and after.deletions == 0


def test_footnotes_are_built_read_and_added_to_an_existing_document() -> None:
    outline = {
        "title": "Notes",
        "blocks": [
            {
                "type": "paragraph",
                "runs": [
                    {"text": "Sales rose 12%", "footnote": "Source: finance report, page 4."},
                    {"text": " this year."},
                ],
            },
            {"type": "paragraph", "text": "Second paragraph."},
        ],
    }
    data = build_document(parse_document(json.dumps(outline)), lambda n: b"")
    assert read_document(data).footnotes == ["Source: finance report, page 4."]
    edited, done = edit_document(data, [Operation(op="add_footnote", index=3, text="A second note.")])
    assert read_document(edited).footnotes == ["Source: finance report, page 4.", "A second note."]
    assert "footnote added" in done[0]
    assert "w:footnoteReference" in docx.Document(io.BytesIO(edited)).element.xml
    with pytest.raises(DocumentError, match="needs text"):
        edit_document(data, [Operation(op="add_footnote", index=1, text=" ")])


def test_word_exports_the_document_to_pdf_when_installed() -> None:
    from forge.documents.preview import libreoffice, to_pdf, word_available

    if not (word_available() or libreoffice()):
        pytest.skip("neither Word nor LibreOffice here")
    assert to_pdf(_built()).startswith(b"%PDF")


def _opened(outline: dict) -> object:
    return docx.Document(io.BytesIO(build_document(parse_document(json.dumps(outline)), lambda path: b"")))


def test_a_letter_has_no_title_tight_address_blocks_and_room_to_sign() -> None:
    """W8, doc-2 w05: the letter had a large title and a gap after every address line."""
    lines = ["Priya Raman", "14 Elm Court", "12 November 2026", "Harbour Music School", "8 Harbour Street"]
    blocks = [{"type": "paragraph", "text": line} for line in lines]
    blocks += [
        {"type": "paragraph", "text": t} for t in ("Dear Sir,", "A sentence.", "Yours sincerely,", "Priya")
    ]
    opened = _opened({"title": "Letter", "kind": "letter", "blocks": blocks})
    paragraphs = opened.paragraphs
    assert all(p.style.name != "Title" for p in paragraphs) and opened.core_properties.title == "Letter"
    assert paragraphs[0].paragraph_format.space_after == 0 and paragraphs[3].paragraph_format.space_after == 0
    assert paragraphs[2].paragraph_format.space_before.pt == 12  # the date stands apart
    assert paragraphs[-1].paragraph_format.space_before.pt == 36
    assert "Page" not in opened.sections[0].footer.paragraphs[0].text


def test_a_newline_is_a_line_break_and_a_callout_shows_the_figure() -> None:
    opened = _opened(
        {
            "title": "Proposal",
            "blocks": [
                {"type": "paragraph", "text": "8 Harbour Street\nMillbridge"},
                {"type": "callout", "label": "Total price", "value": "GBP 8,400", "note": "excluding VAT"},
            ],
        }
    )
    address = next(p for p in opened.paragraphs if "Harbour" in p.text)
    assert "<w:br/>" in address._p.xml and address.text == "8 Harbour Street\nMillbridge"
    box = opened.tables[0].rows[0].cells[0]
    assert [p.text for p in box.paragraphs] == ["Total price", "GBP 8,400", "excluding VAT"]
    assert box.paragraphs[1].runs[0].bold and box.paragraphs[1].runs[0].font.size.pt == 20


def test_tables_get_a_shaded_header_right_aligned_numbers_and_a_bold_total() -> None:
    """W8, doc-2 w04: a plain grid with the total price no different from any other row."""
    table = {
        "type": "table",
        "header": ["Phase", "Description", "Price (GBP)"],
        "rows": [["Design", "Page layouts and a visual design", "2,800"], ["Total", "All phases", "8,400"]],
    }
    opened = _opened({"title": "Proposal", "blocks": [table]})
    made = opened.tables[0]
    assert 'w:fill="1F3864"' in made.rows[0].cells[0]._tc.xml
    assert (
        made.rows[1].cells[2].paragraphs[0].alignment == 2
        and made.rows[1].cells[0].paragraphs[0].alignment is None
    )
    assert made.rows[2].cells[2].paragraphs[0].runs[0].bold
    assert made.rows[1].cells[1].width > made.rows[1].cells[2].width


def test_paragraphs_around_tables_and_a_bold_subject_line_are_not_problems() -> None:
    """D-532, doc-3: a table, a callout and one empty paragraph were counted as a run of spacing, and a letter's bold subject as a fake heading."""
    outline = {
        "title": "Proposal",
        "blocks": [
            {"type": "table", "header": ["Phase", "Price"], "rows": [["Build", "3,500"]]},
            {"type": "callout", "label": "Total", "value": "GBP 3,500"},
            {"type": "paragraph", "text": ""},
            {"type": "paragraph", "runs": [{"text": "Subject: a place for Anika", "bold": True}]},
        ],
    }
    data = build_document(parse_document(json.dumps(outline)), lambda path: b"")
    assert describe_problems(check_document(data)).startswith("No structure")
    spacing = {"title": "T", "blocks": [{"type": "paragraph", "text": ""}] * 3}
    assert "empty paragraphs" in describe_problems(
        check_document(build_document(parse_document(json.dumps(spacing)), lambda path: b""))
    )


def test_an_unlabelled_bold_subject_above_the_salutation_is_not_a_fake_heading() -> None:
    """D-541, doc-4 w05: a British letter puts its bold subject line just above "Dear ..." without "Re:"."""

    def problems(blocks: list[dict[str, object]]) -> str:
        outline = {"title": "Letter", "kind": "letter", "blocks": blocks}
        return describe_problems(
            check_document(build_document(parse_document(json.dumps(outline)), lambda path: b""))
        )

    subject = {"type": "paragraph", "runs": [{"text": "Place for my daughter Anika", "bold": True}]}
    assert problems([subject, {"type": "paragraph", "text": "Dear Admissions Office,"}]).startswith(
        "No structure"
    )
    assert "looks like a heading" in problems(
        [subject, {"type": "paragraph", "text": "Our school opened in 1990."}]
    )
    assert "looks like a heading" in problems(
        [{"type": "paragraph", "text": "Intro."}, subject]
    )  # the last paragraph
    assert (
        problems([subject, subject, {"type": "paragraph", "text": "Dear Sir,"}]).count("looks like a heading")
        == 1
    )
