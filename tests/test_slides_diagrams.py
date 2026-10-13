"""P2.6: diagram recipes, connectors glued to shapes, groups, align and distribute."""

from __future__ import annotations

import io
import json
import re
import zipfile

import pytest
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.errors import BuildError
from forge.slides.layout import LayoutError, expand_elements
from forge.slides.spec import SpecError, parse_deck
from forge.slides.validate import validate_package

BOX = {"x": 0.8, "y": 2.0, "w": 11.7, "h": 4.6}


def _build(elements: list[dict]) -> tuple[bytes, list[str]]:  # type: ignore[type-arg]
    deck = {
        "title": "Deck",
        "slides": [{"layout": "freeform", "title": "Diagrams state the idea quickly", "elements": elements}],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda name: b"")
    return built.data, [w for s in built.slides for w in s.warnings]


def _xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("ppt/slides/slide1.xml").decode("utf-8")


RECIPES = {
    "process": ["Plan", "Build", "Test", "Ship"],
    "cycle": ["Measure", "Analyse", "Improve", "Repeat"],
    "hierarchy": [
        {"title": "CEO", "children": [{"title": "Eng"}, {"title": "Sales", "children": [{"title": "EMEA"}]}]}
    ],
    "matrix": ["Quick wins", "Big bets", "Fill-ins", "Money pits"],
    "timeline": [
        {"date": "Jan", "title": "Kick-off"},
        {"date": "Jun", "title": "Pilot"},
        {"date": "Dec", "title": "Launch"},
    ],
    "funnel": [
        {"title": "Visitors", "text": "10,000"},
        {"title": "Sign-ups", "text": "2,400"},
        {"title": "Customers", "text": "310"},
    ],
    "roadmap": [
        {"title": "Now", "items": ["A", "B"]},
        {"title": "Next", "items": ["C"]},
        {"title": "Later", "items": ["D"]},
    ],
}


@pytest.mark.parametrize("recipe", list(RECIPES))
def test_every_recipe_builds_a_valid_slide_without_warnings(recipe: str) -> None:
    data, warnings = _build([{"type": "diagram", "recipe": recipe, **BOX, "items": RECIPES[recipe]}])
    assert warnings == [] and validate_package(data) == []
    assert len(list(Presentation(io.BytesIO(data)).slides[0].shapes)) >= 4


def test_process_connectors_are_glued_to_the_boxes() -> None:
    data, _ = _build([{"type": "diagram", "recipe": "process", **BOX, "items": ["Plan", "Build", "Ship"]}])
    xml = _xml(data)
    assert len(re.findall(r"<a:stCxn ", xml)) == 2 and len(re.findall(r"<a:endCxn ", xml)) == 2
    assert 'idx="3"' in xml and 'idx="1"' in xml  # right side of one box to the left side of the next


def test_cycle_closes_the_loop_and_can_have_a_centre() -> None:
    data, warnings = _build(
        [{"type": "diagram", "recipe": "cycle", **BOX, "items": RECIPES["cycle"], "center": "Learn"}]
    )
    assert len(re.findall(r"<a:stCxn ", _xml(data))) == 4 and "Learn" in _xml(data) and warnings == []


def test_hierarchy_has_one_box_per_node_and_needs_a_single_root() -> None:
    elements = expand_elements(
        [{"type": "diagram", "recipe": "hierarchy", **BOX, "items": RECIPES["hierarchy"]}], 1
    )
    assert len([e for e in elements if e["type"] == "text"]) == 4
    with pytest.raises(LayoutError, match="one root item"):
        expand_elements([{"type": "diagram", "recipe": "hierarchy", **BOX, "items": ["a", "b"]}], 1)


def test_diagram_mistakes_are_reported_in_words() -> None:
    with pytest.raises(LayoutError, match="unknown recipe"):
        expand_elements([{"type": "diagram", "recipe": "spiral", **BOX, "items": ["a"]}], 1)
    with pytest.raises(LayoutError, match="needs x, y, w and h"):
        expand_elements([{"type": "diagram", "recipe": "process", "items": ["a", "b"]}], 1)
    with pytest.raises(LayoutError, match="needs 2 to 6 items"):
        expand_elements([{"type": "diagram", "recipe": "process", **BOX, "items": ["a"]}], 1)
    with pytest.raises(LayoutError, match="exactly 4 items"):
        expand_elements([{"type": "diagram", "recipe": "matrix", **BOX, "items": ["a", "b"]}], 1)
    with pytest.raises(LayoutError, match="needs a title"):
        expand_elements([{"type": "diagram", "recipe": "process", **BOX, "items": [{"text": "x"}, "b"]}], 1)


def test_a_line_can_be_glued_between_any_two_shapes_with_ids() -> None:
    elements = [
        {"type": "shape", "id": "a", "x": 1, "y": 2, "w": 2, "h": 1, "text": "A"},
        {"type": "shape", "id": "b", "shape": "oval", "x": 6, "y": 3, "w": 2, "h": 1.5, "text": "B"},
        {
            "type": "line",
            "start": {"id": "a", "side": "right"},
            "end": {"id": "b", "side": "left"},
            "kind": "elbow",
            "arrow": True,
        },
    ]
    data, warnings = _build(elements)
    xml = _xml(data)
    assert 'idx="3"' in xml and 'idx="2"' in xml and "bentConnector3" in xml and "tailEnd" in xml
    assert warnings == [] and validate_package(data) == []


def test_lines_need_known_ids_or_coordinates() -> None:
    with pytest.raises(BuildError, match="no element with id 'ghost'"):
        _build([{"type": "line", "start": {"id": "ghost"}, "end": {"id": "ghost"}}])
    with pytest.raises(SpecError, match="needs x1, y1, x2, y2"):
        _build([{"type": "line"}])
    data, _ = _build([{"type": "line", "x1": 1, "y1": 2, "x2": 5, "y2": 2}])
    assert validate_package(data) == []


def test_duplicate_ids_are_an_error() -> None:
    boxes = [
        {"type": "shape", "id": "a", "x": 1, "y": 2, "w": 1, "h": 1},
        {"type": "shape", "id": "a", "x": 3, "y": 2, "w": 1, "h": 1},
    ]
    with pytest.raises(BuildError, match="used twice"):
        _build(boxes)


def test_group_binds_shapes_together() -> None:
    elements = [
        {"type": "shape", "id": "a", "x": 1, "y": 2, "w": 1, "h": 1},
        {"type": "shape", "id": "b", "x": 3, "y": 2, "w": 1, "h": 1},
        {"type": "group", "ids": ["a", "b"], "name": "Pair"},
    ]
    data, warnings = _build(elements)
    group = next(s for s in Presentation(io.BytesIO(data)).slides[0].shapes if s.shape_type == 6)
    assert group.name == "Pair" and len(group.shapes) == 2 and warnings == [] and validate_package(data) == []
    with pytest.raises(BuildError, match="no element with id 'zzz'"):
        _build([{"type": "group", "ids": ["zzz", "yyy"]}])


def _cards() -> list[dict]:  # type: ignore[type-arg]
    positions = [(1, 2.0), (2.5, 2.6), (9, 2.2)]
    return [
        {"type": "shape", "id": f"c{i}", "x": x, "y": y, "w": 1, "h": 1}
        for i, (x, y) in enumerate(positions, 1)
    ]


def test_align_and_distribute_move_the_named_elements() -> None:
    steps = [
        {"type": "align", "ids": ["c1", "c2", "c3"], "to": "top"},
        {"type": "distribute", "ids": ["c1", "c2", "c3"], "axis": "horizontal"},
    ]
    shapes = [e for e in expand_elements([*_cards(), *steps], 1) if e["type"] == "shape"]
    assert [e["y"] for e in shapes] == [2.0, 2.0, 2.0]
    assert [e["x"] for e in shapes] == [1, 5, 9]  # equal gaps (3.0) between the three, outer two fixed
    right = expand_elements([*_cards(), {"type": "align", "ids": ["c1", "c3"], "to": "right"}], 1)
    assert next(e["x"] for e in right if e["type"] == "shape") == 9
    middle = expand_elements([*_cards(), {"type": "align", "ids": ["c1", "c2"], "to": "middle"}], 1)
    assert [e["y"] for e in middle if e["type"] == "shape"][:2] == [2.3, 2.3]


def test_align_and_distribute_errors() -> None:
    cards = [
        {"type": "shape", "id": "a", "x": 1, "y": 2, "w": 1, "h": 1},
        {"type": "shape", "id": "b", "x": 3, "y": 2, "w": 1, "h": 1},
    ]
    with pytest.raises(LayoutError, match="no placed element with id 'zzz'"):
        expand_elements([*cards, {"type": "align", "ids": ["a", "zzz"], "to": "left"}], 1)
    with pytest.raises(LayoutError, match="`to` must be"):
        expand_elements([*cards, {"type": "align", "ids": ["a", "b"], "to": "sideways"}], 1)
    with pytest.raises(LayoutError, match="at least three"):
        expand_elements([*cards, {"type": "distribute", "ids": ["a", "b"], "axis": "horizontal"}], 1)
    with pytest.raises(LayoutError, match="at least two"):
        expand_elements([*cards, {"type": "align", "ids": ["a"], "to": "left"}], 1)
