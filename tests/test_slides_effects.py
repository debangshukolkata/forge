"""P2.2: gradients, transparency, shadows, picture frames and background layers on freeform slides."""

from __future__ import annotations

import io
import json
import re
import zipfile

from PIL import Image
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.spec import parse_deck
from forge.slides.validate import validate_package


def _png(colour: tuple[int, int, int] = (40, 90, 160), size: tuple[int, int] = (800, 600)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, colour).save(out, "PNG")
    return out.getvalue()


def _build(elements: list[dict], **slide: object) -> tuple[bytes, list[str]]:  # type: ignore[type-arg]
    deck = {
        "title": "Deck",
        "slides": [
            {"layout": "freeform", "title": "Effects make one idea stand out", "elements": elements, **slide}
        ],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda name: _png())
    return built.data, [w for s in built.slides for w in s.warnings]


def _slide_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("ppt/slides/slide1.xml").decode("utf-8")


def test_a_gradient_card_has_a_gradient_fill_in_the_right_place_and_a_valid_package() -> None:
    card = {
        "type": "shape",
        "x": 1,
        "y": 2,
        "w": 4,
        "h": 2,
        "fill": "accent",
        "fill_to": "#112233",
        "gradient_angle": 0,
    }
    data, warnings = _build([card])
    xml = _slide_xml(data)
    assert "<a:gradFill" in xml and 'ang="0"' in xml and 'val="112233"' in xml
    assert xml.index("<a:gradFill") < xml.index("<a:ln")  # the schema wants the fill before the outline
    assert validate_package(data) == [] and warnings == []


def test_transparency_is_an_alpha_on_the_fill() -> None:
    data, _ = _build([{"type": "shape", "x": 1, "y": 2, "w": 3, "h": 2, "fill": "accent", "opacity": 0.4}])
    assert re.search(r'<a:alpha val="40000"', _slide_xml(data))
    assert validate_package(data) == []


def test_shadows_soft_and_hard() -> None:
    for kind, blur in (("soft", "228600"), ("hard", "0")):
        data, _ = _build(
            [
                {
                    "type": "text",
                    "x": 1,
                    "y": 2,
                    "w": 4,
                    "h": 1,
                    "text": "Card",
                    "fill": "surface",
                    "shadow": kind,
                }
            ]
        )
        xml = _slide_xml(data)
        assert f'blurRad="{blur}"' in xml and "<a:outerShdw" in xml
        assert validate_package(data) == []


def test_a_text_card_can_have_a_gradient_and_the_contrast_check_uses_the_worst_stop() -> None:
    card = {
        "type": "text",
        "x": 1,
        "y": 2,
        "w": 5,
        "h": 1.5,
        "text": "Readable",
        "fill": "#0F1B2D",
        "fill_to": "#E8EEF8",
        "color": "#FFFFFF",
    }
    _, warnings = _build([card])
    assert any("low contrast" in w for w in warnings)  # white text on the pale end
    safe = dict(card, fill_to="#27405F")
    assert not any("low contrast" in w for w in _build([safe])[1])


def test_a_transparent_fill_is_judged_as_it_looks_on_the_slide() -> None:
    card = {
        "type": "text",
        "x": 1,
        "y": 2,
        "w": 5,
        "h": 1.5,
        "text": "Faint",
        "fill": "#0F1B2D",
        "opacity": 0.1,
        "color": "#FFFFFF",
    }
    assert any("low contrast" in w for w in _build([card])[1])  # 10% navy on white is nearly white


def test_picture_frames_and_circle_warning() -> None:
    circle = {
        "type": "image",
        "x": 1,
        "y": 2,
        "w": 2,
        "h": 2,
        "path": "p.png",
        "alt": "A portrait",
        "frame": "circle",
        "fit": "cover",
        "shadow": "soft",
    }
    data, warnings = _build([circle])
    presentation = Presentation(io.BytesIO(data))
    picture = next(s for s in presentation.slides[0].shapes if s.shape_type == 13)
    assert 'prst="ellipse"' in _slide_xml(data) and picture is not None and warnings == []
    rounded = dict(circle, frame="rounded")
    assert 'prst="roundRect"' in _slide_xml(_build([rounded])[0])
    oval = dict(circle, h=3)
    assert any("circle frame needs a square box" in w for w in _build([oval])[1])
    assert validate_package(data) == []


def test_a_gradient_background_sits_behind_everything() -> None:
    text = {"type": "text", "x": 1, "y": 2, "w": 6, "h": 1, "text": "Dark to light", "color": "#FFFFFF"}
    data, warnings = _build(
        [text],
        background="#0F1B2D",
        background_to="#27405F",
        background_angle=45,
        title_box={"color": "#FFFFFF"},
    )
    presentation = Presentation(io.BytesIO(data))
    shapes = list(presentation.slides[0].shapes)
    assert shapes[0].width == presentation.slide_width  # the full-bleed gradient layer is first (at the back)
    assert warnings == [] and validate_package(data) == []
    pale = _build([text], background="#0F1B2D", background_to="#F2F5FA")[1]
    assert any("low contrast" in w for w in pale)


def test_a_background_picture_with_a_scrim_keeps_text_readable_and_without_one_warns() -> None:
    text = {"type": "text", "x": 1, "y": 2, "w": 6, "h": 1, "text": "Over a picture", "color": "#FFFFFF"}
    data, warnings = _build(
        [text], background_image="bg.png", scrim=0.6, background="#101820", title_box={"color": "#FFFFFF"}
    )
    shapes = list(Presentation(io.BytesIO(data)).slides[0].shapes)
    assert (
        shapes[0].shape_type == 13 and shapes[1].width == shapes[0].width
    )  # picture at the back, the scrim above it
    assert warnings == [] and validate_package(data) == []
    assert any("without a scrim" in w for w in _build([text], background_image="bg.png")[1])
