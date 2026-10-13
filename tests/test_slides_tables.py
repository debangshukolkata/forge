"""P2.5: merged cells, widths, alignment, number formats, borders, heat maps and splitting long tables."""

from __future__ import annotations

import io
import json
import re
import zipfile

import pytest
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.errors import BuildError
from forge.slides.spec import parse_deck
from forge.slides.tables import format_number, lay_out_rows
from forge.slides.validate import validate_package

COLUMNS = ["Region", "Revenue", "Growth"]
ROWS = [["North", 1200, 0.12], ["South", 950, -0.03], ["East", 1400, 0.2]]
BOX = {"x": 1, "y": 2.4, "w": 8, "h": 3.5}


def _table(**options: object) -> dict:  # type: ignore[type-arg]
    return {"type": "table", **BOX, "columns": COLUMNS, "rows": ROWS, **options}


def _build(element: dict) -> tuple[bytes, list[str]]:  # type: ignore[type-arg]
    deck = {
        "title": "Deck",
        "slides": [
            {"layout": "freeform", "title": "Tables put exact numbers in view", "elements": [element]}
        ],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda name: b"")
    return built.data, [w for s in built.slides for w in s.warnings]


def _slide_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("ppt/slides/slide1.xml").decode("utf-8")


def _table_shape(data: bytes):  # type: ignore[no-untyped-def]
    return next(s for s in Presentation(io.BytesIO(data)).slides[0].shapes if s.has_table).table


@pytest.mark.parametrize(
    ("value", "pattern", "expected"),
    [
        (1234.5, "#,##0", "1,234"),
        (1234.5, "#,##0.00", "1,234.50"),
        (0.256, "0%", "26%"),
        (0.256, "0.0%", "25.6%"),
        (12, "$#,##0", "$12"),
        (3.14159, "0.00", "3.14"),
        (1500, '0.0"k"', "1500.0k"),
        (7, "", "7"),
        (2.5, "", "2.5"),
    ],
)
def test_number_formats(value: float, pattern: str, expected: str) -> None:
    assert format_number(value, pattern) == expected


def test_layout_counts_merged_cells_and_reports_short_or_long_rows() -> None:
    grid = lay_out_rows(
        [[{"text": "a", "rowspan": 2}, "b", "c"], ["d", "e"], [{"text": "wide", "colspan": 3}]], 3, "slide 1"
    )
    assert (
        grid[0][0]["rowspan"] == 2 and grid[1][0] is None and grid[1][1]["text"] == "d" and grid[2][1] is None
    )
    with pytest.raises(BuildError, match="needs 3 cells"):
        lay_out_rows([["a", "b"]], 3, "slide 1")
    with pytest.raises(BuildError, match="more cells"):
        lay_out_rows([["a", "b", "c", "d"]], 3, "slide 1")


def test_numbers_align_right_text_left_and_formats_apply() -> None:
    data, warnings = _build(_table(formats=["", "$#,##0", "0.0%"]))
    table = _table_shape(data)
    assert table.cell(1, 1).text == "$1,200" and table.cell(2, 2).text == "-3.0%"
    assert str(table.cell(1, 1).text_frame.paragraphs[0].alignment).startswith("RIGHT")
    assert str(table.cell(1, 0).text_frame.paragraphs[0].alignment).startswith("LEFT")
    assert warnings == [] and validate_package(data) == []


def test_explicit_alignment_widths_size_and_cell_styles() -> None:
    cells = [
        [
            {"text": "North", "bold": True, "fill": "#C62828"},
            1200,
            {"text": "up", "italic": True, "align": "center"},
        ]
    ]
    data, _ = _build(
        _table(
            rows=cells,
            align=["center", "left", "right"],
            col_widths=[1, 1, 2],
            size=18,
            first_column_bold=True,
        )
    )
    table = _table_shape(data)
    assert table.columns[2].width == pytest.approx(2 * table.columns[0].width, rel=0.02)
    assert str(table.cell(1, 0).fill.fore_color.rgb) == "C62828"
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold
    assert table.cell(1, 2).text_frame.paragraphs[0].runs[0].font.italic
    assert (
        str(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    )  # light text on a dark fill
    assert validate_package(data) == []


def test_merged_cells_span_rows_and_columns() -> None:
    rows = [[{"text": "Team", "rowspan": 2}, 1, 2], [3, 4], [{"text": "All together", "colspan": 3}]]
    data, warnings = _build(_table(rows=rows))
    table = _table_shape(data)
    assert table.cell(1, 0).is_merge_origin and table.cell(2, 0).is_spanned
    assert table.cell(3, 0).is_merge_origin and table.cell(3, 2).is_spanned
    assert warnings == [] and validate_package(data) == []


@pytest.mark.parametrize(
    ("style", "expect_vertical"), [("striped", False), ("plain", False), ("lines", False), ("grid", True)]
)
def test_border_styles(style: str, expect_vertical: bool) -> None:
    data, _ = _build(_table(style=style))
    xml = _slide_xml(data)
    has_vertical = bool(re.search(r"<a:lnL w=\"9525\"", xml))
    assert has_vertical is expect_vertical
    assert ('<a:lnB w="9525"' in xml) is (style in ("lines", "grid"))
    assert validate_package(data) == []


def test_heat_map_shades_each_column_against_its_own_range() -> None:
    data, _ = _build(_table(heatmap={"columns": ["Revenue", "Growth"], "low": "#FFFFFF", "high": "#C62828"}))
    table = _table_shape(data)
    revenue = [str(table.cell(r, 1).fill.fore_color.rgb) for r in (1, 2, 3)]
    growth = [str(table.cell(r, 2).fill.fore_color.rgb) for r in (1, 2, 3)]
    assert revenue[1] == "FFFFFF" and revenue[2] == "C62828"  # lowest white, highest red
    assert growth[1] == "FFFFFF" and growth[2] == "C62828"  # a tiny column is shaded too
    assert (
        str(table.cell(3, 1).text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    )  # white text on the dark cell
    assert str(table.cell(1, 0).fill.fore_color.rgb) != "C62828"  # a text column stays unshaded


def test_long_tables_split_over_slides_with_the_header_repeated() -> None:
    rows = [[i, f"Item {i}", i * 10] for i in range(1, 16)]
    deck = {
        "title": "Deck",
        "slides": [
            {
                "layout": "table",
                "title": "Fifteen items need two slides",
                "columns": ["#", "Item", "Value"],
                "rows": rows,
                "split_rows": 8,
                "notes": "say this",
            }
        ],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda n: b"")
    presentation = Presentation(io.BytesIO(built.data))
    assert len(presentation.slides) == 2
    first, second = (next(s for s in slide.shapes if s.has_table).table for slide in presentation.slides)
    assert len(first.rows) == 9 and len(second.rows) == 8  # header + 8, header + 7
    assert second.cell(0, 1).text == "Item" and second.cell(1, 0).text == "9"
    assert "continued, 2 of 2" in presentation.slides[1].shapes.title.text_frame.text
    assert presentation.slides[0].has_notes_slide and not presentation.slides[1].has_notes_slide
    assert validate_package(built.data) == []


def test_a_long_table_without_split_still_warns() -> None:
    rows = [[i, f"Item {i}", i] for i in range(1, 13)]
    deck = {
        "title": "Deck",
        "slides": [
            {
                "layout": "table",
                "title": "Twelve rows are too many here",
                "columns": ["#", "Item", "V"],
                "rows": rows,
            }
        ],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda n: b"")
    assert any("split the table" in w for s in built.slides for w in s.warnings)


def test_wrong_option_lengths_are_errors() -> None:
    with pytest.raises(BuildError, match="col_widths needs one number per column"):
        _build(_table(col_widths=[1, 2]))
    with pytest.raises(BuildError, match="formats needs one entry per column"):
        _build(_table(formats=["0"]))
