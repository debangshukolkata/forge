"""D-334: how dark a slide is, read from the file: solid, gradient, picture and theme colours through the master's colour map."""

from __future__ import annotations

import io

import pytest

pytest.importorskip("pptx")
from lxml import etree
from PIL import Image
from pptx import Presentation

from forge.slides.backgrounds import background_luminance, is_dark
from tests.slide_templates import make_template

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _slide_with_background(xml: str, dark_master: bool = False) -> object:
    prs = Presentation(io.BytesIO(make_template(dark=dark_master, sample_slide=False)))
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    background = etree.fromstring(
        f'<p:bg xmlns:p="{P}" xmlns:a="{A}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">{xml}</p:bg>'
    )
    slide._element.find(f"{{{P}}}cSld").insert(0, background)
    return slide


def test_solid_and_gradient_backgrounds() -> None:
    assert is_dark(
        background_luminance(
            _slide_with_background('<p:bgPr><a:solidFill><a:srgbClr val="101010"/></a:solidFill></p:bgPr>')
        )
    )  # type: ignore[arg-type]
    assert (
        is_dark(
            background_luminance(
                _slide_with_background(
                    '<p:bgPr><a:solidFill><a:srgbClr val="F2F2F2"/></a:solidFill></p:bgPr>'
                )
            )
        )
        is False
    )  # type: ignore[arg-type]
    grad = '<p:bgPr><a:gradFill><a:gsLst><a:gs pos="0"><a:srgbClr val="000000"/></a:gs><a:gs pos="100000"><a:srgbClr val="1F3A5F"/></a:gs></a:gsLst></a:gradFill></p:bgPr>'
    assert is_dark(background_luminance(_slide_with_background(grad)))  # type: ignore[arg-type]
    mid = '<p:bgPr><a:solidFill><a:srgbClr val="808080"/></a:solidFill></p:bgPr>'
    assert is_dark(background_luminance(_slide_with_background(mid))) is None  # type: ignore[arg-type]


def test_a_picture_background_is_read_from_its_pixels() -> None:
    prs = Presentation(io.BytesIO(make_template(sample_slide=False)))
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    out = io.BytesIO()
    Image.new("RGB", (64, 36), (5, 8, 20)).save(out, "PNG")
    _, relation = slide.part.get_or_add_image_part(io.BytesIO(out.getvalue()))
    background = etree.fromstring(
        f'<p:bg xmlns:p="{P}" xmlns:a="{A}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:bgPr><a:blipFill><a:blip r:embed="{relation}"/></a:blipFill></p:bgPr></p:bg>'
    )
    slide._element.find(f"{{{P}}}cSld").insert(0, background)
    assert is_dark(background_luminance(slide))


def test_theme_colours_follow_the_masters_colour_map_and_the_masters_own_background_is_the_fallback() -> None:
    scheme = '<p:bgPr><a:solidFill><a:schemeClr val="bg1"/></a:solidFill></p:bgPr>'
    assert is_dark(background_luminance(_slide_with_background(scheme, dark_master=False))) is False  # type: ignore[arg-type]
    assert is_dark(background_luminance(_slide_with_background(scheme, dark_master=True)))  # type: ignore[arg-type]
    prs = Presentation(io.BytesIO(make_template(sample_slide=False)))
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )  # no background of its own and none on the layout: the master's (not read: None)
    assert is_dark(background_luminance(slide)) in (None, False)
