"""D-296: `forge template-probe` tells how a template is read and used, without sending the template anywhere."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge import cli
from forge.slides import preview
from forge.slides.probe import run_probe
from tests.slide_templates import make_corporate_template, make_template


@pytest.fixture(autouse=True)
def no_powerpoint(monkeypatch: pytest.MonkeyPatch) -> None:
    def unavailable(*args: object, **kwargs: object) -> None:
        raise preview.PreviewUnavailable("PowerPoint is not available in this test")

    monkeypatch.setattr(preview, "export_slides", unavailable)


def test_the_probe_builds_a_deck_on_every_sample_and_reports(tmp_path: Path) -> None:
    template = tmp_path / "acme.pptx"
    template.write_bytes(make_corporate_template())
    result = run_probe(template, tmp_path / "out")
    assert result.slides_failed == 0 and result.slides_built >= 20
    deck = Presentation(str(tmp_path / "out" / "probe-deck.pptx"))
    assert len(deck.slides) == result.slides_built
    report = (tmp_path / "out" / "report.txt").read_text(encoding="utf-8")
    assert "Slide 2 [cards]" in report and "Probe slides" in report and "sample 2 [cards] long text" in report
    data = json.loads((tmp_path / "out" / "probe.json").read_text(encoding="utf-8"))
    assert [s["kind"] for s in data["samples"]] == ["title", "cards", "stats", "bullets", "thanks"]
    assert any("no pictures" in issue for issue in data["issues"])  # PowerPoint was not available here


def test_hide_text_keeps_the_structure_but_not_the_words(tmp_path: Path) -> None:
    template = tmp_path / "acme.pptx"
    template.write_bytes(make_corporate_template())
    run_probe(template, tmp_path / "out", hide_text=True)
    report = (tmp_path / "out" / "report.txt").read_text(encoding="utf-8")
    assert "Sample presentation title" not in report and "Pillar" not in report
    assert "[cards]" in report and "chars>" in report


def test_a_template_without_samples_and_a_four_by_three_one_are_explained(tmp_path: Path) -> None:
    plain = tmp_path / "plain.pptx"
    plain.write_bytes(make_template(sample_slide=False))
    result = run_probe(plain, tmp_path / "plain-out")
    assert any("no sample slides" in i for i in result.issues) and result.slides_built >= 9
    narrow = tmp_path / "narrow.pptx"
    narrow.write_bytes(make_template(10, 7.5))
    result = run_probe(narrow, tmp_path / "narrow-out")
    assert any("not 16:9" in i for i in result.issues) and result.slides_built == 0


def test_the_command_prints_where_the_files_are(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    template = tmp_path / "acme.pptx"
    template.write_bytes(make_corporate_template())
    assert cli.main(["template-probe", str(template), "--out", str(tmp_path / "o")]) == 0
    out = capsys.readouterr().out
    assert "probe-deck.pptx" in out and "report.txt" in out
    assert cli.main(["template-probe", str(tmp_path / "missing.pptx")]) == 2
