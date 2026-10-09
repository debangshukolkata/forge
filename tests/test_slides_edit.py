"""Reading and editing existing decks (D-241), with real python-pptx files."""

from __future__ import annotations

import asyncio
import io
import json
import zipfile
from pathlib import Path
from typing import Any

import pytest

pptx = pytest.importorskip("pptx")

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Pt  # noqa: E402

from forge.slides.build import build_deck  # noqa: E402
from forge.slides.edit import apply_operations, parse_operations  # noqa: E402
from forge.slides.open_deck import DeckError, open_deck, save_deck  # noqa: E402
from forge.slides.read import describe_deck  # noqa: E402
from forge.slides.spec import parse_deck  # noqa: E402
from forge.toolkit.base import ToolContext  # noqa: E402
from forge.tools.presentation import EditPresentation, ReadPresentation  # noqa: E402
from forge.workspace.create import create_workspace  # noqa: E402

OUTLINE = {
    "title": "Q3 Review",
    "slides": [
        {"layout": "title", "title": "Q3 Review", "subtitle": "What shipped"},
        {
            "layout": "bullets",
            "title": "We shipped three things",
            "bullets": ["SSO", {"text": "Faster search", "sub": ["210 ms median"]}],
            "notes": "Say thanks to Q3 team.",
        },
        {
            "layout": "chart",
            "title": "Users doubled",
            "kind": "column",
            "categories": ["Jun", "Jul"],
            "series": [{"name": "Users", "values": [12, 26]}],
        },
        {
            "layout": "table",
            "title": "Plan",
            "columns": ["What", "Who"],
            "rows": [["App", "Asha"], ["Export", "Ravi"]],
        },
        {"layout": "image", "title": "Screenshot", "image": "shot.png", "alt": "The dashboard"},
    ],
}


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (300, 150), (30, 90, 200)).save(buffer, "PNG")
    return buffer.getvalue()


def _deck(tmp_path: Path) -> Any:
    built = build_deck(parse_deck(json.dumps(OUTLINE)), lambda name: _png())
    path = tmp_path / "deck.pptx"
    path.write_bytes(built.data)
    return open_deck(path)


def _texts(deck: Any) -> list[str]:
    return [shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]


def _run(deck: Any, operations: list[dict[str, Any]]) -> list[str]:
    return apply_operations(deck, parse_operations(operations))


def test_reading_describes_slides_shapes_tables_charts_pictures_notes_and_layouts(tmp_path: Path) -> None:
    text = describe_deck(_deck(tmp_path))
    assert "Slide 2" in text and "We shipped three things" in text and "notes: Say thanks" in text
    assert "  210 ms median" in text  # a sub-bullet is indented
    assert "series 'Users': 12.0, 26.0" in text and "categories: Jun, Jul" in text
    assert "| What | Who |" in text and "| Export | Ravi |" in text
    assert "alt text: The dashboard" in text
    assert "Layouts this deck can add slides from" in text and '"Title and Content"' in text
    assert "[id " in text  # every shape can be referred to by its id


def test_replace_text_keeps_formatting_and_handles_text_split_across_runs(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    slide = deck.slides[1]
    frame = next(s.text_frame for s in slide.shapes if s.has_text_frame and "SSO" in s.text_frame.text)
    paragraph = frame.paragraphs[0]
    paragraph.runs[0].text = "Single sign"
    extra = paragraph.add_run()
    extra.text = "-on (SSO)"
    extra.font.bold = True
    done = _run(deck, [{"op": "replace_text", "find": "sign-on", "replace": "login"}])
    assert "1 occurrence" in done[0]
    assert paragraph.text == "Single login (SSO)"
    _run(deck, [{"op": "replace_text", "find": "Q3", "replace": "Q4", "notes": True}])
    assert (
        "Q4" in _texts(deck)[0]
        and deck.slides[1].notes_slide.notes_text_frame.text == "Say thanks to Q4 team."
    )
    with pytest.raises(DeckError, match="was not found"):
        _run(deck, [{"op": "replace_text", "find": "nothing like this", "replace": "x"}])


def test_set_text_takes_the_look_of_the_text_it_replaces(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    shape = next(s for s in deck.slides[1].shapes if s.has_text_frame and "SSO" in s.text_frame.text)
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(31)
    _run(deck, [{"op": "set_text", "slide": 2, "shape": shape.shape_id, "text": "One\n  Two\nThree"}])
    paragraphs = shape.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["One", "Two", "Three"] and [p.level for p in paragraphs] == [
        0,
        1,
        0,
    ]
    assert paragraphs[2].runs[0].font.size == Pt(31)  # the first run's size carries to every new paragraph
    with pytest.raises(DeckError, match="no shape"):
        _run(deck, [{"op": "set_text", "slide": 2, "shape": "No such name", "text": "x"}])


def test_delete_move_and_duplicate_slides(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    _run(
        deck,
        [
            {"op": "duplicate_slide", "slide": 2},
            {"op": "move_slide", "slide": 1, "to": 3},
            {"op": "delete_slide", "slide": 6},
        ],
    )
    titles = [s.shapes.title.text_frame.text for s in deck.slides]
    assert titles == [
        "We shipped three things",
        "We shipped three things",
        "Q3 Review",
        "Users doubled",
        "Plan",
    ]
    reopened = open_deck_from_bytes(save_deck(deck))
    assert len(reopened.slides) == 5  # the dropped slide is gone from the saved file
    with pytest.raises(DeckError, match="there is no slide 9"):
        _run(deck, [{"op": "delete_slide", "slide": 9}])


def open_deck_from_bytes(data: bytes) -> Any:
    return Presentation(io.BytesIO(data))


def test_duplicating_a_slide_with_a_picture_copies_the_picture_and_a_chart_is_refused(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    _run(deck, [{"op": "duplicate_slide", "slide": 5}])
    pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 2 and pictures[1].image.size == (300, 150)
    assert (
        "alt text: The dashboard" in describe_deck(open_deck_from_bytes(save_deck(deck))).split("Slide 6")[1]
    )
    with pytest.raises(DeckError, match="chart"):
        _run(deck, [{"op": "duplicate_slide", "slide": 3}])


def test_charts_and_tables_are_updated(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    chart = next(s for s in deck.slides[2].shapes if s.has_chart)
    table = next(s for s in deck.slides[3].shapes if s.has_table)
    _run(
        deck,
        [
            {
                "op": "update_chart",
                "slide": 3,
                "shape": chart.shape_id,
                "categories": ["Aug", "Sep", "Oct"],
                "series": [{"name": "Users", "values": [30, 33, 41]}],
            },
            {
                "op": "update_table",
                "slide": 4,
                "shape": table.name,
                "rows": [["What", "Who"], ["App", "Meera"], ["Export", "Ravi"]],
            },
        ],
    )
    assert list(chart.chart.plots[0].categories) == ["Aug", "Sep", "Oct"] and list(
        chart.chart.plots[0].series[0].values
    ) == [30, 33, 41]
    assert table.table.cell(1, 1).text == "Meera"
    with pytest.raises(DeckError, match="3 rows and 2 columns"):
        _run(deck, [{"op": "update_table", "slide": 4, "shape": table.name, "rows": [["only", "one row"]]}])
    with pytest.raises(DeckError, match="is not a chart"):
        _run(
            deck,
            [
                {
                    "op": "update_chart",
                    "slide": 4,
                    "shape": table.name,
                    "categories": ["a"],
                    "series": [{"name": "s", "values": [1]}],
                }
            ],
        )


def test_slides_are_added_from_the_decks_own_layouts_and_a_template_works(tmp_path: Path) -> None:
    built = build_deck(parse_deck(json.dumps(OUTLINE)), lambda name: _png()).data
    template = tmp_path / "company.potx"
    with zipfile.ZipFile(io.BytesIO(built)) as source, zipfile.ZipFile(template, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.presentation.main+xml", b"presentationml.template.main+xml"
                )
            target.writestr(item, data)
    deck = open_deck(template)  # a .potx opens like a deck
    _run(
        deck,
        [
            {
                "op": "add_slide",
                "layout": "title and content",
                "title": "Next steps",
                "body": ["Ship it", "  by Friday"],
                "notes": "Be brief",
                "at": 1,
            },
            {
                "op": "add_slide",
                "layout": "Title Slide",
                "title": "Hello",
                "placeholders": {"1": "A subtitle"},
            },
        ],
    )
    first = deck.slides[0]
    assert (
        first.shapes.title.text_frame.text == "Next steps"
        and first.notes_slide.notes_text_frame.text == "Be brief"
    )
    body = next(p for p in first.placeholders if p.placeholder_format.idx == 1)
    assert [(p.text, p.level) for p in body.text_frame.paragraphs] == [("Ship it", 0), ("by Friday", 1)]
    last = deck.slides[len(deck.slides) - 1]
    assert [p.placeholder_format.idx for p in last.placeholders] == [
        0,
        1,
    ]  # no empty "Click to add text" left behind
    with pytest.raises(DeckError, match=r'no layout called .*Layouts: "Title Slide"'):
        _run(deck, [{"op": "add_slide", "layout": "Nonexistent layout"}])


def test_a_failing_operation_stops_everything_and_names_itself(tmp_path: Path) -> None:
    with pytest.raises(DeckError, match=r"operation 2 \(delete_slide\) failed.*Nothing was saved"):
        _run(
            _deck(tmp_path),
            [{"op": "replace_text", "find": "Q3", "replace": "Q4"}, {"op": "delete_slide", "slide": 99}],
        )


def test_bad_operations_are_explained() -> None:
    with pytest.raises(DeckError) as raised:
        parse_operations([{"op": "explode"}, {"op": "set_notes", "slide": 1}])
    assert "Operations are:" in str(raised.value)


@pytest.fixture
def context(tmp_path: Path, original_repo: Path) -> ToolContext:
    return ToolContext(workspace=create_workspace(original_repo, tmp_path / "ws", "backend"))


def test_the_tools_read_and_edit_through_the_workspace(context: ToolContext, tmp_path: Path) -> None:
    workspace = context.workspace
    built = build_deck(parse_deck(json.dumps(OUTLINE)), lambda name: _png()).data
    workspace.write_bytes("docs/deck.pptx", built)
    read = asyncio.run(ReadPresentation().run(ReadPresentation.Args(path="docs/deck.pptx"), context))
    assert read.ok and "Slide 1" in read.content and "Layouts this deck" in read.content
    edit = asyncio.run(
        EditPresentation().run(
            EditPresentation.Args(
                path="docs/deck.pptx", operations=[{"op": "replace_text", "find": "Q3", "replace": "Q4"}]
            ),
            context,
        )
    )
    assert edit.ok and "Saved docs/deck.pptx" in edit.content and "replaced" in edit.content
    assert "Q4 Review" in describe_deck(open_deck(workspace.path_of("docs/deck.pptx")))
    assert (
        len([c for c in workspace.changes() if c.path == "docs/deck.pptx"]) == 2
    )  # build + edit, both undoable
    failed = asyncio.run(
        EditPresentation().run(
            EditPresentation.Args(path="docs/deck.pptx", operations=[{"op": "delete_slide", "slide": 50}]),
            context,
        )
    )
    assert not failed.ok and "Nothing was saved" in failed.content
    assert "Q4 Review" in describe_deck(open_deck(workspace.path_of("docs/deck.pptx")))  # untouched


def test_a_template_is_never_overwritten_and_other_files_are_refused(context: ToolContext) -> None:
    workspace = context.workspace
    workspace.write_bytes("company.potx", b"PK")
    refused = asyncio.run(
        EditPresentation().run(EditPresentation.Args(path="company.potx", operations=[]), context)
    )
    assert not refused.ok and "output" in refused.content
    workspace.write_text("old.ppt", "not a deck")
    old = asyncio.run(ReadPresentation().run(ReadPresentation.Args(path="old.ppt"), context))
    assert not old.ok and ".pptx" in old.content
    broken = asyncio.run(ReadPresentation().run(ReadPresentation.Args(path="company.potx"), context))
    assert not broken.ok and "could not be opened" in broken.content
