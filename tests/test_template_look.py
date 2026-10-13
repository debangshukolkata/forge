"""D-331: `template_look.py` gives a short summary, a palettes picture and contact sheets of a large template."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides import look, preview
from forge.slides.importer import import_slides
from tests.slide_templates import make_corporate_template, make_template


def _two_master_template(tmp_path: Path) -> Path:
    first = Presentation(io.BytesIO(make_corporate_template()))
    second = Presentation(
        io.BytesIO(
            make_template(
                accents=("C0392B", "2E86C1", "27AE60", "F39C12", "8E44AD", "566573"),
                fonts=("Cambria", "Verdana"),
            )
        )
    )
    import_slides(first, second, design="keep")
    path = tmp_path / "big.pptx"
    first.save(str(path))
    return path


def test_the_summary_palettes_and_zip_are_written_without_powerpoint(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(look, "available", lambda: False)
    path = _two_master_template(tmp_path)
    result = look.run_look(path, tmp_path / "out")
    summary = (tmp_path / "out" / "summary.txt").read_text(encoding="utf-8")
    assert "2 master(s)" in summary and "master 1" in summary and "master 2" in summary
    assert "accent1=0B6E4F" in summary and "accent1=C0392B" in summary and "Cambria" in summary
    assert "layout pictures (layouts-N.png)" in summary and "sample slide(s)" in summary
    assert (tmp_path / "out" / "palettes.png").is_file() and "palettes.png" in result.files
    assert any("PowerPoint is not available" in n for n in result.notes)
    with zipfile.ZipFile(tmp_path / "out" / "template-look.zip") as bundle:
        assert set(bundle.namelist()) == {"summary.txt", "palettes.png"}
    assert not (tmp_path / "out" / "_work").exists()


def test_layouts_are_taken_round_robin_so_every_master_shows(tmp_path: Path) -> None:
    path = _two_master_template(tmp_path)
    order = look._layout_deck(path, tmp_path / "layouts.pptx")
    masters = [m for m, _ in order]
    assert masters[:4] == [1, 2, 1, 2] and set(masters) == {1, 2}
    deck = Presentation(str(tmp_path / "layouts.pptx"))
    assert len(deck.slides) == len(order) <= look.MAX_LAYOUTS
    assert deck.slides[0].shapes.title is None or deck.slides[0].shapes.title.text_frame.text == order[0][1]


def test_many_sample_slides_are_thinned_evenly(tmp_path: Path) -> None:
    prs = Presentation(io.BytesIO(make_template()))
    for _ in range(100):
        prs.slides.add_slide(prs.slide_layouts[0])
    source = tmp_path / "many.pptx"
    prs.save(str(source))
    kept = look._sample_deck(source, tmp_path / "samples.pptx")
    assert kept <= look.MAX_SAMPLES and kept >= look.MAX_SAMPLES // 2


@pytest.mark.skipif(not preview.available(), reason="PowerPoint is not installed here")
def test_with_powerpoint_the_contact_sheets_are_made(tmp_path: Path) -> None:
    result = look.run_look(_two_master_template(tmp_path), tmp_path / "out")
    names = set(result.files)
    assert "layouts-1.png" in names and "samples-1.png" in names and "template-look.zip" in names
    assert (tmp_path / "out" / "layouts-1.png").stat().st_size > 10_000
