"""Test helpers: company-style templates generated from python-pptx's default template (which is 4:3 and plain)."""

from __future__ import annotations

import io
import re

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

EMU = 914400


def make_template(
    width_in: float = 13.333,
    height_in: float = 7.5,
    accents: tuple[str, ...] = ("0B6E4F", "F2A541", "3C6E9E", "B5446E", "6B5BD2", "5B6675"),
    fonts: tuple[str, str] = ("Georgia", "Arial"),
    dark: bool = False,
    sample_slide: bool = True,
    strip_layouts: tuple[str, ...] = (),
) -> bytes:
    """A template: default layouts stretched to the given size, its own accent colours and fonts, optionally a dark master."""
    prs = Presentation()
    old_w, old_h = prs.slide_width, prs.slide_height
    new_w, new_h = int(width_in * EMU), int(height_in * EMU)
    sx, sy = new_w / old_w, new_h / old_h
    prs.slide_width, prs.slide_height = new_w, new_h
    for holder in (prs.slide_master, *prs.slide_layouts):
        for shape in holder.shapes:
            if shape._element.xpath(
                "./p:spPr/a:xfrm"
            ):  # a placeholder without its own box inherits the (already scaled) master's
                shape.left, shape.top = int(shape.left * sx), int(shape.top * sy)
                shape.width, shape.height = int(shape.width * sx), int(shape.height * sy)
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    for index, colour in enumerate(accents, start=1):
        xml = re.sub(
            rf'(<a:accent{index}>)<a:srgbClr val="[0-9A-Fa-f]{{6}}"/>', rf'\1<a:srgbClr val="{colour}"/>', xml
        )
    xml = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]+"', rf'\1{fonts[0]}"', xml)
    xml = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]+"', rf'\1{fonts[1]}"', xml)
    if dark:
        xml = re.sub(r"(<a:dk1>)<a:sysClr[^>]*/>", r'\1<a:srgbClr val="0E1621"/>', xml)
        xml = re.sub(r"(<a:lt1>)<a:sysClr[^>]*/>", r'\1<a:srgbClr val="F4F6FA"/>', xml)
    theme._blob = xml.encode("utf-8")
    if dark:
        cmap = prs.slide_master._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap"
        )
        cmap.set("bg1", "dk1")
        cmap.set("tx1", "lt1")
    for name in strip_layouts:
        for layout in list(prs.slide_layouts):
            if layout.name == name:
                prs.slide_layouts.remove(layout)
    if sample_slide:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text_frame.text = "SAMPLE SLIDE FROM THE TEMPLATE"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_corporate_template() -> bytes:
    """A template that carries sample slides (title, three cards, big numbers, bullets, thank you), each with the company's
    confidentiality line, as a corporate template does."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation(io.BytesIO(make_template(sample_slide=False)))
    layouts = {layout.name: layout for layout in prs.slide_layouts}

    def confidential(slide: object) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(5), Inches(0.3))  # type: ignore[attr-defined]
        box.text_frame.text = "Acme Corp. Confidential"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(10)

    def column(slide: object, x: float, heading: str, text: str, number: str | None = None) -> None:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(3.6), Inches(4.2)
        )  # type: ignore[attr-defined]
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xEE, 0xF3, 0xF0)
        top = 2.2
        if number:
            n = slide.shapes.add_textbox(Inches(x + 0.2), Inches(top), Inches(3.2), Inches(1.0))  # type: ignore[attr-defined]
            n.text_frame.text = number
            n.text_frame.paragraphs[0].runs[0].font.size = Pt(48)
            top += 1.1
        h = slide.shapes.add_textbox(Inches(x + 0.2), Inches(top), Inches(3.2), Inches(0.6))  # type: ignore[attr-defined]
        h.text_frame.text = heading
        h.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        h.text_frame.paragraphs[0].runs[0].font.bold = True
        t = slide.shapes.add_textbox(Inches(x + 0.2), Inches(top + 0.7), Inches(3.2), Inches(1.6))  # type: ignore[attr-defined]
        t.text_frame.word_wrap = True
        t.text_frame.text = text
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    s = prs.slides.add_slide(layouts["Title Slide"])
    s.shapes.title.text_frame.text = "Sample presentation title"
    s.placeholders[1].text_frame.text = "Sample subtitle goes here"
    confidential(s)
    s = prs.slides.add_slide(layouts["Title Only"])
    s.shapes.title.text_frame.text = "Three pillars of the sample"
    for i, x in enumerate((0.8, 4.85, 8.9)):
        column(s, x, f"Pillar {i + 1}", "A sentence that explains this pillar in plain words.")
    confidential(s)
    s = prs.slides.add_slide(layouts["Title Only"])
    s.shapes.title.text_frame.text = "Numbers that matter"
    for i, x in enumerate((0.8, 4.85, 8.9)):
        column(s, x, "Metric label", "Context for the metric.", number=f"{i + 1}0%")
    confidential(s)
    s = prs.slides.add_slide(layouts["Title and Content"])
    s.shapes.title.text_frame.text = "Sample bullet slide"
    s.placeholders[1].text_frame.text = "First sample point"
    for text in ("Second sample point", "Third sample point"):
        s.placeholders[1].text_frame.add_paragraph().text = text
    confidential(s)
    s = prs.slides.add_slide(layouts["Title Slide"])
    s.shapes.title.text_frame.text = "Thank you"
    s.placeholders[1].text_frame.text = "Questions? name@acme.example"
    confidential(s)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
