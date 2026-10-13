"""P2.3: the bundled icons parse into drawable contours; the SVG path reader handles curves, arcs and relative commands."""

from __future__ import annotations

import io
import json
import math

import pytest
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.icons import IconError, icon_names, load_icon, parse_path, search_icons
from forge.slides.spec import parse_deck
from forge.slides.validate import validate_package


def test_lines_and_relative_moves() -> None:
    (contour,) = parse_path("M1 2 L3 4 h2 v-1 l-1 -1")
    assert contour.points == ((1, 2), (3, 4), (5, 4), (5, 3), (4, 2)) and not contour.closed
    closed = parse_path("M0 0H4V4H0Z")[0]
    assert closed.closed and closed.points[0] == closed.points[-1] == (0, 0)


def test_several_subpaths_and_implicit_line_after_move() -> None:
    contours = parse_path("M0 0 2 0 M5 5 l1 1")
    assert [c.points for c in contours] == [((0, 0), (2, 0)), ((5, 5), (6, 6))]


def test_a_semicircle_arc_has_the_right_radius_and_passes_through_the_top() -> None:
    (contour,) = parse_path("M0 0 A1 1 0 0 1 2 0")
    assert contour.points[-1] == (2, 0)
    for x, y in contour.points:
        assert math.hypot(x - 1, y) == pytest.approx(1, abs=1e-6)
    assert min(y for _, y in contour.points) == pytest.approx(
        -1, abs=0.01
    )  # sweep flag 1: through y = -1 (up on screen)


def test_arc_flags_glued_to_numbers_and_the_large_arc() -> None:
    glued = parse_path("M0 0a1 1 0 01 2 0")[0]  # "01" is large=0, sweep=1
    assert glued.points[-1] == (2, 0)
    large = parse_path("M0 0 A2 2 0 1 1 2 0")[0]
    assert max(abs(y) for _, y in large.points) > 2  # the large arc bulges past the radius


def test_cubic_quadratic_and_smooth_curves_end_at_their_target() -> None:
    assert parse_path("M0 0 C0 10 10 10 10 0")[0].points[-1] == (10, 0)
    assert parse_path("M0 0 Q5 10 10 0")[0].points[-1] == (10, 0)
    smooth = parse_path("M0 0 C0 5 5 5 5 0 S10 -5 10 0")[0]
    assert smooth.points[-1] == (10, 0) and len(smooth.points) > 20


def test_a_bad_path_is_an_error_in_words() -> None:
    with pytest.raises(IconError, match="cannot read a number"):
        parse_path("M1 L2")
    with pytest.raises(IconError, match="must be 0 or 1"):
        parse_path("M0 0 A1 1 0 5 1 2 0")


def test_the_bundles_hold_the_expected_icons() -> None:
    assert len(icon_names("lucide")) > 1500 and len(icon_names("material")) >= 150
    assert {"rocket", "heart", "chart-bar", "users"} <= set(icon_names("lucide"))
    assert {"rocket_launch", "check", "trending_up"} <= set(icon_names("material"))


def test_every_bundled_icon_parses_into_finite_points_inside_the_frame() -> None:
    for library in ("lucide", "material"):
        for name in icon_names(library):
            icon = load_icon(f"{library}:{name}")
            assert icon.contours, f"{library}:{name} has no drawing"
            for contour in icon.contours:
                for x, y in contour.points:
                    assert math.isfinite(x) and math.isfinite(y)
                    assert -0.3 <= x <= 1.3 and -0.3 <= y <= 1.3, (library, name, x, y)


def test_styles_and_names() -> None:
    lucide = load_icon("lucide:rocket")
    assert lucide.stroke and lucide.stroke_fraction == pytest.approx(2 / 24)
    material = load_icon("material:rocket_launch")
    assert not material.stroke and material.name == "material:rocket_launch"
    assert load_icon("rocket").name == "lucide:rocket"  # a bare name tries Lucide first
    assert load_icon("rocket-launch").name.endswith("rocket_launch") or load_icon(
        "rocket-launch"
    ).name.startswith("lucide:")


def test_unknown_icons_suggest_alternatives() -> None:
    with pytest.raises(IconError, match=r"similar: .*rocket"):
        load_icon("rockets")
    with pytest.raises(IconError, match="unknown icon library"):
        load_icon("fontawesome:rocket")


def test_search_finds_icons_by_name_and_by_tag() -> None:
    assert "lucide:rocket" in search_icons("rocket launch")
    assert any("chart" in name for name in search_icons("growth chart"))
    assert any(name.startswith("material:") for name in search_icons("shield security"))
    assert search_icons("zzzzqq") == []


# --- the icon element on a freeform slide ---


def _build(elements: list[dict]) -> tuple[bytes, list[str]]:  # type: ignore[type-arg]
    deck = {
        "title": "Deck",
        "slides": [{"layout": "freeform", "title": "Icons carry one idea each", "elements": elements}],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda name: b"")
    return built.data, [w for s in built.slides for w in s.warnings]


def _shapes(data: bytes) -> list:  # type: ignore[type-arg]
    return list(Presentation(io.BytesIO(data)).slides[0].shapes)


def test_an_outline_icon_is_one_recolourable_freeform_with_alt_text() -> None:
    data, warnings = _build(
        [{"type": "icon", "name": "lucide:rocket", "x": 1, "y": 2, "w": 1, "h": 1, "color": "#C62828"}]
    )
    icon = next(s for s in _shapes(data) if s.name.startswith("Icon "))
    assert str(icon.line.color.rgb) == "C62828" and icon.line.width.pt >= 1
    assert icon._element.nvSpPr.cNvPr.get("descr") == "rocket icon"
    assert warnings == [] and validate_package(data) == []


def test_a_filled_icon_on_a_badge_is_a_circle_with_a_smaller_icon_inside() -> None:
    data, warnings = _build(
        [
            {
                "type": "icon",
                "name": "material:check",
                "x": 2,
                "y": 2,
                "w": 1.2,
                "h": 1.2,
                "color": "#FFFFFF",
                "badge": "accent",
            }
        ]
    )
    shapes = _shapes(data)
    badge = next(s for s in shapes if s.shape_type == 1)  # an auto shape: the circle
    icon = next(s for s in shapes if s.name.startswith("Icon "))
    assert icon.width < badge.width and icon.left > badge.left and icon.top > badge.top
    assert str(icon.fill.fore_color.rgb) == "FFFFFF" and warnings == [] and validate_package(data) == []


def test_unknown_icon_and_non_square_box_are_warnings() -> None:
    _, warnings = _build([{"type": "icon", "name": "rockets-nope", "x": 1, "y": 2, "w": 1, "h": 1}])
    assert any("no icon named" in w for w in warnings)
    _, warnings = _build([{"type": "icon", "name": "rocket", "x": 1, "y": 2, "w": 2, "h": 1}])
    assert any("needs a square box" in w for w in warnings)


def test_several_icons_in_a_row_container_do_not_overlap() -> None:
    row = {
        "type": "row",
        "x": 1,
        "y": 2,
        "w": 6,
        "h": 1,
        "gap": 0.5,
        "children": [
            {"type": "icon", "name": n, "badge": "accent", "color": "#FFFFFF", "length": 1, "cross": 1}
            for n in ("heart", "users", "chart-bar")
        ],
    }
    data, warnings = _build([row])
    assert warnings == [] and validate_package(data) == []


def test_the_find_icon_tool_lists_names() -> None:
    import asyncio

    from forge.toolkit.base import ToolContext
    from forge.tools.presentation import FindIcon, presentation_tools

    assert "find_icon" in {t.name for t in presentation_tools()}
    result = asyncio.run(FindIcon().run(FindIcon.Args(query="growth chart"), ToolContext(workspace=None)))  # type: ignore[arg-type]
    assert result.ok and "lucide:" in result.content
    nothing = asyncio.run(FindIcon().run(FindIcon.Args(query="zzzzqq"), ToolContext(workspace=None)))  # type: ignore[arg-type]
    assert "No icon matched" in nothing.content


def test_the_licence_texts_ship_with_the_data() -> None:
    from importlib import resources

    folder = resources.files("forge") / "data" / "icons"
    assert "ISC License" in (folder / "LICENSE-lucide.txt").read_text(encoding="utf-8")
    assert "Apache License" in (folder / "LICENSE-material-symbols.txt").read_text(encoding="utf-8")
