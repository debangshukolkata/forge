"""Presentation making (D-240): the outline, the builder (the real python-pptx), the fit estimate, the two
tools and the PowerPoint preview (only where PowerPoint exists)."""

from __future__ import annotations

import asyncio
import io
import json
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from PIL import Image  # noqa: E402

from forge.slides import fit, preview  # noqa: E402
from forge.slides.build import BuildError, build_deck  # noqa: E402
from forge.slides.spec import SpecError, parse_deck  # noqa: E402
from forge.toolkit.base import ToolContext  # noqa: E402
from forge.tools.presentation import BuildPresentation, PreviewPresentation  # noqa: E402
from forge.workspace.create import create_workspace  # noqa: E402

OUTLINE = {
    "title": "Q3 Review",
    "theme": "warm",
    "slides": [
        {"layout": "title", "title": "Q3 Review", "subtitle": "What shipped"},
        {
            "layout": "bullets",
            "title": "We shipped three things",
            "bullets": ["SSO", {"text": "Faster search", "sub": ["210 ms median"]}],
            "notes": "Say thanks.",
        },
        {"layout": "stats", "title": "Growth", "items": [{"value": "42%", "label": "weekly users"}]},
        {
            "layout": "chart",
            "title": "Users doubled",
            "kind": "line",
            "categories": ["Jun", "Jul"],
            "series": [{"name": "Users", "values": [12, 26]}],
        },
        {
            "layout": "table",
            "title": "Plan",
            "columns": ["What", "Who"],
            "rows": [["App", "Asha"], ["Export", "Ravi"]],
        },
        {
            "layout": "two_column",
            "title": "Review",
            "left": {"heading": "Good", "bullets": ["a"]},
            "right": {"bullets": ["b"]},
        },
        {"layout": "quote", "quote": "Simple wins.", "by": "A customer"},
        {"layout": "image", "title": "Screenshot", "image": "shot.png", "alt": "The dashboard"},
        {"layout": "section", "title": "Questions"},
        {"layout": "closing", "title": "Thank you"},
    ],
}


def _png(width: int = 400, height: int = 200) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (30, 90, 200)).save(buffer, "PNG")
    return buffer.getvalue()


def test_the_outline_reports_every_problem_in_words() -> None:
    with pytest.raises(SpecError) as raised:
        parse_deck(
            json.dumps(
                {"title": "T", "slides": [{"layout": "bullets", "title": "No bullets"}, {"layout": "bogus"}]}
            )
        )
    message = str(raised.value)
    assert "bullets" in message and "Layouts are:" in message
    with pytest.raises(SpecError, match="not valid JSON"):
        parse_deck("{not json")
    with pytest.raises(SpecError):
        parse_deck(json.dumps({"title": "T", "slides": []}))
    assert (
        parse_deck("title: T\nslides:\n  - {layout: title, title: Hi}\n", as_yaml=True).slides[0].layout
        == "title"
    )


def test_a_deck_with_every_layout_is_a_real_editable_presentation() -> None:
    built = build_deck(parse_deck(json.dumps(OUTLINE)), lambda name: _png())
    assert [s.layout for s in built.slides][:3] == ["title", "bullets", "stats"] and len(built.slides) == 10
    assert all(not s.warnings for s in built.slides)
    from pptx import Presentation

    deck = Presentation(io.BytesIO(built.data))
    assert len(deck.slides) == 10
    for slide in deck.slides:
        assert (
            slide.shapes.title is not None and slide.shapes.title.text_frame.text
        )  # a real title on every slide
    assert deck.slides[1].notes_slide.notes_text_frame.text == "Say thanks."
    kinds = {shape.shape_type for slide in deck.slides for shape in slide.shapes}
    assert any(
        shape.has_chart for slide in deck.slides for shape in slide.shapes
    )  # a native chart, not a picture
    assert any(shape.has_table for slide in deck.slides for shape in slide.shapes)
    assert kinds  # nothing else to say about shape types
    picture = next(s for slide in deck.slides for s in slide.shapes if s.shape_type == 13)
    assert picture._element.nvPicPr.cNvPr.get("descr") == "The dashboard"  # alt text
    assert deck.core_properties.title == "Q3 Review"


def test_text_that_cannot_fit_is_reported_and_text_that_can_is_shrunk() -> None:
    crowded = [
        f"Bullet number {i}: "
        + "a fairly long sentence that goes on and on about nothing in particular. " * 3
        for i in range(9)
    ]
    deck = {"title": "T", "slides": [{"layout": "bullets", "title": "Crowded slide", "bullets": crowded}]}
    warnings = build_deck(parse_deck(json.dumps(deck)), lambda n: b"").slides[0].warnings
    assert any("bullets" in w for w in warnings) and any("too much text" in w for w in warnings)
    size, fits = fit.fit_size([("short", 1.0, False)], 10, 4, 28, 16)
    assert (size, fits) == (28, True)
    size, fits = fit.fit_size([("word " * 400, 1.0, False)], 4, 2, 28, 16)
    assert size == 16 and not fits


def test_bad_charts_and_tables_are_named_by_slide() -> None:
    chart = {
        "layout": "chart",
        "title": "C",
        "categories": ["a", "b"],
        "series": [{"name": "s", "values": [1]}],
    }
    with pytest.raises(BuildError, match=r"slide 1.*series 's' has 1 values"):
        build_deck(parse_deck(json.dumps({"title": "T", "slides": [chart]})), lambda n: b"")
    table = {"layout": "table", "title": "T", "columns": ["a", "b"], "rows": [["only one"]]}
    with pytest.raises(BuildError, match="every table row needs 2 cells"):
        build_deck(parse_deck(json.dumps({"title": "T", "slides": [table]})), lambda n: b"")


@pytest.fixture
def context(tmp_path: Path, original_repo: Path) -> ToolContext:
    return ToolContext(workspace=create_workspace(original_repo, tmp_path / "ws", "backend"))


def test_the_build_tool_writes_the_file_through_the_workspace(context: ToolContext) -> None:
    workspace = context.workspace
    (workspace.repo_dir / "backend" / "shot.png").write_bytes(_png())
    outline = json.loads(json.dumps(OUTLINE))
    outline["slides"][7]["image"] = "backend/shot.png"
    workspace.write_text("docs/deck.json", json.dumps(outline))
    result = asyncio.run(
        BuildPresentation().run(
            BuildPresentation.Args(outline="docs/deck.json", output="docs/deck.pptx"), context
        )
    )
    assert result.ok and "10 slides" in result.content and "No layout warnings" in result.content
    written = workspace.path_of("docs/deck.pptx")
    assert written.read_bytes()[:2] == b"PK"  # a zip, as every .pptx is
    assert any(change.path == "docs/deck.pptx" for change in workspace.changes())  # recorded like any write


def test_the_build_tool_explains_what_is_wrong(context: ToolContext) -> None:
    workspace = context.workspace
    wrong = asyncio.run(
        BuildPresentation().run(BuildPresentation.Args(outline="x.json", output="x.docx"), context)
    )
    assert not wrong.ok and ".pptx" in wrong.content
    missing = asyncio.run(
        BuildPresentation().run(BuildPresentation.Args(outline="nope.json", output="x.pptx"), context)
    )
    assert not missing.ok and "Write it first" in missing.content
    workspace.write_text("bad.json", '{"title": "T", "slides": [{"layout": "bullets", "title": "x"}]}')
    bad = asyncio.run(
        BuildPresentation().run(BuildPresentation.Args(outline="bad.json", output="x.pptx"), context)
    )
    assert not bad.ok and "The outline has problems" in bad.content
    outside = asyncio.run(
        BuildPresentation().run(
            BuildPresentation.Args(outline="bad.json", output="../../escape.pptx"), context
        )
    )
    assert not outside.ok  # the write jail still applies


def test_preview_says_so_when_powerpoint_is_missing(
    context: ToolContext, monkeypatch: pytest.MonkeyPatch
) -> None:
    context.workspace.write_bytes(
        "deck.pptx", build_deck(parse_deck(json.dumps(OUTLINE)), lambda n: _png()).data
    )
    monkeypatch.setattr(preview, "available", lambda: False)
    result = asyncio.run(PreviewPresentation().run(PreviewPresentation.Args(path="deck.pptx"), context))
    assert result.ok and "No pictures could be made" in result.content and "not looked at" in result.content


@pytest.mark.skipif(not preview.available(), reason="PowerPoint is not installed here")
def test_powerpoint_makes_a_picture_of_every_slide(context: ToolContext) -> None:
    context.workspace.write_bytes(
        "deck.pptx", build_deck(parse_deck(json.dumps(OUTLINE)), lambda n: _png()).data
    )
    result = asyncio.run(PreviewPresentation().run(PreviewPresentation.Args(path="deck.pptx"), context))
    assert result.ok and ".forge/reports/slides/deck/slide-1.png" in result.content
    pictures = sorted((context.workspace.forge_dir / "reports" / "slides" / "deck").glob("slide-*.png"))
    assert len(pictures) == 10
    with Image.open(pictures[0]) as first:
        assert first.size == (preview.WIDTH, preview.HEIGHT)
