"""P2.8 (D-316): hidden slides, transitions, entrance animations, sections and comments."""

from __future__ import annotations

import io
import json

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides import features
from forge.slides.build import build_deck
from forge.slides.edit import apply_operations, parse_operations
from forge.slides.open_deck import DeckError
from forge.slides.read import describe_deck
from forge.slides.spec import parse_deck
from forge.slides.validate import validate_package


def _deck(extra: dict | None = None, slides: list[dict] | None = None) -> bytes:  # type: ignore[type-arg]
    outline = {
        "title": "Features",
        "slides": slides
        or [
            {"layout": "title", "title": "Features"},
            {
                "layout": "bullets",
                "title": "One",
                "bullets": ["a", "b"],
                "section": "Part one",
                "transition": "push",
            },
            {"layout": "bullets", "title": "Two", "bullets": ["c"], "hidden": True, "animate": "fade"},
            {"layout": "bullets", "title": "Three", "bullets": ["d"], "section": "Part two"},
        ],
        **(extra or {}),
    }
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: b"")
    assert not validate_package(built.data)
    return built.data


def _open(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def _save(prs: Presentation) -> bytes:
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_the_outline_sets_hidden_transition_section_and_animation() -> None:
    prs = _open(_deck({"transition": "fade"}))
    slides = list(prs.slides)
    assert [features.is_hidden(s) for s in slides] == [False, False, True, False]
    assert [features.transition_of(s) for s in slides] == [
        "fade",
        "push",
        "fade",
        "fade",
    ]  # the deck's, the slide's own
    assert [(s.first_slide, s.name) for s in features.read_sections(prs)] == [
        (1, "Start"),
        (2, "Part one"),
        (4, "Part two"),
    ]
    assert features.animations_of(slides[2])  # the content of slide 3 comes in on a click
    assert features.animations_of(slides[2])[0][2] == "on_click"
    assert not features.animations_of(slides[1])


def test_the_xml_order_of_a_slide_stays_valid() -> None:
    prs = _open(_deck({"transition": "fade"}))
    slide = list(prs.slides)[2]
    tags = [child.tag.split("}")[1] for child in slide._element]
    assert tags.index("cSld") < tags.index("clrMapOvr") < tags.index("transition") < tags.index("timing")


def test_animations_support_starts_and_can_be_cleared() -> None:
    prs = _open(_deck())
    slide = list(prs.slides)[1]
    first = next(s for s in slide.shapes if not s.is_placeholder)
    second = slide.shapes.title
    features.add_animation(slide, first, "wipe", "on_click")
    features.add_animation(slide, second, "appear", "after_previous")
    features.add_animation(slide, first, "fade", "with_previous")
    found = features.animations_of(slide)
    assert [(effect, start) for _, effect, start in found] == [
        ("wipe", "on_click"),
        ("appear", "after_previous"),
        ("fade", "with_previous"),
    ]
    ids = slide._element.xpath("./p:timing//p:cTn/@id")
    assert len(ids) == len(set(ids))  # every timing node has its own id
    assert features.clear_animations(slide) == 3
    assert not features.animations_of(slide)
    with pytest.raises(DeckError):
        features.add_animation(slide, first, "spin", "on_click")


def test_sections_validate_and_survive_slide_edits() -> None:
    prs = _open(_deck())
    with pytest.raises(DeckError, match="first section"):
        features.set_sections(prs, [features.Section("A", 2)])
    with pytest.raises(DeckError, match="rise"):
        features.set_sections(prs, [features.Section("A", 1), features.Section("B", 1)])
    apply_operations(prs, parse_operations([{"op": "delete_slide", "slide": 2}]))
    names = [(s.first_slide, s.name) for s in features.read_sections(prs)]
    assert names == [(1, "Start"), (2, "Part one"), (3, "Part two")]  # the deleted slide leaves its section
    apply_operations(prs, parse_operations([{"op": "set_sections", "sections": []}]))
    assert features.read_sections(prs) == []


def test_edit_operations_for_every_feature_and_duplicate_keeps_them() -> None:
    prs = _open(_deck())
    done = apply_operations(
        prs,
        parse_operations(
            [
                {"op": "hide_slide", "slide": 4},
                {"op": "set_transition", "kind": "wipe", "speed": "fast", "advance_after": 5},
                {"op": "animate", "slide": 3, "shape": "text:c", "effect": "fade"},
                {"op": "add_comment", "slide": 2, "text": "Check this number", "author": "Ada Lovelace"},
                {"op": "duplicate_slide", "slide": 3},
            ]
        ),
    )
    assert len(done) == 5
    slides = list(prs.slides)
    copy = slides[3]  # the copy of the hidden, animated slide 3 sits right after it
    assert (
        features.is_hidden(copy) and features.transition_of(copy) == "wipe" and features.animations_of(copy)
    )
    assert (
        slides[1]
        ._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}transition")
        .get("advTm")
        == "5000"
    )
    assert not validate_package(_save(prs))


def test_comments_are_added_read_back_and_survive_a_save() -> None:
    prs = _open(_deck())
    features.add_comment(prs, 2, "First <note> & more", "Ada Lovelace")
    features.add_comment(prs, 2, "Second", "Ada Lovelace")
    features.add_comment(prs, 3, "Other slide", "Grace Hopper")
    again = _open(_save(prs))
    comments = features.read_comments(again)
    assert [(c.slide, c.author, c.text) for c in comments] == [
        (2, "Ada Lovelace", "First <note> & more"),
        (2, "Ada Lovelace", "Second"),
        (3, "Grace Hopper", "Other slide"),
    ]
    assert not validate_package(_save(again))
    with pytest.raises(DeckError):
        features.add_comment(prs, 99, "x")


def test_reading_a_deck_shows_the_features() -> None:
    prs = _open(_deck())
    features.add_comment(prs, 2, "Look here", "Ada")
    text = describe_deck(prs)
    assert 'section starts: "Part one"' in text and "hidden: skipped" in text
    assert "transition: push" in text and "animation:" in text and "comment by Ada: Look here" in text


def test_a_slide_with_a_chart_can_be_duplicated_and_the_copy_is_independent() -> None:
    """P2.13 (D-319): the chart part and its workbook are cloned, so changing the copy leaves the original alone."""
    prs = _open(
        _deck(
            slides=[
                {"layout": "title", "title": "Charts"},
                {
                    "layout": "chart",
                    "title": "Users doubled",
                    "categories": ["Jan", "Feb"],
                    "series": [{"name": "Users", "values": [1, 2]}],
                },
            ]
        )
    )
    apply_operations(prs, parse_operations([{"op": "duplicate_slide", "slide": 2}]))
    assert len(prs.slides) == 3
    chart_id = next(
        s for s in prs.slides[2].shapes if getattr(s, "has_chart", False) and s.has_chart
    ).shape_id
    apply_operations(
        prs,
        parse_operations(
            [
                {
                    "op": "update_chart",
                    "slide": 3,
                    "shape": chart_id,
                    "categories": ["Jan", "Feb"],
                    "series": [{"name": "Users", "values": [10, 20]}],
                }
            ]
        ),
    )

    def values(slide_number: int) -> list[float]:
        chart = next(
            s for s in prs.slides[slide_number - 1].shapes if getattr(s, "has_chart", False) and s.has_chart
        ).chart
        return list(chart.plots[0].series[0].values)

    assert values(2) == [1, 2] and values(3) == [10, 20]
    saved = _save(prs)
    assert not validate_package(saved)
    names = {p.partname for p in _open(saved).part.package.iter_parts()}
    assert len([n for n in names if str(n).startswith("/ppt/charts/chart")]) == 2
    assert len([n for n in names if str(n).startswith("/ppt/embeddings/")]) == 2
