"""P2.14 (D-320): slides imported from another deck keep their design, decks merge and split, and the package stays valid."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches

from forge.slides import features
from forge.slides.build import build_deck
from forge.slides.edit import apply_operations, parse_operations
from forge.slides.importer import import_slides, merge_decks, split_deck
from forge.slides.open_deck import DeckError
from forge.slides.spec import parse_deck
from forge.slides.validate import validate_package


def _make(theme: str, title: str) -> Presentation:
    outline = {
        "title": title,
        "theme": theme,
        "slides": [
            {"layout": "title", "title": title},
            {
                "layout": "bullets",
                "title": f"Point of {title}",
                "bullets": ["a", "b"],
                "notes": f"Notes of {title}",
            },
            {
                "layout": "chart",
                "title": f"Chart of {title}",
                "categories": ["x", "y"],
                "series": [{"name": "s", "values": [1, 2]}],
            },
        ],
    }
    return Presentation(io.BytesIO(build_deck(parse_deck(json.dumps(outline)), lambda name: b"").data))


def _saved(prs: Presentation) -> bytes:
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _background(slide) -> str:  # type: ignore[no-untyped-def]
    return str(slide.background.fill.fore_color.rgb)


def test_merging_decks_with_different_looks_keeps_each_look_and_makes_a_valid_file() -> None:
    merged, notes = merge_decks([("alpha", _make("clean", "Alpha")), ("beta", _make("midnight", "Beta"))])
    data = _saved(merged)
    assert not validate_package(data) and notes == []
    again = Presentation(io.BytesIO(data))
    assert len(again.slides) == 6 and len(again.slide_masters) == 2
    assert _background(again.slides[0]) != _background(again.slides[3])
    assert (
        again.slides[3].slide_layout.slide_master.part is not again.slides[0].slide_layout.slide_master.part
    )
    assert [(s.name, s.first_slide) for s in features.read_sections(again)] == [("alpha", 1), ("beta", 4)]
    ids = [
        *again.part._element.xpath("./p:sldMasterIdLst/p:sldMasterId/@id"),
        *[i for m in again.slide_masters for i in m._element.xpath("./p:sldLayoutIdLst/p:sldLayoutId/@id")],
    ]
    assert len(ids) == len(set(ids))  # master and layout ids are unique across the deck


def test_the_same_look_is_not_imported_twice() -> None:
    merged, _ = merge_decks([("a", _make("clean", "Alpha")), ("c", _make("clean", "Gamma"))])
    assert len(merged.slide_masters) == 1 and len(merged.slides) == 6
    assert not validate_package(_saved(merged))


def test_charts_pictures_and_notes_come_along_and_are_independent_of_the_source() -> None:
    source = _make("midnight", "Beta")
    target = _make("clean", "Alpha")
    import_slides(target, source, [2, 3])
    assert target.slides[3].notes_slide.notes_text_frame.text == "Notes of Beta"
    chart = next(s for s in target.slides[4].shapes if getattr(s, "has_chart", False) and s.has_chart).chart
    assert list(chart.plots[0].series[0].values) == [1, 2]
    assert not validate_package(_saved(target))
    names = [str(p.partname) for p in Presentation(io.BytesIO(_saved(target))).part.package.iter_parts()]
    assert len(names) == len(set(names))
    assert len([n for n in names if n.startswith("/ppt/charts/chart")]) == 2


def test_import_at_a_position_and_hidden_and_transition_are_kept() -> None:
    source = _make("clean", "Beta")
    features.set_hidden(source.slides[1], True)
    features.set_transition(source.slides[1], "push")
    target = _make("clean", "Alpha")
    import_slides(target, source, [2], at=1)
    first = target.slides[0]
    assert features.is_hidden(first) and features.transition_of(first) == "push"
    assert (
        first.shapes.title.text_frame.text == "Point of Beta"
        and target.slides[1].shapes.title.text_frame.text == "Alpha"
    )


def test_a_different_slide_size_uses_the_target_layouts_and_scales_the_shapes() -> None:
    source = _make("clean", "Wide")
    target = Presentation()  # 4:3, 10 x 7.5 in
    notes = import_slides(target, source, [2], design="keep")
    assert any("slide sizes differ" in n for n in notes)
    slide = target.slides[0]
    assert all(
        s.left + s.width <= target.slide_width + Inches(0.05) for s in slide.shapes if s.left is not None
    )
    assert len(target.slide_masters) == 1


def test_bad_requests_say_what_is_wrong() -> None:
    target, source = _make("clean", "A"), _make("clean", "B")
    with pytest.raises(DeckError, match="no slide 9"):
        import_slides(target, source, [9])
    with pytest.raises(DeckError, match="keep"):
        import_slides(target, source, [1], design="copy")
    with pytest.raises(DeckError, match="at least two"):
        merge_decks([("a", target)])


def test_splitting_keeps_the_chosen_slides_and_their_sections() -> None:
    deck = _make("clean", "Alpha")
    features.set_sections(deck, [features.Section("One", 1), features.Section("Two", 3)])
    part = split_deck(deck, [1, 3])
    assert len(part.slides) == 2 and [s.name for s in features.read_sections(part)] == ["One", "Two"]
    assert not validate_package(_saved(part))
    with pytest.raises(DeckError):
        split_deck(_make("clean", "B"), [])


def test_the_edit_operation_imports_from_a_file_the_caller_opens() -> None:
    other = _make("midnight", "Beta")
    target = _make("clean", "Alpha")
    done = apply_operations(
        target,
        parse_operations([{"op": "import_slides", "source": "other.pptx", "slides": [2], "design": "keep"}]),
        None,
        lambda name: other,
    )
    assert "imported 1 slide(s) from other.pptx" in done[0] and len(target.slides) == 4
    with pytest.raises(DeckError, match="needs access"):
        apply_operations(target, parse_operations([{"op": "import_slides", "source": "x.pptx"}]))


async def test_the_merge_and_split_tools(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.presentation import MergePresentations, SplitPart, SplitPresentation
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    context = ToolContext(workspace=workspace)
    workspace.write_bytes("docs/a.pptx", _saved(_make("clean", "Alpha")), "x")
    workspace.write_bytes("docs/b.pptx", _saved(_make("midnight", "Beta")), "x")
    merged = await MergePresentations().run(
        MergePresentations.Args(paths=["docs/a.pptx", "docs/b.pptx"], output="docs/all.pptx"), context
    )
    assert merged.ok and "6 slides from 2 decks" in merged.content
    split = await SplitPresentation().run(
        SplitPresentation.Args(
            path="docs/all.pptx",
            parts=[
                SplitPart(output="docs/first.pptx", slides=[1, 2]),
                SplitPart(output="docs/second.pptx", slides=[4, 5, 6]),
            ],
        ),
        context,
    )
    assert split.ok and "docs/second.pptx (3 slides)" in split.content
    assert len(Presentation(str(workspace.path_of("docs/second.pptx"))).slides) == 3
