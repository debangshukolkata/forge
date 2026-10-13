"""P2.16 (D-317): a deck read out as Markdown, its pictures as files, and conversion through PowerPoint."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides import features
from forge.slides.build import build_deck
from forge.slides.extract import deck_markdown, extract_pictures
from forge.slides.preview import available
from forge.slides.spec import parse_deck


def _png() -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


def _deck() -> Presentation:
    outline = {
        "title": "Quarterly review",
        "slides": [
            {"layout": "title", "title": "Quarterly review", "notes": "Welcome"},
            {
                "layout": "bullets",
                "title": "Why it matters",
                "bullets": ["Revenue grew", "Costs fell"],
                "notes": "Say it slowly\nthen pause",
                "section": "Results",
                "hidden": True,
            },
            {
                "layout": "table",
                "title": "Regions",
                "columns": ["Region", "Revenue"],
                "rows": [["North", "1 | 2"], ["South", "3"]],
            },
            {
                "layout": "chart",
                "title": "Users doubled",
                "categories": ["Jan", "Feb"],
                "series": [{"name": "Users", "values": [1, 2]}],
            },
            {
                "layout": "image",
                "title": "The product",
                "image": "shot.png",
                "caption": "A screenshot",
                "alt": "Login screen",
            },
        ],
    }
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: _png())
    return Presentation(io.BytesIO(built.data))


def test_markdown_has_headings_bullets_tables_charts_pictures_notes_and_flags() -> None:
    prs = _deck()
    features.add_comment(prs, 3, "Check the numbers", "Ada")
    text = deck_markdown(prs)
    assert "## Slide 2: Why it matters" in text and "# Section: Results" in text
    assert "_hidden slide_" in text
    assert "- Revenue grew" in text and "- Costs fell" in text
    assert "> Notes: Say it slowly" in text and "> then pause" in text
    assert "| Region | Revenue |" in text and r"| North | 1 \| 2 |" in text  # a pipe in a cell is escaped
    assert "| Jan | 1 |" in text and "| Feb | 2 |" in text and "|  | Users |" in text
    assert "![Login screen](pictures/slide-5-" in text
    assert "> Comment by Ada: Check the numbers" in text
    assert "Quarterly review" in text.splitlines()[0]  # the document title
    assert "Slide 1" in text and text.count("## Slide") == 5


def test_footer_and_slide_numbers_are_left_out_of_the_markdown() -> None:
    text = deck_markdown(_deck())
    section = text.split("## Slide 2")[1].split("## Slide 3")[0]
    assert "Quarterly review" not in section  # the footer text of slide 2 is not content


def test_pictures_are_written_with_stable_names(tmp_path: Path) -> None:
    written = extract_pictures(_deck(), tmp_path / "out")
    assert len(written) == 1 and written[0].suffix == ".png" and written[0].name.startswith("slide-5-")
    assert written[0].read_bytes()[:4] == b"\x89PNG"


@pytest.mark.skipif(not available(), reason="PowerPoint is not installed here")
def test_conversion_through_powerpoint_makes_a_pdf_and_an_old_format_readable(tmp_path: Path) -> None:
    from forge.slides.convert import convert_deck

    source = tmp_path / "deck.pptx"
    out = io.BytesIO()
    _deck().save(out)
    source.write_bytes(out.getvalue())
    assert convert_deck(source, "pdf")[:4] == b"%PDF"
    assert convert_deck(source, "pptx")[:2] == b"PK"


async def test_the_tools_read_markdown_and_save_pictures(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.presentation import ConvertPresentation, ReadPresentation
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    context = ToolContext(workspace=workspace)
    out = io.BytesIO()
    _deck().save(out)
    workspace.write_bytes("docs/deck.pptx", out.getvalue(), "x")
    read = await ReadPresentation().run(ReadPresentation.Args(path="docs/deck.pptx", format="markdown"), context)
    assert read.ok and "## Slide 3: Regions" in read.content
    saved = await ConvertPresentation().run(
        ConvertPresentation.Args(path="docs/deck.pptx", to="pictures", output="docs/pics"), context
    )
    assert saved.ok and "Saved 1 picture(s)" in saved.content
    assert any(p.suffix == ".png" for p in workspace.path_of("docs/pics").iterdir())
