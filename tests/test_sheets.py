"""D-339: Excel workbooks are built from an outline, read, edited and checked (W2)."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from forge.sheets.build import build_workbook  # noqa: E402
from forge.sheets.check import check_workbook, describe_findings  # noqa: E402
from forge.sheets.edit import Operation, edit_workbook  # noqa: E402
from forge.sheets.read import describe_workbook, read_workbook  # noqa: E402
from forge.sheets.spec import WorkbookError, parse_workbook  # noqa: E402

OUTLINE = {
    "title": "Budget",
    "sheets": [
        {
            "name": "Assumptions",
            "columns": [{"header": "Item"}, {"header": "Value", "format": "0.0%"}],
            "rows": [["Tax rate", 0.2]],
        },
        {
            "name": "Sales",
            "columns": [
                {"header": "Region", "width": 18},
                {"header": "Q1", "format": "#,##0"},
                {"header": "Q2", "format": "#,##0"},
                {"header": "Tax", "format": "#,##0"},
            ],
            "rows": [
                ["North", 100, 120, "=B2*Assumptions!$B$2"],
                ["South", 80, 90, "=B3*Assumptions!$B$2"],
            ],
            "total_row": {"A": "Total", "B": "SUM", "C": "SUM"},
            "table": False,
            "validations": [{"range": "A2:A9", "choices": ["North", "South"]}],
            "conditional": [{"range": "B2:C3", "rule": "greater_than", "value": 100}],
            "charts": [
                {"type": "column", "title": "Sales", "data": "B1:C3", "categories": "A2:A3", "anchor": "G2"}
            ],
        },
    ],
    "names": {"TaxRate": "Assumptions!$B$2"},
}


def _built() -> bytes:
    return build_workbook(parse_workbook(json.dumps(OUTLINE)))


def test_build_keeps_formulas_formats_names_and_the_chart() -> None:
    data = _built()
    book = openpyxl.load_workbook(io.BytesIO(data))
    sales = book["Sales"]
    assert sales["D2"].value == "=B2*Assumptions!$B$2" and sales["B4"].value == "=SUM(B2:B3)"
    assert sales["B2"].number_format == "#,##0" and sales.freeze_panes == "A2"
    assert len(sales._charts) == 1 and len(sales.data_validations.dataValidation) == 1
    assert "TaxRate" in book.defined_names and book.properties.title == "Budget"
    assert sales.column_dimensions["A"].width == 18


def test_a_bad_outline_or_sheet_name_is_a_clear_error() -> None:
    with pytest.raises(WorkbookError):
        parse_workbook("{not json")
    with pytest.raises(WorkbookError, match="Two sheets"):
        build_workbook(parse_workbook(json.dumps({"sheets": [{"name": "A"}, {"name": "a"}]})))
    with pytest.raises(WorkbookError, match="Sheet name"):
        build_workbook(parse_workbook(json.dumps({"sheets": [{"name": "a/b"}]})))


def test_read_shows_formulas_and_says_when_the_file_was_never_calculated() -> None:
    info = read_workbook(_built())
    text = describe_workbook(info, "budget.xlsx")
    assert "Sheet 'Sales'" in text and "formula =B2*Assumptions!$B$2" in text
    assert "never calculated" in text and "TaxRate" in text and info.charts == 1


def test_check_finds_missing_sheet_circular_reference_and_typed_numbers() -> None:
    outline = {
        "sheets": [
            {
                "name": "A",
                "columns": [{"header": "x"}, {"header": "y"}],
                "rows": [[1, "=B2+1"], ["=Nope!A1", "=A3*1.25"], ["=A3", 2]],
            }
        ]
    }
    findings = check_workbook(build_workbook(parse_workbook(json.dumps(outline))))
    text = describe_findings(findings)
    assert "sheet that does not exist: 'Nope'" in text
    assert "hint" in text and "1.25" in text
    cyc = {"sheets": [{"name": "A", "rows": [["=B1", "=A1"]]}]}
    assert "circular reference" in describe_findings(
        check_workbook(build_workbook(parse_workbook(json.dumps(cyc))))
    )
    assert (
        describe_findings(check_workbook(_built())).startswith("No formula faults")
        or "error" not in describe_findings(check_workbook(_built())).split("\n")[0].split("error")[0]
    )


def test_edit_keeps_formulas_and_charts_and_refuses_a_real_loss() -> None:
    plain = build_workbook(parse_workbook(json.dumps({"sheets": [{"name": "S", "rows": [[1, "=A1*2"]]}]})))
    edited, done = edit_workbook(
        plain,
        [
            Operation(op="set_cells", sheet="S", cells={"A1": 5}),
            Operation(op="append_rows", sheet="S", rows=[[7, "=A2*2"]]),
            Operation(op="format", sheet="S", range="A1:B2", number_format="0.00", bold=True),
            Operation(op="add_sheet", name="New"),
        ],
    )
    book = openpyxl.load_workbook(io.BytesIO(edited))
    assert (
        book["S"]["B1"].value == "=A1*2" and book["S"]["A1"].value == 5 and book["S"]["B2"].value == "=A2*2"
    )
    assert book["S"]["A1"].number_format == "0.00" and "New" in book.sheetnames and len(done) == 4
    kept, notes = edit_workbook(_built(), [Operation(op="set_cells", sheet="Sales", cells={"B2": 1})])
    assert (
        "kept 1 chart(s)" in notes[-1] and len(openpyxl.load_workbook(io.BytesIO(kept))["Sales"]._charts) == 1
    )
    styled = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_built())) as source, zipfile.ZipFile(styled, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr(
            "xl/charts/style1.xml", "<cs:chartStyle/>"
        )  # Excel's style part, which openpyxl drops
    with pytest.raises(WorkbookError, match="Excel styling of 1 chart part"):
        edit_workbook(styled.getvalue(), [Operation(op="set_cells", sheet="Sales", cells={"B2": 1})])
    _, notes = edit_workbook(
        styled.getvalue(), [Operation(op="set_cells", sheet="Sales", cells={"B2": 1})], allow_loss=True
    )
    assert any("LOST" in n for n in notes)
    with pytest.raises(WorkbookError, match="No sheet called"):
        edit_workbook(plain, [Operation(op="set_cells", sheet="Zzz", cells={"A1": 1})])


async def test_the_tools(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.workbook import BuildWorkbook, CheckWorkbook, EditWorkbook, ReadWorkbook
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    workspace.write_text("docs/outline.json", json.dumps(OUTLINE), "x")
    context = ToolContext(workspace=workspace)
    built = await BuildWorkbook().run(
        BuildWorkbook.Args(outline="docs/outline.json", output="docs/b.xlsx"), context
    )
    assert built.ok and "Wrote docs/b.xlsx" in built.content
    read = await ReadWorkbook().run(ReadWorkbook.Args(path="docs/b.xlsx"), context)
    assert read.ok and "Sales" in read.content
    checked = await CheckWorkbook().run(CheckWorkbook.Args(path="docs/b.xlsx"), context)
    assert checked.ok
    edited = await EditWorkbook().run(
        EditWorkbook.Args(
            path="docs/b.xlsx", operations=[Operation(op="set_cells", sheet="Sales", cells={"B2": 1})]
        ),
        context,
    )
    assert edited.ok and "kept 1 chart(s)" in edited.content
    missing = await ReadWorkbook().run(ReadWorkbook.Args(path="docs/none.xlsx"), context)
    assert not missing.ok


def test_excel_recalculation_reports_real_errors_when_excel_is_installed() -> None:
    from forge.sheets import excel

    if not excel.available():
        pytest.skip("Excel is not installed here")
    outline = {
        "sheets": [
            {
                "name": "A",
                "columns": [{"header": "x"}, {"header": "y"}],
                "rows": [[1, "=A2/0"], [2, "=SUM(A2:A3)"]],
            }
        ]
    }
    errors, saved = excel.recalculate(build_workbook(parse_workbook(json.dumps(outline))))
    assert errors == [("A", "B2", "#DIV/0!")] and saved
    values = openpyxl.load_workbook(io.BytesIO(saved), data_only=True)["A"]
    assert values["B3"].value == 3


def test_numbers_that_formulas_depend_on_are_marked_as_inputs_and_unmarked_ones_are_reported() -> None:
    outline = {
        "names": {"Rate": "Assumptions!$B$2"},
        "sheets": [
            {
                "name": "Assumptions",
                "columns": [{"header": "Item"}, {"header": "Value"}],
                "rows": [["Rate", 0.2], ["Unused", 5]],
            },
            {
                "name": "Calc",
                "columns": [{"header": "Amount"}, {"header": "Tax"}],
                "rows": [[100, "=A2*Rate"], [250, "=A3*Rate"]],
            },
        ],
    }
    marked = openpyxl.load_workbook(io.BytesIO(build_workbook(parse_workbook(json.dumps(outline)))))
    assert str(marked["Assumptions"]["B2"].font.color.rgb).endswith("0000FF")  # read through the named range
    assert str(marked["Calc"]["A2"].font.color.rgb).endswith("0000FF")  # read directly by a formula
    assert marked["Assumptions"]["B3"].font.color is None or not str(
        marked["Assumptions"]["B3"].font.color.rgb
    ).endswith("0000FF")  # nothing reads it
    assert "not styled as inputs" not in describe_findings(
        check_workbook(build_workbook(parse_workbook(json.dumps(outline))))
    )
    plain = build_workbook(parse_workbook(json.dumps({**outline, "mark_inputs": False})))
    assert "not styled as inputs" in describe_findings(check_workbook(plain))


def _chart_book(sheet: dict) -> object:
    data = build_workbook(parse_workbook(json.dumps({"sheets": [sheet]})))
    return openpyxl.load_workbook(io.BytesIO(data))[sheet["name"]]


MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def test_a_row_of_monthly_values_is_one_series_not_twelve() -> None:
    """W6, doc-2 w06: a balance row with the month headers as categories was drawn as one series per month."""
    sheet = {
        "name": "Budget",
        "columns": [{"header": "Category"}] + [{"header": month} for month in MONTHS],
        "rows": [
            ["Income"] + [3000] * 12,
            ["Rent"] + [1200] * 12,
            ["Balance"] + [f"={c}2-{c}3" for c in "BCDEFGHIJKLM"],
        ],
        "charts": [
            {
                "type": "column",
                "title": "Monthly balance",
                "data": "B4:M4",
                "categories": "B1:M1",
                "anchor": "B6",
            }
        ],
    }
    chart = _chart_book(sheet)._charts[0]
    assert len(chart.series) == 1 and chart.varyColors is False and chart.legend is None
    assert chart.series[0].tx.v == "Balance"
    assert chart.series[0].cat.numRef is not None or chart.series[0].cat.strRef is not None


def test_one_measure_by_month_is_one_colour_and_a_straight_line() -> None:
    """W6, doc-2 w14: one series per chart, no colour per point, an unsmoothed line; the second chart moves off the first."""
    sheet = {
        "name": "Data",
        "columns": [{"header": "Month"}, {"header": "Revenue"}, {"header": "Customers"}],
        "rows": [[f"2026-0{n}", 40000 + n, 300 + n] for n in range(1, 7)],
        "charts": [
            {
                "type": "column",
                "title": "Revenue by month",
                "data": "B1:B7",
                "categories": "A2:A7",
                "anchor": "F2",
            },
            {
                "type": "line",
                "title": "Customers by month",
                "data": "C1:C7",
                "categories": "A2:A7",
                "anchor": "F10",
            },
        ],
    }
    column, line = _chart_book(sheet)._charts
    assert [len(column.series), len(line.series)] == [1, 1]
    assert column.varyColors is False and line.varyColors is False and line.series[0].smooth is False
    assert line.anchor._from.row > column.anchor._from.row + 15  # moved below, not on top


def test_categories_that_do_not_match_the_values_are_a_clear_error() -> None:
    sheet = {
        "name": "Data",
        "columns": [{"header": "Month"}, {"header": "Revenue"}],
        "rows": [["a", 1], ["b", 2], ["c", 3]],
        "charts": [{"type": "column", "data": "B1:B4", "categories": "A1:A4"}],
    }
    with pytest.raises(WorkbookError, match=r"4 categories .* but 3 value"):
        _chart_book(sheet)


def test_every_sheet_prints_one_page_wide_and_wide_sheets_turn_landscape() -> None:
    """W7, doc-2 w07 and w13: a table split across pages and a chart cut by the margin."""
    narrow = _chart_book({"name": "A", "columns": [{"header": "x"}], "rows": [[1]]})
    assert narrow.page_setup.fitToWidth == 1 and narrow.page_setup.fitToHeight == 0
    assert narrow.sheet_properties.pageSetUpPr.fitToPage and narrow.page_setup.orientation == "portrait"
    assert narrow.print_title_rows == "$1:$1"
    wide = _chart_book(
        {
            "name": "B",
            "columns": [{"header": "Month"}, {"header": "Revenue"}],
            "rows": [["a", 1], ["b", 2]],
            "charts": [{"type": "column", "data": "B1:B3", "categories": "A2:A3", "anchor": "F2"}],
        }
    )
    assert wide.page_setup.orientation == "landscape"


def test_the_check_finds_a_chart_drawn_one_series_per_value_however_the_file_was_made() -> None:
    """D-532, doc-3 w06: a model's own openpyxl script drew a monthly balance as twelve one-value series."""
    from openpyxl.chart import BarChart, Reference

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Balance"] + [100 + n for n in range(12)])
    chart = BarChart()
    chart.add_data(
        Reference(sheet, min_col=2, max_col=13, min_row=1, max_row=1)
    )  # columns, the wrong way for a row
    sheet.add_chart(chart, "B4")
    out = io.BytesIO()
    book.save(out)
    found = describe_findings(check_workbook(out.getvalue()))
    assert "12 series of one value each" in found and "print one page wide" in found
    assert "chart" not in describe_findings(check_workbook(_built()))


def test_long_text_wraps_in_a_narrow_column_and_unsized_columns_fit_their_content() -> None:
    """D-541, doc-4 w07 and w14: the model's own widths cut a task description and a KPI label off on the printout."""
    long_task = "Confirm the venue booking and send the floor plan to catering"
    sheet = _chart_book(
        {
            "name": "Tasks",
            "columns": [{"header": "Task", "width": 20}, {"header": "Owner"}, {"header": "Note"}],
            "rows": [
                [long_task, "Asha", "ok"],
                ["Short", "Bartholomew Featherstonehaugh", "ok"],
                ["=CONCATENATE(A2, A3, A2, A3, A2, A3, A2, A3)", "x", "=1"],
            ],
        }
    )
    assert sheet["A2"].alignment.wrap_text and not sheet["A3"].alignment.wrap_text
    assert sheet.column_dimensions["A"].width == 20  # the model's width is kept
    assert sheet.column_dimensions["B"].width >= len("Bartholomew Featherstonehaugh")
    assert sheet.column_dimensions["C"].width == 10 and not sheet["C2"].alignment.wrap_text
    assert not sheet["A4"].alignment.wrap_text  # a formula is not text to wrap


def test_constant_date_formulas_in_a_date_column_become_real_dates() -> None:
    """D-541, doc-4b w07: the model wrote every due date as =DATE(2026,11,4); a reader of the file saw formulas, not dates."""
    from datetime import datetime

    sheet = _chart_book(
        {
            "name": "Tasks",
            "columns": [{"header": "Due", "format": "yyyy-mm-dd"}, {"header": "Other"}],
            "rows": [
                ["=DATE(2026,11,4)", "=DATE(2026,11,4)"],
                ["2026-11-06", "x"],
                ["=DATE(2026,13,1)", "x"],
                ["=TODAY()", "x"],
            ],
        }
    )
    assert sheet["A2"].value == datetime(2026, 11, 4) and sheet["A3"].value == datetime(2026, 11, 6)
    assert sheet["B2"].value == "=DATE(2026,11,4)"  # not a date column: left alone
    assert sheet["A4"].value == "=DATE(2026,13,1)" and sheet["A5"].value == "=TODAY()"


def test_a_chart_on_a_summary_sheet_can_draw_another_sheets_numbers_and_an_empty_range_is_caught() -> None:
    """D-543, doc-5b w14: the charts sat on the KPI sheet, their ranges did not name the Data sheet, and they drew empty cells."""
    months = [[month, 100 + n] for n, month in enumerate(MONTHS[:6])]
    summary = {
        "name": "KPI",
        "title": "Summary",
        "columns": [{"header": "KPI"}, {"header": "Value"}],
        "rows": [["Best", "=MAX(Data!B2:B7)"]],
        "charts": [{"type": "column", "data": "Data!B1:B7", "categories": "Data!A2:A7", "anchor": "D2"}],
    }
    data_sheet = {
        "name": "Data",
        "title": "Monthly data",
        "columns": [{"header": "Month"}, {"header": "Revenue"}],
        "rows": months,
    }
    made = build_workbook(
        parse_workbook(json.dumps({"sheets": [summary, data_sheet]}))
    )  # the chart's sheet comes first
    chart = openpyxl.load_workbook(io.BytesIO(made))["KPI"]._charts[0]
    assert chart.series[0].val.numRef.f == "'Data'!$B$5:$B$10"  # moved below the Data sheet's title too
    categories = (
        chart.series[0].cat.numRef or chart.series[0].cat.strRef
    )  # openpyxl writes numRef for any range
    assert categories.f == "'Data'!$A$5:$A$10"
    assert "holds no numbers" not in describe_findings(check_workbook(made))

    summary["charts"] = [
        {"type": "column", "data": "C2:C7", "categories": "A2:A7", "anchor": "D2"}
    ]  # as in w14
    blank = build_workbook(parse_workbook(json.dumps({"sheets": [summary, data_sheet]})))
    assert "holds no numbers" in describe_findings(check_workbook(blank))
    with pytest.raises(WorkbookError, match="names a sheet that does not exist"):
        summary["charts"] = [{"type": "column", "data": "Sales!B1:B7", "anchor": "D2"}]
        build_workbook(parse_workbook(json.dumps({"sheets": [summary, data_sheet]})))


def test_the_total_row_is_one_band_with_rules_and_formatted_numbers() -> None:
    """D-556, doc-5c w02: the judge called the total row unfinished: only the named cells bold, no rule, sums in General."""
    sheet = _chart_book(
        {
            "name": "Orders",
            "columns": [{"header": "Item"}, {"header": "Qty"}, {"header": "Amount"}, {"header": "Note"}],
            "rows": [["Kettle", 2, 19.5], ["Toaster", 1, 24.0]],
            "total_row": {"A": "Total", "B": "SUM", "C": "SUM"},
        }
    )
    row = [sheet.cell(row=4, column=column) for column in range(1, 5)]
    assert all(
        cell.font.bold and cell.fill.fgColor.rgb.endswith("F2F2F2") for cell in row
    )  # the Note column too
    assert all(cell.border.top.style == "thin" and cell.border.bottom.style == "double" for cell in row)
    assert (
        row[1].number_format == "#,##0" and row[2].number_format == "#,##0.00"
    )  # whole numbers, then money-like


def test_a_negative_bar_keeps_its_colour_and_the_labels_sit_below_the_chart() -> None:
    """D-556, doc-6 w02: a negative amount was drawn as a white outlined box and 'Toaster' was printed inside it."""
    sheet = _chart_book(
        {
            "name": "Orders",
            "columns": [{"header": "Product"}, {"header": "Amount"}],
            "rows": [["Kettle", 1200.5], ["Toaster", -300], ["Blender", 950]],
            "charts": [{"type": "column", "data": "B1:B4", "categories": "A2:A4", "anchor": "D2"}],
        }
    )
    chart = sheet._charts[0]
    assert all(series.invertIfNegative is False for series in chart.series)
    assert chart.x_axis.tickLblPos == "low"
