"""P2.7 (D-315): a generated deck has a real master, theme and footer placeholders, so it re-themes in PowerPoint."""

from __future__ import annotations

import io
import json

import pytest

pytest.importorskip("pptx")
from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.spec import parse_deck
from forge.slides.themes import THEMES
from forge.slides.validate import validate_package

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _build(theme: str = "clean") -> Presentation:
    outline = {
        "title": "Quarterly review",
        "theme": theme,
        "footer": "Acme internal",
        "slides": [
            {"layout": "title", "title": "Quarterly review"},
            {"layout": "bullets", "title": "Point", "bullets": ["a", "b"]},
        ],
    }
    built = build_deck(parse_deck(json.dumps(outline)), lambda name: b"")
    assert not validate_package(built.data)
    return Presentation(io.BytesIO(built.data))


def _theme_xml(prs: Presentation):  # type: ignore[no-untyped-def]
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE

    return etree.fromstring(prs.slide_master.part.part_related_by(RELATIONSHIP_TYPE.THEME).blob)


@pytest.mark.parametrize("name", sorted(THEMES))
def test_the_theme_part_carries_the_deck_colours_and_fonts(name: str) -> None:
    theme = THEMES[name]
    xml = _theme_xml(_build(name))
    colour = lambda slot: xml.xpath(f"//a:clrScheme/a:{slot}/a:srgbClr/@val", namespaces=NS)[0]  # noqa: E731
    assert colour("dk1") == theme.text.upper() and colour("lt1") == theme.background.upper()
    assert colour("accent1") == theme.palette[0].lstrip("#").upper()
    assert xml.xpath("//a:majorFont/a:latin/@typeface", namespaces=NS) == [theme.title_font]
    assert xml.xpath("//a:minorFont/a:latin/@typeface", namespaces=NS) == [theme.body_font]


def test_layouts_are_sixteen_by_nine_and_only_the_useful_four_remain() -> None:
    prs = _build()
    assert [layout.name for layout in prs.slide_layouts] == [
        "Title Slide",
        "Title and Content",
        "Title Only",
        "Blank",
    ]
    for shape in prs.slide_master.placeholders:
        assert shape.left + shape.width <= prs.slide_width
    title = prs.slide_master.placeholders[0]
    assert title.width > prs.slide_width * 0.8  # not the 4:3 width of the Office template


def test_footer_and_slide_number_are_real_placeholders_that_follow_the_master() -> None:
    prs = _build()
    slide = prs.slides[1]
    kinds = {shape.placeholder_format.type.name: shape for shape in slide.placeholders}
    assert kinds["FOOTER"].text_frame.text == "Acme internal"
    assert kinds["SLIDE_NUMBER"].text_frame.text == "2"
    assert kinds["SLIDE_NUMBER"]._element.xpath(".//a:fld[@type='slidenum']")
    master_footer = next(
        s for s in prs.slide_master.placeholders if s.placeholder_format.type.name == "FOOTER"
    )
    assert kinds["FOOTER"].top == master_footer.top  # no position of its own


def test_the_master_has_the_theme_background_and_text_colours_by_scheme() -> None:
    prs = _build("clean")
    master = prs.slide_master._element
    assert master.xpath("//p:bg//a:schemeClr/@val") == ["bg1"]
    assert master.xpath("//p:titleStyle//a:schemeClr/@val") == ["tx1"]
