"""P2.17 (D-318): the review of a deck finds what can be measured, and a well-built deck passes without problems."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from forge.slides.build import build_deck
from forge.slides.review import review_deck, summarize
from forge.slides.spec import parse_deck


def _messages(findings: list) -> str:  # type: ignore[type-arg]
    return " | ".join(f.render() for f in findings)


def _bad_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for number in range(5):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text_frame.text = "Overview" if number else "Plan"
        slide.shapes.title.left = Inches(0.5 + number * 0.4)  # jumps around
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
        frame = box.text_frame
        for index in range(9):  # nine bullets of many words
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = "word " * 12
            run.font.size = Pt(10)
            run.font.name = f"Font {number}"
            run.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    off = prs.slides[0].shapes.add_textbox(Inches(12), Inches(6), Inches(3), Inches(1))
    off.text_frame.text = "off the slide"
    return prs


def test_a_bad_deck_gets_findings_in_every_area() -> None:
    text = _messages(review_deck(_bad_deck(), minutes=2))
    for needle in (
        "is a label, not a message",
        "is used on 4 slides",
        "bullets in one box",
        "words on the slide",
        "font families",
        "text as small as 10 pt",
        "against its background",
        "reaches past the edge",
        "title is not where it is",
        "no speaker notes",
        "more than one slide a minute",
        "the document has no title",
    ):
        assert needle in text, needle


def test_the_summary_lists_problems_first_with_a_fix_on_each_line() -> None:
    summary = summarize(review_deck(_bad_deck()))
    lines = summary.splitlines()
    assert lines[0].startswith("Review: ") and "problem(s)" in lines[0]
    assert all(" -> " in line for line in lines[1:])
    first_suggestion = next(
        i for i, line in enumerate(lines) if ": suggestion:" in line or "] suggestion:" in line
    )
    assert all("] problem:" in line for line in lines[1:first_suggestion])


def test_a_deck_built_by_debcraft_has_no_problems() -> None:
    outline = {
        "title": "Why we should move to the new platform",
        "slides": [
            {"layout": "title", "title": "Why we should move to the new platform"},
            {
                "layout": "bullets",
                "title": "The old platform costs us a day a week",
                "bullets": ["Releases take two days", "Incidents repeat", "Hiring is hard"],
                "notes": "Say how the day a week was counted, then ask for questions.",
            },
            {
                "layout": "chart",
                "title": "Release time fell after the pilot",
                "categories": ["Before", "After"],
                "series": [{"name": "Days", "values": [2, 0.5]}],
                "notes": "The pilot ran for six weeks on two teams, as agreed.",
            },
            {"layout": "closing", "title": "Decide by the end of the month"},
        ],
    }
    prs = Presentation(io.BytesIO(build_deck(parse_deck(json.dumps(outline)), lambda n: b"").data))
    problems = [f for f in review_deck(prs) if f.level == "problem"]
    assert problems == [], _messages(problems)


async def test_the_tool_reviews_a_file(original_repo: Path, tmp_path: Path) -> None:
    from forge.toolkit.base import ToolContext
    from forge.tools.presentation import ReviewPresentation
    from forge.workspace.create import create_workspace

    workspace = create_workspace(original_repo, tmp_path / "ws", "backend")
    out = io.BytesIO()
    _bad_deck().save(out)
    workspace.write_bytes("docs/bad.pptx", out.getvalue(), "x")
    result = await ReviewPresentation().run(
        ReviewPresentation.Args(path="docs/bad.pptx"), ToolContext(workspace=workspace)
    )
    assert result.ok and "problem(s)" in result.content and "bullets in one box" in result.content
