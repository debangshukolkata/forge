"""P1.4: rich text inside one paragraph, links, line and letter spacing, vertical text."""

from __future__ import annotations

import io
import json

from pptx import Presentation

from forge.slides.build import build_deck
from forge.slides.richtext import RunSpec, parse_runs, plain_text
from forge.slides.spec import parse_deck
from forge.slides.validate import validate_package


def test_markup_becomes_runs() -> None:
    runs = parse_runs(
        "Revenue grew **42%** in *Q3*, see [the report](https://example.org/q3) and {accent}note{/}."
    )
    assert runs == [
        RunSpec("Revenue grew "),
        RunSpec("42%", bold=True),
        RunSpec(" in "),
        RunSpec("Q3", italic=True),
        RunSpec(", see "),
        RunSpec("the report", link="https://example.org/q3"),
        RunSpec(" and "),
        RunSpec("note", colour="accent"),
        RunSpec("."),
    ]
    assert plain_text("**a** and {#C62828}b{/}") == "a and b"


def test_nested_and_unmatched_markup() -> None:
    assert parse_runs("{accent}**big** news{/}") == [
        RunSpec("big", bold=True, colour="accent"),
        RunSpec(" news", colour="accent"),
    ]
    assert plain_text("2 * 3 = 6 and a*b") == "2 * 3 = 6 and a*b"  # stray asterisks stay literal
    assert parse_runs("") == [RunSpec("")]


def _built(element: dict) -> tuple[object, list[str]]:  # type: ignore[type-arg]
    deck = {
        "title": "Deck",
        "slides": [{"layout": "freeform", "title": "Rich text works inside one line", "elements": [element]}],
    }
    built = build_deck(parse_deck(json.dumps(deck)), lambda n: b"")
    return Presentation(io.BytesIO(built.data)), [w for slide in built.slides for w in slide.warnings]


def _paragraph(presentation: object):  # type: ignore[no-untyped-def]
    slide = presentation.slides[0]  # type: ignore[attr-defined]
    box = next(s for s in slide.shapes if s.has_text_frame and "Revenue" in s.text_frame.text)
    return box, box.text_frame.paragraphs[0]


def test_a_rich_line_is_several_runs_with_their_own_style_and_link() -> None:
    element = {
        "type": "text",
        "x": 1,
        "y": 2,
        "w": 8,
        "h": 1,
        "rich": True,
        "text": "Revenue **grew** {accent}fast{/} [here](https://example.org)",
    }
    presentation, warnings = _built(element)
    _, paragraph = _paragraph(presentation)
    runs = paragraph.runs
    assert [r.text for r in runs] == ["Revenue ", "grew", " ", "fast", " ", "here"]
    assert runs[1].font.bold and not runs[0].font.bold
    assert runs[3].font.color.rgb is not None and runs[3].font.color.rgb != runs[0].font.color.rgb
    assert runs[5].hyperlink.address == "https://example.org"
    assert warnings == []


def test_rich_is_off_by_default_so_asterisks_are_kept() -> None:
    presentation, _ = _built({"type": "text", "x": 1, "y": 2, "w": 8, "h": 1, "text": "Revenue **grew**"})
    _, paragraph = _paragraph(presentation)
    assert [r.text for r in paragraph.runs] == ["Revenue **grew**"]


def test_spacing_and_vertical_text() -> None:
    element = {
        "type": "text",
        "x": 1,
        "y": 1,
        "w": 1,
        "h": 4,
        "text": "Revenue",
        "line_spacing": 1.5,
        "letter_spacing": 2,
        "vertical": "up",
        "size": 20,
    }
    presentation, _ = _built(element)
    box, paragraph = _paragraph(presentation)
    assert paragraph.line_spacing == 1.5
    assert paragraph.runs[0]._r.rPr.get("spc") == "200"
    assert box.text_frame._txBody.bodyPr.get("vert") == "vert270"


def test_an_unknown_colour_name_is_a_warning_not_a_crash() -> None:
    element = {"type": "text", "x": 1, "y": 2, "w": 8, "h": 1, "rich": True, "text": "Revenue {nonsense}x{/}"}
    _, warnings = _built(element)
    assert any("not a theme colour" in w for w in warnings)


def test_a_rich_deck_is_a_valid_package() -> None:
    deck = {
        "title": "Deck",
        "slides": [
            {
                "layout": "freeform",
                "title": "Rich text works inside one line",
                "elements": [
                    {
                        "type": "text",
                        "x": 1,
                        "y": 2,
                        "w": 8,
                        "h": 1,
                        "rich": True,
                        "text": "Revenue **grew** [here](https://example.org)",
                        "letter_spacing": 1,
                    }
                ],
            }
        ],
    }
    assert validate_package(build_deck(parse_deck(json.dumps(deck)), lambda n: b"").data) == []
